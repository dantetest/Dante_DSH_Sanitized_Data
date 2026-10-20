"""DOCX 解析器与破坏性导出器。

parse
    用 python-docx 打开文档，把正文段落、表格（单元格）、内容控件、页眉页脚
    归一化成 UDM 的 text Block；表格另有一个 kind="table" 的 Block 供前端排版。
    **Block.text 严格等于该段落 ``"".join(run.text for run in p.runs)``**，
    保证检测出的字符区间在导出时可以原样映射回 run。

export
    复制原文件到成品路径，在副本上用 python-docx 真正改写文字（破坏性重写）：
    - ``black``：用 “█” 覆盖原长度，原文字彻底消失；
    - ``surrogate`` / ``mask``：直接写入 replacement，长度可变；
    - ``drop``：删除该区间文字；
    随后清空 core properties、docProps/app.xml 里的公司/标题信息，删除批注、
    修订节点、自定义属性、CustomXML、缩略图等隐藏部件（zip 级重写），再保存。

    额外做一次「值级传播」：检测器漏掉的 block 里若仍出现同一实体值，也按同一
    决策覆盖（宁可多涂不可漏），并写 warning 说明覆盖处数。

已知边界（会写 warning，不静默丢弃）：
    - python-docx 的 ``Paragraph.runs`` 不含超链接内的 run，超链接显示文字不参与
      检测；域代码、文本框、脚注尾注、嵌套表格、内嵌对象/图片同样只告警。
    - python-docx 没有自定义属性 API，``docProps/custom.xml`` 在 zip 级删除。
"""
from __future__ import annotations

import os
import posixpath
import re
import shutil
import tempfile
import time
import uuid
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any

from PIL import Image, ImageDraw

from ..registry import Parser, register_parser
from ..udm import (
    Block, Cell, Decision, Document, Entity, ExportError, ExportResult,
    Page, ParseContext,
)

# ── 可选依赖：环境里必装，但缺失时降级成 warning 而不是让插件加载失败 ──
try:  # pragma: no cover - 环境已装
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _DOCX_OK = True
    _DOCX_ERR = ""
except Exception as _exc:  # pragma: no cover
    docx = None  # type: ignore[assignment]
    qn = None  # type: ignore[assignment]
    Table = None  # type: ignore[assignment]
    Paragraph = None  # type: ignore[assignment]
    _DOCX_OK = False
    _DOCX_ERR = repr(_exc)

try:  # pragma: no cover - 环境已装
    from lxml import etree

    _LXML_OK = True
except Exception:  # pragma: no cover
    etree = None  # type: ignore[assignment]
    _LXML_OK = False

A4_WIDTH = 595.0
A4_HEIGHT = 842.0
BLACK_CHAR = "█"

# 修订 / 批注相关节点：直接删除（w:ins 里的内容属于未接受的修订，一并删除更安全）
_REVISION_TAGS = frozenset(
    {
        "ins", "del", "moveFrom", "moveTo",
        "commentRangeStart", "commentRangeEnd", "commentReference",
        "rPrChange", "pPrChange", "sectPrChange", "tblPrChange", "trPrChange",
        "tcPrChange", "numberingChange",
        "moveFromRangeStart", "moveFromRangeEnd",
        "moveToRangeStart", "moveToRangeEnd",
        "cellIns", "cellDel", "cellMerge",
        "customXmlInsRangeStart", "customXmlInsRangeEnd",
        "customXmlDelRangeStart", "customXmlDelRangeEnd",
        "customXmlMoveFromRangeStart", "customXmlMoveFromRangeEnd",
        "customXmlMoveToRangeStart", "customXmlMoveToRangeEnd",
    }
)

# settings.xml 里可能泄漏本机路径 / 口令哈希 / 修订记录的节点
_SETTINGS_DROP_TAGS = frozenset(
    {"documentProtection", "writeProtection", "attachedTemplate",
     "mailMerge", "docVars", "trackChanges", "rsids"}
)

# zip 级需要删除的部件（小写匹配）
_REMOVE_PART_EXACT = frozenset(
    {
        "word/comments.xml", "word/commentsextended.xml", "word/commentsids.xml",
        "word/commentsextensible.xml", "word/people.xml",
        "docprops/custom.xml", "docprops/thumbnail.jpeg", "docprops/thumbnail.jpg",
        "docprops/thumbnail.png", "docprops/thumbnail.wmf", "docprops/thumbnail.emf",
    }
)
_REMOVE_PART_RE = (
    re.compile(r"^word/comments.*\.xml$"),
    re.compile(r"^docprops/thumbnail\..*$"),
    re.compile(r"^customxml/.*$"),
)

# 需要从中提取作者/公司等元数据的 app.xml 标签
_APP_META_TAGS = {
    "Company": "company",
    "Manager": "manager",
    "Template": "template",
    "HyperlinkBase": "hyperlinkbase",
}
_CORE_META_FIELDS = (
    "author", "last_modified_by", "title", "subject", "keywords",
    "comments", "category", "content_status", "identifier", "language",
)


# ══════════════════════════════ 小工具 ══════════════════════════════

def _localname(tag: Any) -> str:
    """去掉 XML 命名空间，只留本地名。"""
    text = tag if isinstance(tag, str) else str(tag)
    if text.startswith("{"):
        return text.rsplit("}", 1)[-1]
    return text


def _paragraph_text(paragraph: Any) -> str:
    """段落文字：严格用 run 文本拼接（与导出偏移一致）。"""
    return "".join((run.text or "") for run in paragraph.runs)


def _safe_style_name(paragraph: Any) -> str:
    try:
        return str(paragraph.style.name or "")
    except Exception:
        return ""


def _cell_text(cell: Any) -> str:
    """单元格文字：各段落 run 文本用 \\n 连接（导出按同样方式换算偏移）。"""
    return "\n".join(_paragraph_text(p) for p in cell.paragraphs)


def _fmt_meta_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    try:
        return str(value)
    except Exception:
        return ""


def _replacement_text(decision: Decision, length: int) -> str | None:
    """把 Decision 翻译成实际写入的字符串；keep 返回 None。"""
    if decision.strategy == "keep":
        return None
    if decision.strategy == "drop":
        return ""
    if decision.black or decision.strategy == "black" or decision.replacement is None:
        return BLACK_CHAR * max(0, length)
    return decision.replacement


def _propagation_rules(
    entities: list[Entity], decisions: dict[str, Decision], warnings: list[str]
) -> dict[str, str]:
    """值级传播规则：同一实体值在检测范围外的其它出现也按同一决策覆盖。

    启发式检测器可能漏掉页眉、表格、批注锚点等 block，而出口自检按「成品里
    是否还有该值」判定残留；这里做一次防御性传播，宁可多涂也不漏。
    单字符值不传播（会把整篇文档涂花），只记 warning。
    """
    rules: dict[str, str] = {}
    short: set[str] = set()
    for entity in entities:
        if entity.status == "rejected" or not entity.value:
            continue
        decision = decisions.get(entity.id) or entity.decision
        if decision is None:
            decision = Decision(strategy="black", black=True)
        if decision.strategy == "keep":
            continue
        if len(entity.value) < 2:
            short.add(entity.value)
            continue
        replacement = _replacement_text(decision, len(entity.value))
        if replacement is None:
            continue
        existing = rules.get(entity.value)
        if existing is not None and existing != replacement:
            warnings.append(
                f"同一值「{entity.value[:12]}」存在不同的替换决策，已按涂黑处理"
            )
            rules[entity.value] = BLACK_CHAR * len(entity.value)
            continue
        rules[entity.value] = replacement
    if short:
        warnings.append(f"有 {len(short)} 个单字符实体值未做值级传播，请人工复核")
    return rules


def _propagate_in_blocks(
    block_by_id: dict[str, Block],
    tasks: dict[str, list[tuple[int, int, str]]],
    rules: dict[str, str],
) -> int:
    """把值级传播补进替换任务表（按原始 block 偏移），返回新增处数。"""
    added = 0
    ordered = sorted(rules.items(), key=lambda item: len(item[0]), reverse=True)
    for block_id, block in block_by_id.items():
        if block.kind != "text" or not block.text:
            continue
        text = block.text
        occupied = tasks.get(block_id, [])
        additions: list[tuple[int, int, str]] = []
        for value, replacement in ordered:
            start = 0
            while True:
                index = text.find(value, start)
                if index < 0:
                    break
                end = index + len(value)
                overlaps = any(
                    not (end <= old_start or index >= old_end)
                    for old_start, old_end, _old in occupied
                ) or any(
                    not (end <= old_start or index >= old_end)
                    for old_start, old_end, _old in additions
                )
                if not overlaps:
                    additions.append((index, end, replacement))
                start = end
        if additions:
            tasks.setdefault(block_id, []).extend(additions)
            added += len(additions)
    return added


def _remove_element(element: Any) -> None:
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _replace_file(source: Path, target: Path) -> None:
    """原子替换；Windows 上成品可能被杀软/索引器短暂占用，失败后退化为覆盖写。"""
    for attempt in range(6):
        try:
            os.replace(source, target)
            return
        except OSError as exc:
            if attempt == 5:
                last_error = exc
                break
            time.sleep(0.05 * (attempt + 1))
    try:
        with open(source, "rb") as reader, open(target, "wb") as writer:
            shutil.copyfileobj(reader, writer)
        source.unlink(missing_ok=True)
    except Exception as exc:
        raise OSError(f"替换 {target} 失败：{last_error}；覆盖写也失败：{exc}") from exc


def _strip_revisions(root: Any) -> int:
    """删除 root 下所有修订/批注标记节点，返回删除个数。"""
    removed = 0
    for element in list(root.iter()):
        try:
            tag = _localname(element.tag)
        except Exception:
            continue
        if tag in _REVISION_TAGS:
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)
                removed += 1
    return removed


def _clean_settings(document: Any) -> int:
    """清掉 settings.xml 里的保护口令、附加模板、邮件合并数据源等。"""
    try:
        settings = document.settings.element
    except Exception:
        return 0
    removed = 0
    for element in list(settings.iter()):
        if _localname(element.tag) in _SETTINGS_DROP_TAGS:
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)
                removed += 1
    return removed


# ══════════════════════════════ 导出：run 级替换 ══════════════════════════════

def _replace_in_paragraph(paragraph: Any, start: int, end: int, new: str) -> bool:
    """把段落字符区间 [start, end) 换成 new，必要时拆分/合并 run。

    段落文本按 ``"".join(run.text for run in paragraph.runs)`` 计算；替换后的新文字
    写入区间起始 run，保留该 run 的字体格式，其余覆盖到的 run 删除或截断。
    """
    runs = list(paragraph.runs)
    if not runs:
        return False
    texts = [(run.text or "") for run in runs]
    total = sum(len(text) for text in texts)
    start = max(0, min(int(start), total))
    end = max(0, min(int(end), total))
    if start > end:
        start, end = end, start

    # 纯插入：找插入点所在 run，在它内部插入
    if start == end:
        if not new:
            return True
        acc = 0
        for run, text in zip(runs, texts):
            length = len(text)
            if start <= acc + length:
                offset = start - acc
                run.text = text[:offset] + new + text[offset:]
                return True
            acc += length
        # 落在段落末尾
        runs[-1].text = texts[-1] + new
        return True

    # 定位起始 run i（覆盖 start 的 run）
    acc = 0
    i = -1
    i_off = 0
    for idx, text in enumerate(texts):
        length = len(text)
        if length and start < acc + length:
            i, i_off = idx, start - acc
            break
        acc += length
    if i < 0:
        return False

    # 定位结束 run j（覆盖 end-1 的 run）
    acc = 0
    j = -1
    j_off = 0
    for idx, text in enumerate(texts):
        length = len(text)
        if length and end <= acc + length:
            j, j_off = idx, end - acc
            break
        acc += length
    if j < 0 or j < i:
        return False

    prefix = texts[i][:i_off]
    suffix = texts[j][j_off:]
    if i == j:
        runs[i].text = prefix + new + suffix
        return True

    # 删除 (i, j) 之间的 run，再处理首尾
    for idx in range(j - 1, i, -1):
        _remove_element(runs[idx]._element)
    runs[i].text = prefix + new
    runs[j].text = suffix
    return True


def _replace_in_cell(cell: Any, start: int, end: int, new: str) -> bool:
    """把单元格（多段落以 \\n 连接）的字符区间换成 new。"""
    paragraphs = list(cell.paragraphs)
    if not paragraphs:
        return False
    texts = [_paragraph_text(p) for p in paragraphs]
    total = sum(len(t) for t in texts) + max(0, len(texts) - 1)
    start = max(0, min(int(start), total))
    end = max(0, min(int(end), total))
    if start > end:
        start, end = end, start

    pos = 0
    placed = False
    for paragraph, text in zip(paragraphs, texts):
        length = len(text)
        p_start, p_end = pos, pos + length
        pos = p_end + 1  # 段落间的 "\n" 占一个字符
        if end < p_start or start > p_end:
            continue
        lo = max(start, p_start)
        hi = min(end, p_end)
        if hi < lo:
            continue
        if not placed:
            local_start = 0 if start <= p_start else lo - p_start
            local_end = hi - p_start
            _replace_in_paragraph(paragraph, local_start, local_end, new)
            placed = True
        else:
            _replace_in_paragraph(paragraph, 0, hi - p_start, "")
    return placed


# ══════════════════════════════ zip 级加固 ══════════════════════════════

def _resolve_rel_target(rels_name: str, target: str) -> str:
    """把 .rels 里的相对 Target 解析成包内绝对部件名。"""
    rel_name = rels_name.replace("\\", "/")
    if target.startswith("/"):
        return posixpath.normpath(target.lstrip("/"))
    owner_dir = posixpath.dirname(posixpath.dirname(rel_name))
    return posixpath.normpath(posixpath.join(owner_dir, target))


def _clean_rels(name: str, data: bytes, removed: set[str]) -> tuple[bytes, int, int]:
    """删除指向已移除部件的 Relationship；外部链接关系保留（打印告警）。"""
    if not _LXML_OK:
        return data, 0, 0
    try:
        root = etree.fromstring(data)
    except Exception:
        return data, 0, 0
    removed_lower = {item.lower() for item in removed}
    dropped = 0
    external = 0
    for rel in list(root):
        if _localname(rel.tag) != "Relationship":
            continue
        target = rel.get("Target") or ""
        if rel.get("TargetMode") == "External":
            external += 1
            continue
        try:
            resolved = _resolve_rel_target(name, target).lower()
        except Exception:
            continue
        if resolved in removed_lower:
            root.remove(rel)
            dropped += 1
    if not dropped:
        return data, 0, external
    try:
        out = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    except Exception:
        return data, 0, external
    return out, dropped, external


def _clean_content_types(data: bytes, removed: set[str]) -> tuple[bytes, int]:
    """删除 [Content_Types].xml 里指向已移除部件的 Override。"""
    if not _LXML_OK:
        return data, 0
    try:
        root = etree.fromstring(data)
    except Exception:
        return data, 0
    removed_lower = {item.lower() for item in removed}
    dropped = 0
    for element in list(root):
        if _localname(element.tag) != "Override":
            continue
        part_name = (element.get("PartName") or "").lstrip("/").lower()
        if part_name in removed_lower:
            root.remove(element)
            dropped += 1
    if not dropped:
        return data, 0
    try:
        out = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    except Exception:
        return data, 0
    return out, dropped


def _clean_app_xml(data: bytes) -> bytes:
    """清空 docProps/app.xml 里的公司/经理/模板/超链接基址，删除标题大纲缓存。"""
    if not _LXML_OK:
        return data
    try:
        root = etree.fromstring(data)
    except Exception:
        return data
    changed = False
    for element in list(root):
        tag = _localname(element.tag)
        if tag in _APP_META_TAGS:
            if (element.text or "").strip():
                element.text = ""
                changed = True
        elif tag in ("TitlesOfParts", "HeadingPairs"):
            root.remove(element)
            changed = True
    if not changed:
        return data
    try:
        return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    except Exception:
        return data


def _harden_zip(path: Path, warnings: list[str]) -> None:
    """删除 python-docx 无法管理的敏感部件，并清理 app.xml / rels / 内容类型。"""
    if not _LXML_OK:
        warnings.append("缺少 lxml，无法做 zip 级元数据清理（自定义属性/批注/缩略图可能残留）")
        return
    try:
        with zipfile.ZipFile(path) as archive:
            names = archive.namelist()
    except Exception as exc:
        warnings.append(f"zip 级元数据清理失败（文件可能仍含批注/自定义属性）：{exc}")
        return

    removed: set[str] = set()
    for name in names:
        low = name.lower()
        if low in _REMOVE_PART_EXACT or any(p.match(low) for p in _REMOVE_PART_RE):
            removed.add(name)

    tmp_path = path.with_name(path.name + ".harden.tmp")
    try:
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(
            tmp_path, "w", zipfile.ZIP_DEFLATED
        ) as target:
            rels_dropped = 0
            external = 0
            overrides_dropped = 0
            for item in source.infolist():
                name = item.filename
                if name in removed:
                    continue
                data = source.read(name)
                if name.endswith(".rels"):
                    data, dropped, ext = _clean_rels(name, data, removed)
                    rels_dropped += dropped
                    external += ext
                elif name == "[Content_Types].xml":
                    data, dropped = _clean_content_types(data, removed)
                    overrides_dropped += dropped
                elif name == "docProps/app.xml":
                    data = _clean_app_xml(data)
                target.writestr(item, data)
        _replace_file(tmp_path, path)
    except Exception as exc:  # 清理失败不能留下半个文件
        warnings.append(f"zip 级元数据清理失败：{exc}")
        try:
            tmp_path.unlink(missing_ok=True)
        except Exception:
            pass
        return

    if removed:
        warnings.append(f"已删除 {len(removed)} 个隐藏部件（批注/自定义属性/缩略图/CustomXML）")
    if external:
        warnings.append(f"文档含 {external} 个外部链接关系，可能暴露本机路径，已告警")
    if rels_dropped or overrides_dropped:
        warnings.append(
            f"已同步清理 {rels_dropped} 条失效关系、{overrides_dropped} 条内容类型声明"
        )


# ═══════════════════ Office 内嵌媒体自动脱敏（P1-B，DOCX/XLSX 共用） ═══════════════════
#
# 背景：P0 只检测已加载的 media/embeddings 并让用户二次确认，放行后原图保留。
# P1-B：
#   1) 按 zip namelist 全量枚举（media/embeddings/charts/drawings/activeX/…
#      + 文本框/域/脚注尾注），写 doc.meta["embedded_media"] / ["hidden_objects"]；
#   2) 导出时对可解码图片自动跑 mediascan（face/id/plate/qr/barcode/chat/sign），
#      有命中就 PIL 涂黑 → 同格式重编码 → zip 级替换（同一部件名，OOXML 结构不变）；
#   3) 图表/OLE/activeX/文本框/不可解码图片（emf/wmf/svg…）无法安全自动脱敏：
#      默认保留并需要用户确认（count>0 触发既有 media_confirmation_required）；
#      strict_office_media=True 时尽量剥离（删部件 + 清 rels/内容类型 + 去 XML 引用）；
#   4) 导出后重开成品 zip，对被替换图片重扫 mediascan，命中必须 0，否则 ExportError。
# 非目标（P1.5）：内嵌图片里的纯文字 OCR 规则脱敏、图片 EXIF/GPS 元数据剥离。

_OFFICE_IMAGE_EXTS = frozenset(
    {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".webp"})
_OFFICE_UNSUPPORTED_IMAGE_EXTS = frozenset(
    {".emf", ".wmf", ".svg", ".ico", ".tga", ".eps", ".psd", ".heic", ".heif",
     ".avif", ".jxl", ".pcx", ".pict", ".pct"})


def _env_flag(name: str, default: bool) -> bool:
    raw = os.environ.get(name)
    if raw is None or str(raw).strip() == "":
        return bool(default)
    return str(raw).strip().lower() not in ("0", "false", "no", "off", "none")


def _env_positive_int(name: str, default: int) -> int:
    try:
        value = int(os.environ.get(name, "") or default)
        return value if value > 0 else int(default)
    except (TypeError, ValueError):
        return int(default)


#: 单次导出最多自动处理多少张图片（防超大文档拖垮导出；超出会明确 warning）
OFFICE_MEDIA_MAX_IMAGES = _env_positive_int("YINZHAN_OFFICE_MEDIA_MAX_IMAGES", 60)
#: 单张图片像素上限（超过则跳过自动处理并 warning）
OFFICE_MEDIA_MAX_PIXELS = _env_positive_int("YINZHAN_OFFICE_MEDIA_MAX_PIXELS", 40_000_000)
#: 单次导出累计处理像素上限
OFFICE_MEDIA_TOTAL_PIXELS = _env_positive_int(
    "YINZHAN_OFFICE_MEDIA_TOTAL_PIXELS", 300_000_000)
#: 扫描前缩放上限（Barcode 需要细节，过大时缩放会漏检 → 会 warning）
OFFICE_MEDIA_SCAN_MAX_SIDE = _env_positive_int("YINZHAN_OFFICE_SCAN_MAX_SIDE", 3000)
#: 每次导出最多为多少张图片跑 OCR（身份证/车牌/聊天截图类目标需要 OCR 文本框）
OFFICE_MEDIA_OCR_MAX_IMAGES = _env_positive_int("YINZHAN_OFFICE_OCR_MAX_IMAGES", 30)
#: 是否默认自动脱敏内嵌图片（导出 option 优先；env 可整体关闭）
OFFICE_MEDIA_REDACT_DEFAULT = _env_flag("YINZHAN_OFFICE_MEDIA", True)
#: 是否默认 strict 处理（导出 option 优先）
OFFICE_MEDIA_STRICT_DEFAULT = _env_flag("YINZHAN_OFFICE_STRICT", False)

# ── P1.5 内嵌图片「纯文字」OCR 规则脱敏（media_text.redact_image_text）──
# 只对 helper 支持的位图格式调用；预算按“整篇文档”统计，超预算保留原图 + warning。
OFFICE_TEXT_SUFFIXES = (".png", ".jpg", ".jpeg", ".bmp", ".webp")
#: 文字 OCR 总开关（默认开；`YINZHAN_OFFICE_TEXT_OCR=0` 可整体关闭）
OFFICE_TEXT_OCR_DEFAULT = _env_flag("YINZHAN_OFFICE_TEXT_OCR", True)


def _office_text_enabled() -> bool:
    """每次导出时动态读取文字 OCR 开关（env 可覆盖，测试/排障用）。"""
    return _env_flag("YINZHAN_OFFICE_TEXT_OCR", OFFICE_TEXT_OCR_DEFAULT)


def _office_text_budget() -> tuple[int, float]:
    """单次导出的文字 OCR 总预算：(最多图片数, 最长秒数)；0 秒表示不限。"""
    max_images = _env_positive_int("YINZHAN_OFFICE_TEXT_MAX_IMAGES", 30)
    try:
        max_seconds = float(os.environ.get("YINZHAN_OFFICE_TEXT_MAX_SECONDS", "") or 60.0)
    except (TypeError, ValueError):
        max_seconds = 60.0
    return max_images, max(0.0, max_seconds)


def _office_text_verify_seconds() -> float:
    """导出后文字自检的时间预算（秒）；0 表示不限。"""
    try:
        return max(0.0, float(
            os.environ.get("YINZHAN_OFFICE_TEXT_VERIFY_SECONDS", "") or 20.0))
    except (TypeError, ValueError):
        return 20.0


def _office_redact_text_bytes(data: bytes, part: str,
                              workdir: Path | None) -> tuple[bytes | None, int, list[str]]:
    """调用 Lead 提供的 media_text helper；异常由调用方兜底。"""
    from .. import media_text

    suffix = Path(str(part)).suffix.lower()
    return media_text.redact_image_text(data, suffix=suffix, workdir=workdir)


def _office_media_option(options: dict | None, key: str, default: bool) -> bool:
    raw = (options or {}).get(key)
    if raw is None:
        return bool(default)
    if isinstance(raw, str):
        return raw.strip().lower() not in ("0", "false", "no", "off", "none")
    return bool(raw)


def _xml_has_textbox(data: bytes) -> bool:
    """判断 drawing/容器 XML 里是否有文本框/形状（无法安全自动脱敏的隐藏对象）。"""
    low = data[:200000].lower()
    if b"txbxcontent" in low or b"txbx" in low or b"txbody" in low:
        return True
    if not _LXML_OK:
        return False
    try:
        root = etree.fromstring(data)
    except Exception:
        return False
    for element in root.iter():
        try:
            if _localname(element.tag).lower() in ("txbxcontent", "txbx", "txbody"):
                return True
        except Exception:
            continue
    return False


def _count_office_textbox_elements(archive: zipfile.ZipFile, fmt: str) -> int:
    """统计包内 XML 里的文本框元素数量（docx: w:txbxContent；xlsx: xdr:txBody）。"""
    if not _LXML_OK:
        return 0
    targets: list[str] = []
    for name in archive.namelist():
        low = name.lower()
        if fmt == "docx":
            if low.startswith(("word/document.xml", "word/header", "word/footer",
                               "word/footnotes.xml", "word/endnotes.xml")):
                if low.endswith(".xml"):
                    targets.append(name)
        else:
            if low.startswith("xl/drawings/") and low.endswith(".xml"):
                targets.append(name)
    total = 0
    for name in targets:
        try:
            data = archive.read(name)
        except Exception:
            continue
        low = data[:200000].lower()
        if fmt == "docx":
            total += low.count(b"txbxcontent")
        else:
            # 只把“形状/文本框”算隐藏对象；纯图片锚点（xdr:pic）不算
            if b"txbody" in low or b"<xdr:sp" in low:
                try:
                    root = etree.fromstring(data)
                    total += sum(
                        1 for element in root.iter()
                        if _localname(element.tag).lower() in ("txbody", "sp"))
                except Exception:
                    total += max(1, low.count(b"txbody"))
    return total


def _classify_office_parts(archive: zipfile.ZipFile, fmt: str) -> dict[str, Any]:
    """按 zip 全量 namelist 分类：可自动脱敏图片 / 不可解码图片 / 隐藏对象。"""
    names = [str(name) for name in archive.namelist()]
    prefix = "word/" if fmt == "docx" else "xl/"
    media_prefix = prefix + "media/"
    images: list[str] = []
    unsupported: list[str] = []
    hidden: list[str] = []
    footnotes: list[str] = []
    endnotes: list[str] = []
    for name in names:
        low = name.lower()
        base = low.rsplit("/", 1)[-1]
        ext = ("." + base.rsplit(".", 1)[-1]) if "." in base else ""
        if low.startswith(media_prefix):
            if ext in _OFFICE_IMAGE_EXTS:
                images.append(name)
            else:
                unsupported.append(name)
            continue
        if "/embeddings/" in low:
            hidden.append(name)
            continue
        if low.startswith(prefix + "charts/"):
            hidden.append(name)
            continue
        if "activex" in low:
            hidden.append(name)
            continue
        if fmt == "docx":
            if low.endswith("word/glossary/document.xml"):
                hidden.append(name)
            elif low.startswith("word/drawings/") and low.endswith(".xml"):
                try:
                    if _xml_has_textbox(archive.read(name)):
                        hidden.append(name)
                except Exception:
                    continue
            elif base == "footnotes.xml":
                footnotes.append(name)
            elif base == "endnotes.xml":
                endnotes.append(name)
        else:
            if low.startswith("xl/drawings/") and low.endswith(".xml"):
                try:
                    if _xml_has_textbox(archive.read(name)):
                        hidden.append(name)
                except Exception:
                    continue
            elif low.startswith("xl/externallinks/"):
                hidden.append(name)
    try:
        textbox_count = _count_office_textbox_elements(archive, fmt)
    except Exception:
        textbox_count = 0
    return {
        "images": images,
        "unsupported_images": unsupported,
        "hidden_objects": hidden,
        "footnotes": footnotes,
        "endnotes": endnotes,
        "textbox_count": int(textbox_count),
        "parts": images + unsupported + hidden,
    }


def _build_office_media_meta(classified: dict[str, Any], *, fmt: str,
                             fields: int = 0) -> dict[str, Any]:
    """构造 doc.meta["embedded_media"]；空文档返回 {}。"""
    images = list(classified.get("images") or [])
    unsupported = list(classified.get("unsupported_images") or [])
    hidden = list(classified.get("hidden_objects") or []) + unsupported
    textbox_count = int(classified.get("textbox_count") or 0)
    if textbox_count:
        hidden.append(f"textbox:{'word/document.xml' if fmt == 'docx' else 'xl/drawings/*'}#{textbox_count}")
    # 去重但保持顺序
    hidden = list(dict.fromkeys(str(item) for item in hidden))
    redact_enabled = bool(OFFICE_MEDIA_REDACT_DEFAULT)
    if not redact_enabled:
        # 全局关闭自动脱敏时：图片全部计入确认数（否则会静默保留原图）
        hidden = list(dict.fromkeys(hidden + images))
    footnotes = list(classified.get("footnotes") or [])
    endnotes = list(classified.get("endnotes") or [])
    if not any((images, hidden, textbox_count, fields, footnotes, endnotes)):
        return {}
    return {
        "count": len(hidden),
        "parts": hidden[:50],
        "images": images[:50],
        "image_count": len(images),
        "hidden_objects": hidden[:50],
        "hidden_object_count": len(hidden),
        "auto_redact": redact_enabled,
        "unsupported_images": unsupported[:20],
        "textboxes": textbox_count,
        "fields": int(fields or 0),
        "footnotes": len(footnotes),
        "endnotes": len(endnotes),
        # P1.5：图片内纯文字会走 media_text.redact_image_text（OCR+规则+黑块）
        "text_ocr_supported": _office_text_enabled(),
    }


def _append_office_media_warnings(document: Document, meta: dict[str, Any],
                                  *, fmt: str = "docx") -> None:
    if not meta:
        return
    label = "文档" if fmt == "docx" else "工作簿"
    image_count = int(meta.get("image_count") or 0)
    hidden_count = int(meta.get("hidden_object_count") or 0)
    if image_count:
        document.warnings.append(
            f"{label}含 {image_count} 张内嵌图片，导出时将自动扫描并涂黑人脸/证件/"
            "车牌/二维码/条码/聊天截图等目标，并对图片内文字做 OCR 规则脱敏（P1.5）")
    if hidden_count:
        document.warnings.append(
            f"{label}含 {hidden_count} 个无法自动脱敏的隐藏对象"
            "（图表/文本框/OLE/activeX/不可解码图片等），导出前需确认；"
            "strict_office_media=True 可尝试剥离")
    if meta.get("unsupported_images"):
        document.warnings.append(
            f"{label}含 {len(meta['unsupported_images'])} 张不可解码图片"
            "（emf/wmf/svg 等），已计入隐藏对象")
    if int(meta.get("textboxes") or 0):
        document.warnings.append(
            f"{label}含 {meta['textboxes']} 个文本框（文字未纳入检测），已计入隐藏对象")
    if int(meta.get("fields") or 0):
        document.warnings.append(
            f"{label}含 {meta['fields']} 处域代码（如 HYPERLINK/INCLUDETEXT），未纳入检测")


def _office_apply_orientation(image: Image.Image) -> Image.Image:
    """按 EXIF Orientation 转正（内嵌手机照片常见）；失败原样返回。"""
    try:
        from .image_parser import _apply_orientation

        return _apply_orientation(image)
    except Exception:
        return image


def _office_scan_pil(pil_image: Image.Image, *, ocr: Any,
                     warnings: list[str]) -> tuple[list[dict], list[str], float]:
    """对一张内嵌图片跑 mediascan；返回 (marks 全图坐标, warnings, 缩放比例)。"""
    from .. import mediascan

    if not mediascan.has_enabled():
        return [], warnings, 1.0
    try:
        rgb = pil_image.convert("RGB")
    except Exception as exc:
        warnings.append(f"内嵌图片转 RGB 失败：{exc}")
        return [], warnings, 1.0
    width, height = rgb.size
    scan_scale = 1.0
    scan_image = rgb
    if max(width, height) > OFFICE_MEDIA_SCAN_MAX_SIDE:
        scan_scale = OFFICE_MEDIA_SCAN_MAX_SIDE / float(max(width, height))
        scan_image = rgb.resize(
            (max(1, int(round(width * scan_scale))), max(1, int(round(height * scan_scale)))),
            Image.LANCZOS)
        warnings.append(
            f"内嵌图片 {width}×{height} 超过扫描上限 {OFFICE_MEDIA_SCAN_MAX_SIDE}px，"
            "已缩放后扫描（二维码/条码等细节可能漏检）")
    text_boxes: list[dict] = []
    if ocr is not None:
        try:
            if getattr(ocr, "available", False):
                lines = ocr.recognize(scan_image, min_score=0.35) or []
                for line in lines:
                    box = getattr(line, "box", None)
                    text = str(getattr(line, "text", "") or "")
                    if box:
                        text_boxes.append({"text": text, "rect": box})
        except Exception as exc:
            warnings.append(f"内嵌图片 OCR 失败：{exc}")
    try:
        marks, scan_warns = mediascan.analyze(
            scan_image, page=0, dpi=0.0,
            options={"text_boxes": text_boxes, "fmt": "image",
                     "context": "office_embedded"})
    except Exception as exc:
        warnings.append(f"内嵌图片扫描失败：{exc}")
        return [], warnings, scan_scale
    warnings.extend(scan_warns)
    if scan_scale != 1.0 and marks:
        inv = 1.0 / scan_scale
        for mark in marks:
            try:
                mark["rect"] = tuple(float(v) * inv for v in mark["rect"])
            except Exception:
                continue
    return marks, warnings, scan_scale


def _office_paint_marks(pil_image: Image.Image, marks: list[dict],
                        warnings: list[str]) -> tuple[Image.Image, list[str]]:
    """在图片上按命中框涂黑（不透明），返回 (新图, 命中类型列表)。"""
    if not marks:
        return pil_image, []
    try:
        if pil_image.mode == "RGBA":
            canvas = pil_image
            fill = (0, 0, 0, 255)
        elif pil_image.mode in ("RGB", "L"):
            canvas = pil_image
            fill = 0 if pil_image.mode == "L" else (0, 0, 0)
        else:
            canvas = pil_image.convert("RGB")
            fill = (0, 0, 0)
    except Exception as exc:
        warnings.append(f"内嵌图片涂黑准备失败：{exc}")
        return pil_image, []
    draw = ImageDraw.Draw(canvas)
    width, height = canvas.size
    kinds: list[str] = []
    for mark in marks:
        try:
            x0, y0, x1, y1 = (float(v) for v in (mark.get("rect") or ())[:4])
        except Exception:
            continue
        pad = 2
        x0 = max(0, x0 - pad)
        y0 = max(0, y0 - pad)
        x1 = min(width, x1 + pad)
        y1 = min(height, y1 + pad)
        if x1 <= x0 or y1 <= y0:
            continue
        try:
            draw.rectangle([int(x0), int(y0), int(x1) - 1, int(y1) - 1], fill=fill)
            kinds.append(str(mark.get("kind") or "?"))
        except Exception:
            continue
    return canvas, kinds


def _office_encode_image(pil_image: Image.Image, part_name: str) -> bytes:
    """按原扩展名重编码（不携带 EXIF/ICC）；失败抛 ValueError。"""
    ext = Path(str(part_name)).suffix.lower()
    buffer = BytesIO()
    image = pil_image
    try:
        if ext in (".jpg", ".jpeg"):
            if image.mode not in ("RGB", "L"):
                image = image.convert("RGB")
            image.save(buffer, format="JPEG", quality=92, optimize=True)
        elif ext == ".png":
            image.save(buffer, format="PNG", optimize=True)
        elif ext == ".gif":
            image.convert("P", palette=Image.ADAPTIVE).save(buffer, format="GIF")
        elif ext == ".bmp":
            image.convert("RGB").save(buffer, format="BMP")
        elif ext == ".webp":
            image.save(buffer, format="WEBP", quality=92)
        elif ext in (".tif", ".tiff"):
            image.save(buffer, format="TIFF", compression="tiff_lzw")
        else:
            raise ValueError(f"不支持的图片扩展名 {ext}")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"重编码失败：{exc}") from exc
    data = buffer.getvalue()
    if not data:
        raise ValueError("重编码结果为空")
    return data


def _office_blank_image(pil_image: Image.Image, part_name: str) -> bytes:
    """strict 模式：把无命中图片置为纯黑（保留尺寸/格式，内容彻底不可见）。"""
    mode = "RGBA" if pil_image.mode == "RGBA" else "RGB"
    color = (0, 0, 0, 255) if mode == "RGBA" else (0, 0, 0)
    blank = Image.new(mode, pil_image.size, color)
    return _office_encode_image(blank, part_name)


def _office_strip_strict_xml(data: bytes, warnings: list[str]) -> tuple[bytes, int]:
    """strict：删除 XML 里的文本框/OLE/图表引用元素（best-effort）。"""
    if not _LXML_OK:
        return data, 0
    try:
        root = etree.fromstring(data)
    except Exception:
        return data, 0
    drop_tags = {"txbxcontent", "object", "oleobject", "chart"}
    removed = 0
    for element in list(root.iter()):
        try:
            if _localname(element.tag).lower() in drop_tags:
                parent = element.getparent()
                if parent is not None:
                    parent.remove(element)
                    removed += 1
        except Exception:
            continue
    if not removed:
        return data, 0
    try:
        return (etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                               standalone=True), removed)
    except Exception as exc:
        warnings.append(f"strict XML 清理失败：{exc}")
        return data, 0


def _office_zip_rewrite(path: Path, *, replacements: dict[str, bytes],
                        removals: set[str], xml_transforms: list | None,
                        warnings: list[str]) -> int:
    """zip 级替换图片/删除部件，并同步清理 .rels 与 [Content_Types].xml。"""
    if not replacements and not removals and not xml_transforms:
        return 0
    tmp_path = path.with_name(path.name + ".media.tmp")
    rels_dropped = 0
    overrides_dropped = 0
    xml_changed = 0
    removed_lower = {str(name).lower() for name in removals}
    try:
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(
                tmp_path, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                name = item.filename
                low = name.lower()
                if low in removed_lower:
                    continue
                if name in replacements:
                    target.writestr(item, replacements[name])
                    continue
                try:
                    data = source.read(name)
                except Exception:
                    data = b""
                if name.endswith(".rels") and removals:
                    data, dropped, _external = _clean_rels(name, data, removals)
                    rels_dropped += dropped
                elif name == "[Content_Types].xml" and removals:
                    data, dropped = _clean_content_types(data, removals)
                    overrides_dropped += dropped
                elif xml_transforms:
                    for predicate, transform in xml_transforms:
                        try:
                            if predicate(low):
                                data, changed = transform(data, warnings)
                                xml_changed += changed
                                break
                        except Exception:
                            continue
                target.writestr(item, data)
        _replace_file(tmp_path, path)
    except Exception as exc:
        try:
            tmp_path.unlink(missing_ok=True)
        except Exception:
            pass
        raise ExportError(f"Office 内嵌媒体 zip 重写失败：{exc}") from exc
    if rels_dropped or overrides_dropped:
        warnings.append(
            f"strict 剥离：已清理 {rels_dropped} 条关系、{overrides_dropped} 条内容类型声明")
    return xml_changed


def _office_verify_images(path: Path, parts: list[str], *, ocr: Any,
                          warnings: list[str]) -> list[str]:
    """导出后自检：重新提取被替换图片再跑 mediascan，命中必须为 0。"""
    issues: list[str] = []
    if not parts:
        return issues
    try:
        archive = zipfile.ZipFile(path)
    except Exception as exc:
        return [f"成品 zip 无法打开：{exc}"]
    try:
        for part in parts:
            try:
                raw = archive.read(part)
            except Exception as exc:
                issues.append(f"{part} 无法读取：{exc}")
                continue
            try:
                with Image.open(BytesIO(raw)) as image:
                    image.load()
                    pil = _office_apply_orientation(image)
            except Exception as exc:
                issues.append(f"{part} 无法解码：{exc}")
                continue
            marks, _scan_warns, _scale = _office_scan_pil(pil, ocr=ocr, warnings=[])
            if marks:
                kinds = sorted({str(m.get("kind") or "?") for m in marks})
                issues.append(f"{part} 仍有 {len(marks)} 个敏感目标（{','.join(kinds)}）")
    finally:
        try:
            archive.close()
        except Exception:
            pass
    return issues


def _office_process_media(path: Path, *, fmt: str, options: dict | None,
                          warnings: list[str]) -> dict[str, Any]:
    """导出时自动脱敏内嵌图片 +（strict）剥离隐藏对象；含导出后自检。"""
    options = options or {}
    redact = _office_media_option(options, "redact_office_media",
                                  OFFICE_MEDIA_REDACT_DEFAULT)
    strict = _office_media_option(options, "strict_office_media",
                                  OFFICE_MEDIA_STRICT_DEFAULT)
    stats: dict[str, Any] = {
        "images_total": 0, "images_redacted": 0, "images_clean": 0,
        "images_blanked": 0, "images_skipped": 0, "images_failed": 0,
        "hits_by_kind": {}, "hidden_total": 0, "hidden_stripped": 0,
        "hidden_kept": 0, "redact_office_media": redact,
        "strict_office_media": strict,
        # P1.5 文字 OCR 统计
        "text_images_processed": 0, "text_images_redacted": 0,
        "text_hits": 0, "text_skipped": 0,
        "text_verify_checked": 0, "text_verify_hits": 0, "text_verify_skipped": 0,
    }
    try:
        archive = zipfile.ZipFile(path)
    except Exception as exc:
        warnings.append(f"Office 内嵌媒体处理跳过（zip 无法打开）：{exc}")
        return stats
    replacements: dict[str, bytes] = {}
    replaced_parts: list[str] = []
    removals: set[str] = set()
    text_workdir: Path | None = None
    try:
        classified = _classify_office_parts(archive, fmt)
        images = list(classified.get("images") or [])
        hidden_parts = list(classified.get("hidden_objects") or [])
        unsupported = list(classified.get("unsupported_images") or [])
        textbox_count = int(classified.get("textbox_count") or 0)
        stats["images_total"] = len(images)
        stats["hidden_total"] = len(hidden_parts) + len(unsupported) + textbox_count
        ocr = None
        if redact and images:
            try:
                from ..ocr import get_ocr

                ocr = get_ocr()
                if not getattr(ocr, "available", False):
                    warnings.append("OCR 不可用：内嵌图片只做人脸/二维码/条码类扫描")
                    ocr = None
            except Exception as exc:
                warnings.append(f"OCR 服务不可用：内嵌图片只做视觉/码类扫描（{exc}）")
                ocr = None
        ocr_budget = OFFICE_MEDIA_OCR_MAX_IMAGES
        total_pixels = 0
        # P1.5 文字 OCR：整篇文档预算（图片数/耗时），超预算保留原图 + warning
        text_max_images, text_max_seconds = _office_text_budget()
        text_verify_seconds = _office_text_verify_seconds()
        text_used = 0
        text_elapsed = 0.0
        text_budget_warned = False
        text_enabled = _office_text_enabled()
        text_replaced_parts: list[str] = []
        for index, part in enumerate(images):
            if index >= OFFICE_MEDIA_MAX_IMAGES:
                skipped = len(images) - index
                stats["images_skipped"] += skipped
                warnings.append(
                    f"内嵌图片超过自动处理上限 {OFFICE_MEDIA_MAX_IMAGES} 张，"
                    f"剩余 {skipped} 张未处理，请人工复核")
                break
            try:
                raw = archive.read(part)
            except Exception as exc:
                stats["images_failed"] += 1
                warnings.append(f"内嵌图片 {part} 读取失败，未脱敏：{exc}")
                continue
            try:
                with Image.open(BytesIO(raw)) as image:
                    image.load()
                    pil = _office_apply_orientation(image)
                    pil = pil.copy()
            except Exception as exc:
                stats["images_failed"] += 1
                unsupported.append(part)
                stats["hidden_total"] += 1
                warnings.append(f"内嵌图片 {part} 无法解码，已计入隐藏对象：{exc}")
                continue
            pixels = int(pil.width) * int(pil.height)
            if (pixels > OFFICE_MEDIA_MAX_PIXELS
                    or total_pixels + pixels > OFFICE_MEDIA_TOTAL_PIXELS):
                stats["images_skipped"] += 1
                warnings.append(
                    f"内嵌图片 {part} 尺寸 {pil.width}×{pil.height} 超过自动处理上限，"
                    "未脱敏，请人工复核")
                continue
            total_pixels += pixels
            allow_ocr = ocr is not None and ocr_budget > 0
            if allow_ocr:
                ocr_budget -= 1
            marks, _scan_warns, _scale = _office_scan_pil(
                pil, ocr=ocr if allow_ocr else None, warnings=warnings)
            visual_kinds: list[str] = []
            visual_encoded: bytes | None = None
            if marks:
                painted, visual_kinds = _office_paint_marks(pil, marks, warnings)
                try:
                    visual_encoded = _office_encode_image(painted, part)
                except Exception as exc:
                    stats["images_failed"] += 1
                    warnings.append(f"内嵌图片 {part} 重编码失败，未脱敏：{exc}")
                    continue
                replacements[part] = visual_encoded
            # ── P1.5：图片内纯文字 OCR 规则脱敏（在视觉涂黑之后，叠加执行）──
            text_hits = 0
            suffix = Path(str(part)).suffix.lower()
            if redact and suffix in OFFICE_TEXT_SUFFIXES:
                if not text_enabled:
                    stats["text_skipped"] += 1
                elif (text_max_images > 0 and text_used >= text_max_images) or (
                        text_max_seconds > 0 and text_elapsed >= text_max_seconds):
                    stats["text_skipped"] += 1
                    if not text_budget_warned:
                        text_budget_warned = True
                        warnings.append(
                            "内嵌图片文字 OCR 预算已用尽"
                            f"（{text_max_images} 张/{text_max_seconds:.0f}s），"
                            "剩余图片只做视觉/码类脱敏，文字内容需人工复核")
                else:
                    if text_workdir is None:
                        text_workdir = Path(tempfile.mkdtemp(prefix="yz_office_text_"))
                    candidate = visual_encoded if visual_encoded is not None else raw
                    started = time.time()
                    try:
                        text_new, text_hits, text_warns = _office_redact_text_bytes(
                            candidate, part, text_workdir)
                    except Exception as exc:
                        text_new, text_hits, text_warns = None, 0, [
                            f"内嵌图片 {part} 文字 OCR 失败：{exc}"]
                    text_elapsed += max(0.0, time.time() - started)
                    text_used += 1
                    stats["text_images_processed"] += 1
                    warnings.extend(str(item) for item in (text_warns or []))
                    if text_new is not None:
                        replacements[part] = text_new
                        text_replaced_parts.append(part)
                        stats["text_hits"] += int(text_hits or 0)
                        if text_hits:
                            stats["text_images_redacted"] += 1
            # 决定最终成品图片字节
            if part in replacements:
                if part not in replaced_parts:
                    replaced_parts.append(part)
                if visual_kinds:
                    stats["images_redacted"] += 1
                    for kind in visual_kinds:
                        stats["hits_by_kind"][kind] = (
                            stats["hits_by_kind"].get(kind, 0) + 1)
            elif strict:
                try:
                    replacements[part] = _office_blank_image(pil, part)
                    replaced_parts.append(part)
                    stats["images_blanked"] += 1
                except Exception as exc:
                    stats["images_failed"] += 1
                    warnings.append(f"内嵌图片 {part} strict 置空失败：{exc}")
            else:
                stats["images_clean"] += 1
        # 隐藏对象：strict 剥离，默认保留并 warning
        for part in hidden_parts + unsupported:
            if part in archive.namelist():
                removals.add(part)
        if strict and removals:
            xml_transforms = [
                (lambda low: bool(re.match(
                    r"^(word|xl)/(document|header\d*|footer\d*|footnotes|endnotes)\.xml$"
                    r"|^xl/drawings/", low)), _office_strip_strict_xml),
            ]
        else:
            xml_transforms = []
        if strict and textbox_count:
            xml_transforms = xml_transforms or [
                (lambda low: bool(re.match(
                    r"^(word|xl)/(document|header\d*|footer\d*|footnotes|endnotes)\.xml$"
                    r"|^xl/drawings/", low)), _office_strip_strict_xml),
            ]
        if redact or strict:
            _office_zip_rewrite(path, replacements=replacements,
                               removals=removals if strict else set(),
                               xml_transforms=xml_transforms, warnings=warnings)
        # 以“成品实际状态”重算保留/剥离的隐藏对象数（python-docx/openpyxl 保存时
        # 可能已静默丢弃未被引用的部件，不能按源文件计数糊弄用户）
        try:
            with zipfile.ZipFile(path) as check_archive:
                after = _classify_office_parts(check_archive, fmt)
            left_hidden = (len(after.get("hidden_objects") or [])
                           + len(after.get("unsupported_images") or [])
                           + int(after.get("textbox_count") or 0))
        except Exception as exc:
            if strict:
                raise ExportError(f"strict 成品隐藏对象复检失败：{exc}") from exc
            left_hidden = stats["hidden_total"]
        if strict:
            stats["hidden_stripped"] = max(0, stats["hidden_total"] - left_hidden)
            stats["hidden_kept"] = 0
            if left_hidden:
                raise ExportError(
                    f"strict_office_media=True 但成品仍有 {left_hidden} 个隐藏对象，"
                    "无法安全剥离，已阻断导出")
        else:
            stats["hidden_kept"] = left_hidden
            if left_hidden:
                warnings.append(
                    f"保留 {left_hidden} 个无法自动脱敏的隐藏对象"
                    "（图表/文本框/OLE/activeX 等），请确认可外发；"
                    "如需剥离请用 strict_office_media=True")
    finally:
        try:
            archive.close()
        except Exception:
            pass
        if text_workdir is not None:
            try:
                shutil.rmtree(text_workdir, ignore_errors=True)
            except Exception:
                pass
    # 自检 1：被替换图片重扫视觉/码类必须 0 命中
    if redact and replaced_parts:
        ocr_check = None
        try:
            from ..ocr import get_ocr

            ocr_check = get_ocr()
            if not getattr(ocr_check, "available", False):
                ocr_check = None
        except Exception:
            ocr_check = None
        leftovers = _office_verify_images(path, replaced_parts, ocr=ocr_check,
                                          warnings=warnings)
        if leftovers:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass
            raise ExportError("导出后 Office 内嵌图片复检失败：" + "；".join(leftovers))
    # 自检 2（P1.5）：被替换图片重跑文字 helper，敏感文字命中必须 0
    if redact and text_replaced_parts:
        verify_started = time.time()
        verify_ocr_ok = False
        try:
            from ..ocr import get_ocr as _get_ocr_probe

            _probe = _get_ocr_probe()
            verify_ocr_ok = bool(getattr(_probe, "available", False))
        except Exception:
            verify_ocr_ok = False
        text_issues: list[str] = []
        try:
            with zipfile.ZipFile(path) as check_archive:
                for part in text_replaced_parts:
                    if (text_verify_seconds > 0
                            and time.time() - verify_started > text_verify_seconds):
                        remaining = max(0, len(text_replaced_parts)
                                        - stats["text_verify_checked"]
                                        - stats["text_verify_skipped"])
                        stats["text_verify_skipped"] += remaining
                        warnings.append(
                            f"内嵌图片文字 OCR 自检超时（>{text_verify_seconds:.0f}s），"
                            f"{remaining} 张未能确认文字残留，请人工复核")
                        break
                    try:
                        data = check_archive.read(part)
                    except Exception as exc:
                        stats["text_verify_skipped"] += 1
                        warnings.append(f"内嵌图片文字自检无法读取 {part}：{exc}")
                        continue
                    if not verify_ocr_ok:
                        stats["text_verify_skipped"] += 1
                        warnings.append(
                            f"内嵌图片 {part} 文字残留无法确认（OCR 不可用），请人工复核")
                        continue
                    try:
                        verify_new, verify_hits, verify_warns = _office_redact_text_bytes(
                            data, part, None)
                    except Exception as exc:
                        stats["text_verify_skipped"] += 1
                        warnings.append(f"内嵌图片文字自检失败 {part}：{exc}")
                        continue
                    stats["text_verify_checked"] += 1
                    warnings.extend("文字自检：" + str(item) for item in (verify_warns or []))
                    if verify_new is not None or int(verify_hits or 0) > 0:
                        stats["text_verify_hits"] += max(1, int(verify_hits or 0))
                        text_issues.append(part)
        except Exception as exc:
            text_issues.append(f"{type(exc).__name__}: {exc}")
        if text_issues:
            try:
                path.unlink(missing_ok=True)
            except OSError:
                pass
            raise ExportError(
                "导出后 Office 内嵌图片文字复检失败：" + "；".join(text_issues[:6]))
    if not redact and stats["images_total"]:
        stats["text_skipped"] = max(stats["text_skipped"], stats["images_total"])
        warnings.append(
            f"自动脱敏已关闭（redact_office_media=False）：{stats['images_total']} 张内嵌图片"
            "原样保留，存在外发泄漏风险")
    if redact and not text_enabled and stats["images_total"]:
        warnings.append(
            "内嵌图片文字 OCR 已关闭（YINZHAN_OFFICE_TEXT_OCR=0）："
            "图片内文字未处理，请人工复核")
    if stats["images_redacted"]:
        kinds = "、".join(f"{k}×{v}" for k, v in sorted(stats["hits_by_kind"].items()))
        warnings.append(
            f"已自动脱敏 {stats['images_redacted']} 张内嵌图片"
            + (f"（命中：{kinds}）" if kinds else ""))
    if stats["images_blanked"]:
        warnings.append(f"strict：{stats['images_blanked']} 张无敏感命中的内嵌图片已置空白")
    if stats["images_clean"]:
        warnings.append(
            f"{stats['images_clean']} 张内嵌图片未检出视觉/码类目标，已保留原样"
            "（文字 OCR 结果见下一条，建议人工复核）")
    if stats["text_images_processed"]:
        warnings.append(
            f"内嵌图片文字 OCR：处理 {stats['text_images_processed']} 张，"
            f"命中 {stats['text_hits']} 处，脱敏 {stats['text_images_redacted']} 张")
    if stats["text_skipped"]:
        warnings.append(
            f"{stats['text_skipped']} 张内嵌图片文字 OCR 跳过（预算/格式/未开启），"
            "文字内容请人工复核")
    return stats


def _office_media_parse_meta(path: Path, fmt: str, *,
                             fields: int = 0) -> dict[str, Any]:
    """解析时枚举包内媒体并构造 meta（失败返回 {}，绝不带崩解析）。"""
    try:
        with zipfile.ZipFile(path) as archive:
            classified = _classify_office_parts(archive, fmt)
    except Exception:
        return {}
    return _build_office_media_meta(classified, fmt=fmt, fields=fields)


# ══════════════════════════════ 解析辅助 ══════════════════════════════

def _read_app_meta(path: Path) -> dict[str, str]:
    """从 docProps/app.xml 读取公司/经理/模板等属性（失败返回空）。"""
    meta: dict[str, str] = {}
    if not _LXML_OK:
        return meta
    try:
        with zipfile.ZipFile(path) as archive:
            if "docProps/app.xml" not in archive.namelist():
                return meta
            root = etree.fromstring(archive.read("docProps/app.xml"))
    except Exception:
        return meta
    for element in root:
        key = _APP_META_TAGS.get(_localname(element.tag))
        if not key:
            continue
        value = (element.text or "").strip()
        if value:
            meta[key] = value
    return meta


def _collect_core_meta(document: Any) -> dict[str, str]:
    """收集非空 core properties（导出时全部清空）。"""
    meta: dict[str, str] = {}
    try:
        properties = document.core_properties
    except Exception:
        return meta
    for field_name in _CORE_META_FIELDS:
        try:
            value = _fmt_meta_value(getattr(properties, field_name, ""))
        except Exception:
            continue
        if value.strip():
            meta[field_name] = value
    return meta


# ══════════════════════════════ 解析器 ══════════════════════════════

class DocxParser(Parser):
    """DOCX 流式文档解析器（flow 类格式，不提供位图渲染）。"""

    format_id = "docx"
    extensions = (".docx",)
    can_render = False

    # ── 判型 ──
    @classmethod
    def sniff(cls, path: Path) -> float:
        try:
            if Path(path).suffix.lower() in cls.extensions:
                return 1.0
            with open(path, "rb") as handle:
                head = handle.read(2)
            if head == b"PK":
                with zipfile.ZipFile(path) as archive:
                    if "word/document.xml" in archive.namelist():
                        return 0.9
        except Exception:
            pass
        return 0.0

    # ── 解析 ──
    def parse(self, path: Path, ctx: ParseContext) -> Document:
        path = Path(path)
        document = Document(
            doc_id="d_" + uuid.uuid4().hex[:8],
            path=str(path),
            fmt=self.format_id,
        )
        page = Page(index=0, width=A4_WIDTH, height=A4_HEIGHT)
        document.pages.append(page)

        if not _DOCX_OK:
            document.warnings.append(f"缺少 python-docx，DOCX 无法解析：{_DOCX_ERR}")
            document.pages.clear()
            return document

        try:
            parsed = docx.Document(str(path))
        except Exception as exc:  # 解析失败降级，不向用户抛异常
            document.pages.clear()
            document.warnings.append(f"DOCX 解析失败：{exc}")
            return document

        counters = {"blocks": 0, "hyperlinks": 0, "nested_tables": 0}
        blocks = page.blocks

        def add_block(
            kind: str = "text",
            text: str = "",
            ref: dict | None = None,
            attrs: dict | None = None,
        ) -> Block:
            block = Block(
                id=f"b{counters['blocks']}",
                kind=kind,  # type: ignore[arg-type]
                text=text,
                ref=dict(ref or {}),
                page=0,
                attrs=dict(attrs or {}),
            )
            counters["blocks"] += 1
            blocks.append(block)
            return block

        def scan_paragraph(paragraph: Any, ref: dict, attrs_extra: dict | None = None) -> str:
            text = _paragraph_text(paragraph)
            try:
                if len(paragraph._p.hyperlink_lst or []):
                    counters["hyperlinks"] += 1
            except Exception:
                pass
            if text.strip():
                attrs = {"style": _safe_style_name(paragraph)}
                if attrs_extra:
                    attrs.update(attrs_extra)
                add_block("text", text, ref, attrs)
            return text

        def scan_cell(cell: Any, table_index: int, row: int, col: int) -> str:
            text = _cell_text(cell)
            if getattr(cell, "tables", None):
                counters["nested_tables"] += len(cell.tables)
            attrs = {"table": table_index}
            add_block("text", text, {"t": table_index, "r": row, "c": col}, attrs)
            return text

        def scan_table(table: Any, table_index: int) -> None:
            try:
                rows = len(table.rows)
                cols = len(table.columns)
            except Exception:
                rows, cols = 0, 0
            table_block = add_block(
                "table", "", {"t": table_index}, {"rows": rows, "cols": cols}
            )
            # 合并单元格会在 row.cells 里重复出现：按底层 <w:tc> 去重。
            # 注意必须同时保存元素引用，否则 lxml 代理被回收后 id() 会被复用。
            seen: dict[int, Any] = {}
            cells: list[Cell] = []
            for row_index, row in enumerate(table.rows):
                for col_index, cell in enumerate(row.cells):
                    tc = cell._tc
                    key = id(tc)
                    if key in seen:
                        continue
                    seen[key] = tc
                    text = scan_cell(cell, table_index, row_index, col_index)
                    cells.append(Cell(text=text, row=row_index, col=col_index))
            table_block.cells = cells

        def scan_sdt(element: Any, body_index: int) -> None:
            try:
                content = element.find(qn("w:sdtContent"))
            except Exception:
                return
            if content is None:
                return
            paragraph_index = 0
            for child in content.iterchildren():
                tag = _localname(child.tag)
                if tag == "p":
                    scan_paragraph(
                        Paragraph(child, parsed),
                        {"sdt": [body_index, paragraph_index]},
                        {"part": "sdt"},
                    )
                    paragraph_index += 1
                elif tag in ("tbl", "sdt"):
                    counters["nested_tables"] += 1

        try:
            body = parsed.element.body
            paragraph_index = 0
            table_index = 0
            for body_index, child in enumerate(body.iterchildren()):
                tag = _localname(child.tag)
                if tag == "p":
                    scan_paragraph(Paragraph(child, parsed), {"p": paragraph_index})
                    paragraph_index += 1
                elif tag == "tbl":
                    scan_table(Table(child, parsed), table_index)
                    table_index += 1
                elif tag == "sdt":
                    scan_sdt(child, body_index)
            # 页眉页脚（含首页/奇偶页变体）
            for section_index, section in enumerate(parsed.sections):
                slots = [("header", "header"), ("footer", "footer")]
                try:
                    if section.different_first_page_header_footer:
                        slots += [
                            ("header_first", "first_page_header"),
                            ("footer_first", "first_page_footer"),
                        ]
                except Exception:
                    pass
                try:
                    if parsed.settings.odd_and_even_pages_header_footer:
                        slots += [
                            ("header_even", "even_page_header"),
                            ("footer_even", "even_page_footer"),
                        ]
                except Exception:
                    pass
                for part_name, attr_name in slots:
                    header_footer = getattr(section, attr_name, None)
                    if header_footer is None:
                        continue
                    try:
                        if header_footer.is_linked_to_previous:
                            continue  # 继承自上一节，已扫过
                    except Exception:
                        pass
                    try:
                        for paragraph_index, paragraph in enumerate(header_footer.paragraphs):
                            scan_paragraph(
                                paragraph,
                                {"part": part_name, "section": section_index,
                                 "p": paragraph_index},
                                {"part": part_name},
                            )
                        tables = getattr(header_footer, "tables", None)
                        if tables:
                            counters["nested_tables"] += len(tables)
                    except Exception as exc:
                        document.warnings.append(
                            f"页眉/页脚解析失败（第 {section_index} 节 {part_name}）：{exc}"
                        )
        except Exception as exc:  # 结构异常时保留已解析内容
            document.warnings.append(f"DOCX 结构遍历中断，已保留部分内容：{exc}")

        # 未覆盖子特性的告警（契约要求：不能静默丢弃）
        try:
            field_count = len(body.xpath(".//w:instrText"))
        except Exception:
            field_count = 0
        try:
            textbox_count = len(body.xpath(".//w:txbxContent"))
        except Exception:
            textbox_count = 0
        if counters["hyperlinks"]:
            document.warnings.append(
                f"文档含 {counters['hyperlinks']} 处超链接，超链接显示文字未纳入检测（python-docx runs 不含超链接 run）"
            )
        if counters["nested_tables"]:
            document.warnings.append(
                f"文档含 {counters['nested_tables']} 个嵌套表格/内容控件表格，未展开检测"
            )

        # P1-B：zip 全量枚举内嵌媒体/隐藏对象（含页眉页脚里的文本框、图表、OLE 等）
        # 内嵌图片导出时会自动扫描脱敏；图表/文本框/OLE 等无法自动脱敏 → 需确认。
        media_meta: dict[str, Any] = {}
        try:
            media_meta = _office_media_parse_meta(path, "docx", fields=field_count)
        except Exception as exc:  # 枚举失败不影响正文解析，但必须告警
            document.warnings.append(f"内嵌媒体枚举失败：{exc}")
        try:
            with zipfile.ZipFile(path) as archive:
                zip_names = [str(name).lower() for name in archive.namelist()]
        except Exception:
            zip_names = []
        if any("glossary" in name for name in zip_names):
            document.warnings.append("文档含词汇表（构建基块），其中文字未纳入检测")
        if any(name.endswith("footnotes.xml") or name.endswith("endnotes.xml")
               for name in zip_names):
            document.warnings.append("文档含脚注/尾注，其中文字未纳入检测")

        meta = _collect_core_meta(parsed)
        meta.update(_read_app_meta(path))
        if media_meta:
            meta["embedded_media"] = media_meta
            meta["hidden_objects"] = list(media_meta.get("hidden_objects") or [])
        document.meta = meta
        _append_office_media_warnings(document, media_meta, fmt="docx")
        return document

    # ── 导出 ──
    def export(
        self,
        doc: Document,
        entities: list[Entity],
        decisions: dict[str, Decision],
        out_path: Path,
        options: dict,
    ) -> ExportResult:
        warnings: list[str] = []
        source = Path(doc.path)
        out_path = Path(out_path)
        try:
            out_path.parent.mkdir(parents=True, exist_ok=True)
        except Exception:
            pass
        if not source.is_file():
            raise ExportError(f"原文件不存在：{source}")
        shutil.copyfile(source, out_path)

        if not _DOCX_OK:
            warnings.append(f"缺少 python-docx，导出的 DOCX 未做任何脱敏：{_DOCX_ERR}")
            return ExportResult(
                path=str(out_path), warnings=warnings, size=out_path.stat().st_size
            )

        try:
            parsed = docx.Document(str(out_path))
        except Exception as exc:
            warnings.append(f"DOCX 打开失败，已原样复制（未脱敏）：{exc}")
            return ExportResult(
                path=str(out_path), warnings=warnings, size=out_path.stat().st_size
            )

        block_by_id = {block.id: block for _page, block in doc.all_blocks()}
        tasks: dict[str, list[tuple[int, int, str]]] = {}
        unresolved = 0
        for entity in entities:
            if entity.status == "rejected" or not entity.value:
                continue
            decision = decisions.get(entity.id) or entity.decision
            if decision is None:
                decision = Decision(strategy="black", black=True)
            if decision.strategy == "keep":
                continue
            for span in entity.spans:
                if not span.is_text:
                    unresolved += 1
                    continue
                block = block_by_id.get(span.block_id or "")
                if block is None or block.kind != "text":
                    unresolved += 1
                    continue
                length = max(0, int(span.end) - int(span.start))
                replacement = _replacement_text(decision, length)
                if replacement is None:
                    continue
                tasks.setdefault(span.block_id, []).append(
                    (int(span.start), int(span.end), replacement)
                )

        # 值级传播：检测范围外若还有同值出现，也按同一决策覆盖（防漏）
        propagation = _propagation_rules(entities, decisions, warnings)
        if propagation:
            propagated = _propagate_in_blocks(block_by_id, tasks, propagation)
            if propagated:
                warnings.append(
                    f"检测范围外另有 {propagated} 处同值出现，已按同一决策一并覆盖"
                )

        for block_id, items in tasks.items():
            block = block_by_id[block_id]
            for start, end, replacement in sorted(items, key=lambda item: item[0], reverse=True):
                try:
                    applied = self._apply_span(parsed, block, start, end, replacement)
                except Exception:
                    applied = False
                if not applied:
                    unresolved += 1
        if unresolved:
            warnings.append(
                f"有 {unresolved} 处标记无法映射到 DOCX 的 run，请人工复核成品"
            )

        # 修订 / 批注 / 保护口令等 XML 级清理
        removed_revisions = 0
        try:
            for part in parsed.part.package.parts:
                element = getattr(part, "element", None)
                if element is None:
                    continue
                if _localname(element.tag) in ("document", "hdr", "ftr", "footnotes", "endnotes"):
                    removed_revisions += _strip_revisions(element)
            if removed_revisions:
                warnings.append(
                    f"已删除 {removed_revisions} 个修订/批注标记节点（含未接受的修订内容）"
                )
            removed_settings = _clean_settings(parsed)
            if removed_settings:
                warnings.append(
                    f"已删除 {removed_settings} 个设置项（文档保护口令/附加模板/邮件合并数据源等）"
                )
        except Exception as exc:
            warnings.append(f"清理修订/批注节点失败：{exc}")

        # core properties 清空
        cleared = 0
        try:
            properties = parsed.core_properties
            for field_name in _CORE_META_FIELDS:
                try:
                    if _fmt_meta_value(getattr(properties, field_name, "")).strip():
                        cleared += 1
                    setattr(properties, field_name, "")
                except Exception:
                    continue
            if cleared:
                warnings.append(f"已清空 {cleared} 项文档核心属性（作者/标题/关键词等）")
        except Exception as exc:
            warnings.append(f"清空 core properties 失败：{exc}")

        try:
            parsed.save(str(out_path))
        except Exception as exc:
            warnings.append(f"DOCX 保存失败，已回退为原样复制（未脱敏）：{exc}")
            shutil.copyfile(source, out_path)
            return ExportResult(
                path=str(out_path), warnings=warnings, size=out_path.stat().st_size
            )

        _harden_zip(out_path, warnings)
        # P1-B：内嵌图片自动脱敏 + （strict）隐藏对象剥离 + 导出后自检
        try:
            stats = _office_process_media(out_path, fmt="docx",
                                          options=options or {}, warnings=warnings)
            doc.extra["office_media"] = stats
        except ExportError:
            try:
                out_path.unlink(missing_ok=True)
            except OSError:
                pass
            raise
        return ExportResult(
            path=str(out_path), warnings=warnings, size=out_path.stat().st_size
        )

    # ── 导出：按 ref 定位并替换 ──
    @staticmethod
    def _apply_span(parsed: Any, block: Block, start: int, end: int, new: str) -> bool:
        ref = block.ref or {}
        if "t" in ref:
            table = parsed.tables[int(ref["t"])]
            cell = table.cell(int(ref["r"]), int(ref["c"]))
            return _replace_in_cell(cell, start, end, new)
        if "sdt" in ref:
            paragraph = _find_sdt_paragraph(parsed, ref["sdt"])
            return _replace_in_paragraph(paragraph, start, end, new) if paragraph else False
        if "part" in ref:
            paragraph = _find_header_footer_paragraph(parsed, ref)
            return _replace_in_paragraph(paragraph, start, end, new) if paragraph else False
        if "p" in ref:
            paragraphs = parsed.paragraphs
            index = int(ref["p"])
            if index < 0 or index >= len(paragraphs):
                return False
            return _replace_in_paragraph(paragraphs[index], start, end, new)
        return False


_SECTION_HF_ATTR = {
    "header": "header",
    "footer": "footer",
    "header_first": "first_page_header",
    "footer_first": "first_page_footer",
    "header_even": "even_page_header",
    "footer_even": "even_page_footer",
}


def _find_sdt_paragraph(parsed: Any, path: Any) -> Any | None:
    """按 ref={"sdt": [body_child_index, paragraph_index]} 找到内容控件里的段落。"""
    try:
        body_index, paragraph_index = int(path[0]), int(path[1])
        children = list(parsed.element.body.iterchildren())
        content = children[body_index].find(qn("w:sdtContent"))
        if content is None:
            return None
        paragraphs = content.findall(qn("w:p"))
        if paragraph_index < 0 or paragraph_index >= len(paragraphs):
            return None
        return Paragraph(paragraphs[paragraph_index], parsed)
    except Exception:
        return None


def _find_header_footer_paragraph(parsed: Any, ref: dict) -> Any | None:
    """按 ref={"part": ..., "section": i, "p": j} 找到页眉/页脚段落。"""
    try:
        section = parsed.sections[int(ref["section"])]
        header_footer = getattr(section, _SECTION_HF_ATTR[str(ref["part"])])
        paragraphs = header_footer.paragraphs
        index = int(ref["p"])
        if index < 0 or index >= len(paragraphs):
            return None
        return paragraphs[index]
    except Exception:
        return None


register_parser(DocxParser)
