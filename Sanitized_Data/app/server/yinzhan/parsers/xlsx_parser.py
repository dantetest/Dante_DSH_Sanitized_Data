"""XLSX / XLSM 解析器与破坏性导出器。

parse
    用 openpyxl 打开工作簿（data_only=False，保留公式文本），可见工作表的每个
    非空单元格生成一个 text Block（ref 里带 sheet / cell / row / col），一个可见
    工作表一页 Page；隐藏工作表不检测，只记 warning 与 meta。

export
    复制原文件后用 openpyxl 修改：black → 单元格清空 + 纯黑填充；surrogate/mask →
    直接写入 replacement（公式单元格会被替换成文本）；数值单元格（origin=
    ``numeric_cell``，见 detectors/numeric_det.py）按数值写回，保留 number_format，
    使 Excel 公式/求和仍可重算；随后清空文档属性、自定义属性、定义名称、批注、
    外部链接，删除隐藏工作表，并打开全量重算。
    另做一次「值级传播」：检测器漏掉的单元格里若仍出现同一实体值，也按同一决策
    覆盖（防漏），并写 warning。

已知边界（均写 warning）：
    - openpyxl 保存会丢弃公式缓存值（正是我们想要的防残留行为），但 Excel 打开
      成品时需要重算；
    - .xlsm 的 VBA 宏不会保留，成品实际按普通 .xlsx 内容保存；
    - 图表 / 图片 / 透视表等对象里可能藏有原文，本解析器不做 OCR，只告警。
"""
from __future__ import annotations

import datetime as _dt
import shutil
import uuid
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from ..registry import Parser, register_parser
from ..udm import (
    Block, Decision, Document, Entity, ExportError, ExportResult, Page, ParseContext,
)

try:  # pragma: no cover - 环境已装
    import openpyxl
    from openpyxl.styles import PatternFill

    _OPENPYXL_OK = True
    _OPENPYXL_ERR = ""
except Exception as _exc:  # pragma: no cover
    openpyxl = None  # type: ignore[assignment]
    PatternFill = None  # type: ignore[assignment]
    _OPENPYXL_OK = False
    _OPENPYXL_ERR = repr(_exc)

A4_WIDTH = 595.0
A4_HEIGHT = 842.0
BLACK_CHAR = "█"
# 纯黑底：与前端「涂黑 #0D0B09」一致（alpha=FF）
BLACK_FILL_RGB = "FF0D0B09"

# openpyxl 属性名 -> meta 键（meta 只收集非空敏感项，供出口自检核对）
_PROP_FIELDS: tuple[tuple[str, str], ...] = (
    ("creator", "creator"),
    ("lastModifiedBy", "last_modified_by"),
    ("title", "title"),
    ("subject", "subject"),
    ("description", "description"),
    ("keywords", "keywords"),
    ("category", "category"),
    ("contentStatus", "content_status"),
    ("identifier", "identifier"),
    ("language", "language"),
)
_PROP_FIELDS_CLEAR = tuple(name for name, _key in _PROP_FIELDS)

# 保存后理论上不该再出现的部件（openpyxl 会重建包，出现即说明有残留）
_SUSPICIOUS_PARTS = (
    "docprops/custom.xml",
    "docprops/thumbnail",
    "xl/comments",
    "xl/threadedcomments",
    "xl/persons/",
)


def _meta_text(value: Any) -> str:
    if value is None:
        return ""
    try:
        return str(value).strip()
    except Exception:
        return ""


def _replacement_text(decision: Decision, length: int) -> str | None:
    """把 Decision 翻译成实际写入的单元格文本；keep 返回 None。"""
    if decision.strategy == "keep":
        return None
    if decision.strategy == "drop":
        return ""
    if decision.black or decision.strategy == "black" or decision.replacement is None:
        return BLACK_CHAR * max(0, length)
    return decision.replacement


def _cell_attrs(cell: Any, formula: bool) -> dict[str, Any]:
    """单元格元数据（只读，供 numeric_cell 检测器使用；不改变单元格值）。"""
    value = getattr(cell, "value", None)
    numeric = isinstance(value, (int, float)) and not isinstance(value, bool)
    is_date = isinstance(value, (_dt.datetime, _dt.date, _dt.time))
    return {
        "formula": bool(formula),
        "numeric": bool(numeric),
        "date": bool(is_date),
        "number_format": str(getattr(cell, "number_format", "") or ""),
    }


def _numeric_meta(entity: Entity) -> dict[str, Any] | None:
    """读取数值单元格实体在 numeric_det 里登记的元数据（不存在返回 None）。"""
    try:
        from ..detectors import numeric_det

        return numeric_det.get_numeric_meta(entity.id)
    except Exception:
        return None


def _propagation_rules(
    entities: list[Entity], decisions: dict[str, Decision], warnings: list[str]
) -> dict[str, str]:
    """值级传播规则：同一实体值在检测范围外的其它单元格也按同一决策覆盖。

    检测器可能漏掉某些单元格，而出口自检按「成品里是否还有该值」判定残留；
    这里做防御性传播。单字符值不传播，只记 warning。
    """
    rules: dict[str, str] = {}
    short: set[str] = set()
    for entity in entities:
        if entity.status == "rejected" or not entity.value:
            continue
        decision = decisions.get(entity.id) or entity.decision
        if decision is None:
            decision = Decision(strategy="black", black=True)
        if decision.strategy == "keep":
            continue
        if len(entity.value) < 2:
            short.add(entity.value)
            continue
        replacement = _replacement_text(decision, len(entity.value))
        if replacement is None:
            continue
        existing = rules.get(entity.value)
        if existing is not None and existing != replacement:
            warnings.append(
                f"同一值「{entity.value[:12]}」存在不同的替换决策，已按涂黑处理"
            )
            rules[entity.value] = BLACK_CHAR * len(entity.value)
            continue
        rules[entity.value] = replacement
    if short:
        warnings.append(f"有 {len(short)} 个单字符实体值未做值级传播，请人工复核")
    return rules


class XlsxParser(Parser):
    """Excel 工作簿解析器（flow 类格式，不提供位图渲染）。"""

    format_id = "xlsx"
    extensions = (".xlsx", ".xlsm")
    can_render = False

    # ── 判型 ──
    @classmethod
    def sniff(cls, path: Path) -> float:
        try:
            if Path(path).suffix.lower() in cls.extensions:
                return 1.0
            with open(path, "rb") as handle:
                head = handle.read(2)
            if head == b"PK":
                with zipfile.ZipFile(path) as archive:
                    if "xl/workbook.xml" in archive.namelist():
                        return 0.9
        except Exception:
            pass
        return 0.0

    # ── 解析 ──
    def parse(self, path: Path, ctx: ParseContext) -> Document:
        path = Path(path)
        document = Document(
            doc_id="d_" + uuid.uuid4().hex[:8],
            path=str(path),
            fmt=self.format_id,
        )
        if not _OPENPYXL_OK:
            document.warnings.append(f"缺少 openpyxl，XLSX 无法解析：{_OPENPYXL_ERR}")
            return document
        try:
            workbook = openpyxl.load_workbook(str(path), data_only=False)
        except Exception as exc:  # 失败降级：空 Document + warning
            document.warnings.append(f"XLSX 解析失败：{exc}")
            return document

        try:
            meta: dict[str, Any] = {}
            try:
                properties = workbook.properties
            except Exception:
                properties = None
            if properties is not None:
                for attr, key in _PROP_FIELDS:
                    try:
                        value = _meta_text(getattr(properties, attr, ""))
                    except Exception:
                        continue
                    if value:
                        meta[key] = value

            visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
            hidden = [sheet.title for sheet in workbook.worksheets
                      if sheet.sheet_state != "visible"]
            meta["sheets"] = [sheet.title for sheet in visible]
            if hidden:
                meta["hidden_sheets"] = hidden
                document.warnings.append(
                    f"工作簿含 {len(hidden)} 个隐藏工作表（{', '.join(hidden[:5])}），未纳入检测，导出时删除"
                )

            try:
                external_links = len(workbook._external_links or [])
            except Exception:
                external_links = 0
            if external_links:
                document.warnings.append(f"工作簿含 {external_links} 个外部链接，导出时清除")

            block_counter = 0
            formula_count = 0
            object_count = 0
            for page_index, sheet in enumerate(visible):
                try:
                    cells = getattr(sheet, "_cells", None)
                    if isinstance(cells, dict):
                        candidates = list(cells.values())
                    else:  # 退化路径：按维度迭代（超大格式化表可能较慢）
                        candidates = [cell for row in sheet.iter_rows() for cell in row]
                except Exception:
                    candidates = []
                values: list[tuple[int, int, Any]] = []
                for cell in candidates:
                    try:
                        value = cell.value
                        row = int(cell.row)
                        col = int(cell.column)
                    except Exception:
                        continue
                    if value is None or str(value) == "":
                        continue
                    values.append((row, col, cell))
                values.sort(key=lambda item: (item[0], item[1]))

                blocks: list[Block] = []
                max_row = 0
                max_col = 0
                for row, col, cell in values:
                    value = cell.value
                    text = str(value)
                    formula = getattr(cell, "data_type", None) == "f"
                    if formula:
                        formula_count += 1
                    max_row = max(max_row, row)
                    max_col = max(max_col, col)
                    blocks.append(
                        Block(
                            id=f"b{block_counter}",
                            kind="text",
                            text=text,
                            ref={
                                "sheet": sheet.title,
                                "cell": str(cell.coordinate),
                                "row": row,
                                "col": col,
                            },
                            page=page_index,
                            attrs=_cell_attrs(cell, formula),
                        )
                    )
                    block_counter += 1
                try:
                    object_count += len(getattr(sheet, "_charts", []) or [])
                    object_count += len(getattr(sheet, "_images", []) or [])
                    object_count += len(getattr(sheet, "_pivots", []) or [])
                except Exception:
                    pass
                document.pages.append(
                    Page(
                        index=page_index,
                        width=max(A4_WIDTH, min(max_col * 64.0, 10000.0)),
                        height=max(A4_HEIGHT, min(max_row * 20.0, 100000.0)),
                        blocks=blocks,
                        ref={"sheet": sheet.title},
                    )
                )

            if formula_count:
                document.warnings.append(
                    f"工作簿含 {formula_count} 个公式单元格；导出保存后公式缓存值会被丢弃（防止残留原文）"
                )
            if object_count:
                document.warnings.append(
                    f"工作簿含 {object_count} 个图表/图片/透视对象，其中文字未纳入检测"
                )
            # P1-B：zip 全量枚举内嵌媒体/隐藏对象（内嵌图片导出时自动扫描脱敏；
            # 图表/OLE/文本框等无法自动脱敏 → 计入 hidden_objects 需确认）
            media_meta: dict[str, Any] = {}
            try:
                from .docx_parser import (_append_office_media_warnings,
                                          _office_media_parse_meta)

                media_meta = _office_media_parse_meta(path, "xlsx")
                if media_meta:
                    meta["embedded_media"] = media_meta
                    meta["hidden_objects"] = list(media_meta.get("hidden_objects") or [])
                _append_office_media_warnings(document, media_meta, fmt="xlsx")
            except Exception as exc:
                document.warnings.append(f"内嵌媒体枚举失败：{exc}")
            document.meta = meta
        except Exception as exc:  # 结构异常时保留已解析页
            document.warnings.append(f"XLSX 工作表遍历中断，已保留部分内容：{exc}")
        return document

    # ── 导出 ──
    def export(
        self,
        doc: Document,
        entities: list[Entity],
        decisions: dict[str, Decision],
        out_path: Path,
        options: dict,
    ) -> ExportResult:
        warnings: list[str] = []
        source = Path(doc.path)
        out_path = Path(out_path)
        try:
            out_path.parent.mkdir(parents=True, exist_ok=True)
        except Exception:
            pass
        if not source.is_file():
            raise ExportError(f"原文件不存在：{source}")
        shutil.copyfile(source, out_path)

        if not _OPENPYXL_OK:
            out_path.unlink(missing_ok=True)
            raise ExportError(f"缺少 openpyxl，无法安全导出 XLSX：{_OPENPYXL_ERR}")

        try:
            workbook = openpyxl.load_workbook(str(out_path), data_only=False)
        except Exception as exc:
            out_path.unlink(missing_ok=True)
            raise ExportError(f"XLSX 打开失败，未产出成品：{exc}") from exc

        # ── 1. 单元格级替换 ──
        block_by_id = {block.id: block for _page, block in doc.all_blocks()}
        black_cells: set[tuple[str, str]] = set()
        cell_edits: dict[tuple[str, str], list[tuple[int, int, str]]] = {}
        # 数值单元格（numeric_det 识别）：写回数字而不是字符串，保留 number_format
        numeric_writes: dict[tuple[str, str], tuple[str, dict[str, Any]]] = {}
        unresolved = 0
        for entity in entities:
            if entity.status == "rejected" or not entity.value:
                continue
            decision = decisions.get(entity.id) or entity.decision
            if decision is None:
                decision = Decision(strategy="black", black=True)
            if decision.strategy == "keep":
                continue
            for span in entity.spans:
                if not span.is_text:
                    unresolved += 1
                    continue
                block = block_by_id.get(span.block_id or "")
                if block is None or block.kind != "text":
                    unresolved += 1
                    continue
                ref = block.ref or {}
                sheet_name = str(ref.get("sheet") or "")
                coordinate = str(ref.get("cell") or "")
                if not sheet_name or not coordinate:
                    unresolved += 1
                    continue
                key = (sheet_name, coordinate)
                length = max(0, int(span.end) - int(span.start))
                replacement = _replacement_text(decision, length)
                if replacement is None:
                    continue
                if decision.black or decision.strategy == "black" or decision.strategy == "drop":
                    # 单元格无法做字符级涂黑，统一整格处理
                    if decision.strategy == "black" or decision.black:
                        black_cells.add(key)
                    else:
                        cell_edits.setdefault(key, []).append((0, 10**12, ""))
                else:
                    numeric_meta = _numeric_meta(entity)
                    if entity.origin == "numeric_cell" and numeric_meta is not None:
                        # 数值列：整格写回数字/字符串，不做字符级替换
                        numeric_writes[key] = (replacement, numeric_meta)
                        continue
                    cell_edits.setdefault(key, []).append(
                        (int(span.start), int(span.end), replacement)
                    )

        # 值级传播：检测范围外若还有同值出现，也按同一决策覆盖（防漏）
        propagation = _propagation_rules(entities, decisions, warnings)
        if propagation:
            ordered = sorted(
                propagation.items(), key=lambda item: len(item[0]), reverse=True
            )
            propagated = 0
            for block in block_by_id.values():
                if block.kind != "text" or not block.text:
                    continue
                ref = block.ref or {}
                sheet_name = str(ref.get("sheet") or "")
                coordinate = str(ref.get("cell") or "")
                if not sheet_name or not coordinate:
                    continue
                key = (sheet_name, coordinate)
                if key in black_cells or key in numeric_writes:
                    continue
                edits = cell_edits.setdefault(key, [])
                for value, replacement in ordered:
                    start = 0
                    while True:
                        index = block.text.find(value, start)
                        if index < 0:
                            break
                        end = index + len(value)
                        if not any(
                            not (end <= old_start or index >= old_end)
                            for old_start, old_end, _old in edits
                        ):
                            edits.append((index, end, replacement))
                            propagated += 1
                        start = end
                if not edits:
                    cell_edits.pop(key, None)
            if propagated:
                warnings.append(
                    f"检测范围外另有 {propagated} 处同值出现，已按同一决策一并覆盖"
                )

        sheet_names = set(workbook.sheetnames)
        fill = PatternFill(fill_type="solid", fgColor=BLACK_FILL_RGB)
        for sheet_name, coordinate in black_cells:
            if sheet_name not in sheet_names:
                unresolved += 1
                continue
            try:
                cell = workbook[sheet_name][coordinate]
                cell.value = None
                cell.fill = fill
            except Exception:
                unresolved += 1
        # 数值单元格：写回数字（AMOUNT/其他金额类）或字符串（账号类）
        numeric_written = 0
        for (sheet_name, coordinate), (replacement, meta) in numeric_writes.items():
            if sheet_name not in sheet_names:
                unresolved += 1
                continue
            try:
                cell = workbook[sheet_name][coordinate]
            except Exception:
                unresolved += 1
                continue
            kind = str(meta.get("kind") or "amount")
            try:
                if kind in ("account", "date"):
                    cell.value = str(replacement)
                else:
                    cleaned = str(replacement).replace(",", "").replace("，", "").strip()
                    dec = Decimal(cleaned)
                    decimals = int(meta.get("decimals") or 0)
                    if decimals == 0 and "." not in cleaned:
                        cell.value = int(dec)
                    else:
                        cell.value = float(dec)
                numeric_written += 1
            except (InvalidOperation, ValueError, TypeError):
                # 解析不出数字：退回按文本写入，绝不静默丢值
                try:
                    cell.value = str(replacement)
                    numeric_written += 1
                except Exception:
                    unresolved += 1
        if numeric_written:
            warnings.append(
                f"已按数值写回 {numeric_written} 个单元格（保留 number_format，公式可重算）"
            )
        for (sheet_name, coordinate), edits in cell_edits.items():
            if sheet_name not in sheet_names:
                unresolved += 1
                continue
            try:
                cell = workbook[sheet_name][coordinate]
            except Exception:
                unresolved += 1
                continue
            current = "" if cell.value is None else str(cell.value)
            for start, end, replacement in sorted(edits, key=lambda item: item[0], reverse=True):
                start = max(0, min(start, len(current)))
                end = max(start, min(end, len(current)))
                current = current[:start] + replacement + current[end:]
            try:
                cell.value = current
            except Exception:
                unresolved += 1
        if unresolved:
            warnings.append(f"有 {unresolved} 处标记无法映射到 Excel 单元格，请人工复核成品")

        # ── 2. 删除隐藏工作表（隐藏内容默认不检测，留着就是残留） ──
        try:
            visible = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
            hidden_sheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state != "visible"]
            if hidden_sheets and visible:
                for sheet in hidden_sheets:
                    workbook.remove(sheet)
                try:
                    workbook.active = 0
                except Exception:
                    pass
                warnings.append(
                    f"已删除 {len(hidden_sheets)} 个隐藏工作表：{', '.join(s.title for s in hidden_sheets[:5])}"
                )
            elif hidden_sheets:
                warnings.append("工作簿所有工作表均为隐藏，无法安全删除，请人工确认")
        except Exception as exc:
            warnings.append(f"删除隐藏工作表失败：{exc}")

        # ── 3. 定义名称 / 批注 / 外部链接 / 自定义属性 ──
        try:
            defined_count = len(workbook.defined_names)
            workbook.defined_names.clear()
            for sheet in workbook.worksheets:
                names = getattr(sheet, "defined_names", None)
                if names is not None and len(names):
                    defined_count += len(names)
                    names.clear()
            if defined_count:
                warnings.append(f"已删除 {defined_count} 个定义名称（打印区域/命名区域等）")
        except Exception as exc:
            warnings.append(f"删除定义名称失败：{exc}")

        try:
            removed_comments = 0
            for sheet in workbook.worksheets:
                for cell in list(getattr(sheet, "_cells", {}).values()):
                    try:
                        if cell.comment is not None:
                            cell.comment = None
                            removed_comments += 1
                    except Exception:
                        continue
                try:
                    sheet._comments = []
                    if getattr(sheet, "legacy_drawing", None):
                        sheet.legacy_drawing = None
                except Exception:
                    pass
            if removed_comments:
                warnings.append(f"已删除 {removed_comments} 条单元格批注（含批注作者）")
        except Exception as exc:
            warnings.append(f"删除批注失败：{exc}")

        try:
            external = len(getattr(workbook, "_external_links", []) or [])
            if external:
                workbook._external_links.clear()
                warnings.append(f"已清除 {external} 个外部链接引用")
        except Exception as exc:
            warnings.append(f"清除外部链接失败：{exc}")

        try:
            custom = getattr(workbook, "custom_doc_props", None)
            if custom is not None and len(custom):
                count = len(custom)
                custom.props.clear()
                warnings.append(f"已清空 {count} 项自定义文档属性")
        except Exception as exc:
            warnings.append(f"清空自定义属性失败：{exc}")

        try:
            security = getattr(workbook, "security", None)
            if security is not None and (getattr(security, "workbookPassword", None)
                                         or getattr(security, "lockStructure", None)):
                security.workbookPassword = None
                security.lockStructure = False
                warnings.append("已移除工作簿结构保护口令")
        except Exception as exc:
            warnings.append(f"移除工作簿保护失败：{exc}")

        # ── 4. 文档属性清空 + 强制重算 ──
        cleared = 0
        try:
            properties = workbook.properties
            for attr in _PROP_FIELDS_CLEAR:
                try:
                    if _meta_text(getattr(properties, attr, "")):
                        cleared += 1
                    setattr(properties, attr, "")
                except Exception:
                    continue
            if cleared:
                warnings.append(f"已清空 {cleared} 项工作簿属性（作者/标题/关键词等）")
        except Exception as exc:
            warnings.append(f"清空工作簿属性失败：{exc}")

        try:
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.calcMode = "auto"
        except Exception:
            pass

        if Path(doc.path).suffix.lower() == ".xlsm":
            warnings.append("xlsm 的 VBA 宏不会保留；成品按普通 xlsx 内容保存，建议使用 .xlsx 扩展名")

        try:
            workbook.save(str(out_path))
        except Exception as exc:
            out_path.unlink(missing_ok=True)
            raise ExportError(f"XLSX 保存失败，未产出成品：{exc}") from exc

        # 公式缓存丢失提示（无论原文是否有公式，导出语义都依赖重算）
        warnings.append("成品不保留公式缓存值（openpyxl 语义），Excel 打开时会自动重算")

        # P1-B：内嵌图片自动脱敏 + （strict）隐藏对象剥离 + 导出后自检
        try:
            from .docx_parser import _office_process_media

            stats = _office_process_media(out_path, fmt="xlsx",
                                          options=options or {}, warnings=warnings)
            doc.extra["office_media"] = stats
        except ExportError:
            try:
                out_path.unlink(missing_ok=True)
            except OSError:
                pass
            raise

        # ── 5. 防御性检查：确认没有意外保留的注释/缩略图部件 ──
        try:
            with zipfile.ZipFile(out_path) as archive:
                leftovers = [
                    name for name in archive.namelist()
                    if any(part in name.lower() for part in _SUSPICIOUS_PARTS)
                ]
            if leftovers:
                warnings.append(
                    f"成品仍含可疑部件：{', '.join(leftovers[:5])}，请人工复核"
                )
        except Exception:
            pass

        return ExportResult(
            path=str(out_path), warnings=warnings, size=out_path.stat().st_size
        )


register_parser(XlsxParser)
