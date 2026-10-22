# -*- coding: utf-8 -*-
"""财务域 P0 测试：finance-cn v2 规则 / 校验器 / 金额一致性引擎。

运行（仓库任意目录，必须 -X utf8）：

    python -X utf8 app/tests/test_finance.py

也可用 pytest 收集（函数名 test_*）。只写系统临时目录，不污染仓库 data/。

覆盖（对应 task-9 验收）：
1. 校验器正反例：USCC / 发票代码 / 发票号码 / SWIFT-BIC / 组织机构代码 /
   工商注册号 / 开户许可证 / 审计报告编号 / 纳税人识别号；
2. finance-cn v2 规则落盘与 validate 名全部已注册、v1 五条规则保留；
3. 财务字段端到端识别（营业执照字段 / 发票 / 银行流水 / USCC 正反例，
   USCC 非法码被校验位拦截）；
4. 金额一致性：合计=分项和、同一金额多写法同值、组内统一因子 k、
   中文大写金额同步、单位/千分位/小数位保形、比率不缩放；
5. 既有 5 个样例上"新旧财务规则对比"：不新增误报、不丢失既有命中；
6. 文本导出 + 出口自检 residual=0，原文金额不残留。
"""

from __future__ import annotations

import os
import re
import sys
import tempfile
import time
import traceback
from decimal import Decimal
from pathlib import Path

# 必须在 import yinzhan.config 之前把数据目录指向临时目录
_TMP = Path(tempfile.mkdtemp(prefix="yz_fin_test_"))
os.environ.setdefault("YINZHAN_DATA_DIR", str(_TMP / "data"))
os.environ.setdefault("YINZHAN_OUTPUT_DIR", str(_TMP / "out"))

REPO = Path(__file__).resolve().parents[2]
SERVER = REPO / "app" / "server"
FIXTURES = Path(__file__).resolve().parent / "fixtures" / "finance"
SAMPLES = REPO / "samples"
sys.path.insert(0, str(SERVER))

from yinzhan import pipeline, registry, rules as rules_mod          # noqa: E402
from yinzhan.detectors.regex_det import _VALIDATORS, uscc_checksum  # noqa: E402
from yinzhan.parsers.text_parser import TextParser                   # noqa: E402
from yinzhan.replacers import amount_consistency as ac               # noqa: E402
from yinzhan.udm import Entity, ParseContext, Rule                   # noqa: E402

#: finance-cn v2 新增规则 id（用于"新旧规则对比"时回退到 v1 行为）
NEW_RULE_IDS = frozenset({
    "cn.biz_reg_no", "cn.org_code", "cn.tax_no", "cn.legal_rep",
    "cn.reg_address", "cn.license_no", "cn.reg_capital",
    "cn.account_license_no", "cn.swift_bic", "cn.bank_routing_no",
    "cn.pay_platform_id", "cn.invoice_code", "cn.invoice_no",
    "cn.invoice_person", "cn.contract_no", "cn.audit_report_no",
    "cn.audit_cpa", "cn.account_name",
})

#: v1 既有 5 条（行为必须保持不变）
V1_RULE_IDS = frozenset({
    "cn.bank_account", "cn.customer_no", "cn.credit_line",
    "cn.txn_no", "cn.iban",
})

_NUM_RE = re.compile(r"-?\d[\d,]*(?:\.\d+)?")
_TP = TextParser()


# ────────────────────────── 小工具 ──────────────────────────


def _ctx() -> ParseContext:
    return ParseContext(ocr=None, workdir=str(_TMP / "work"))


def _packs() -> dict:
    return rules_mod.load_packs()


def _detect_text(fixture: str):
    """解析并检测一个文本夹具，返回 (doc, detection, session)。"""
    doc = _TP.parse(FIXTURES / fixture, _ctx())
    rules = rules_mod.compile_rules(_packs(), "text")
    det = pipeline.detect(doc, rules, "standard")
    session = pipeline.Session()
    pipeline.decide(det.entities, det.rules, session)
    return doc, det, session


def _plan_of(doc, det, session) -> dict[str, str]:
    return ac.build_amount_plan(det.entities, det.rules, session)


def _num(text: str) -> Decimal | None:
    """从替换文本里取第一个数字（去掉千分位）。"""
    match = _NUM_RE.search(text or "")
    if not match:
        return None
    return Decimal(match.group(0).replace(",", ""))


def _find(entities, etype: str | None = None, label: str | None = None,
          value_contains: str | None = None):
    out = []
    for entity in entities:
        if etype and entity.etype != etype:
            continue
        if label and entity.label != label:
            continue
        if value_contains and value_contains not in entity.value:
            continue
        out.append(entity)
    return out


def _one(entities, etype: str | None = None, label: str | None = None,
         value_contains: str | None = None):
    found = _find(entities, etype, label, value_contains)
    assert len(found) == 1, (
        f"期望恰好 1 个实体 etype={etype} label={label} "
        f"value_contains={value_contains!r}，实际 {len(found)}："
        f"{[(e.etype, e.label, e.value) for e in found]}")
    return found[0]


def _one_value_num(entities, etype: str, expected: Decimal):
    """按金额数值精确挑一个实体（避免 "1,300,000.00" 命中 "11,300,000.00"）。"""
    found = [entity for entity in entities
             if entity.etype == etype and _num(entity.value) == expected]
    assert len(found) == 1, (
        f"期望恰好 1 个 {etype} 数值={expected}，实际："
        f"{[(e.etype, e.value) for e in found]}")
    return found[0]


def _entity(eid: str, value: str, etype: str = "AMOUNT", status: str = "accepted") -> Entity:
    return Entity(id=eid, etype=etype, label="金额", value=value, spans=[],
                  level="L2", confidence=0.99, origin="regex", status=status)


def _rule(rid: str, role: str = "", strategy: str = "surrogate") -> Rule:
    replace = {"strategy": strategy}
    if role:
        replace["amount_role"] = role
    return Rule(id=rid, entity="AMOUNT", label="金额", level="L2",
                scope=["text"], detect=[{"type": "regex", "pattern": ".+"}],
                replace=replace)


def _rules_for(entities) -> dict[str, Rule]:
    """默认给每个实体配一条 line 角色规则。"""
    return {entity.id: _rule("test." + entity.id) for entity in entities}


# ────────────────────────── 1. 校验器正反例 ──────────────────────────


def test_validators_uscc():
    assert _VALIDATORS["uscc_checksum"]("91350100M000100Y43") is True
    assert _VALIDATORS["uscc_checksum"]("91310000MA1FL1XX0G") is True
    # 律所执业许可证号：18 位同形，但校验位非法（误报正例）
    assert _VALIDATORS["uscc_checksum"]("31310000E12345678X") is False
    assert _VALIDATORS["uscc_checksum"]("31310000E123456789") is False
    assert uscc_checksum("1234") is False


def test_validators_invoice():
    v_code = _VALIDATORS["fin_invoice_code"]
    assert v_code("031002100311") is True          # 12 位，0 开头
    assert v_code("131002100311") is False         # 12 位，非 0 开头
    assert v_code("03100210031X") is False         # 含字母
    assert v_code("3100210031") is True            # 10 位历史票种
    assert v_code("0310021003116") is False        # 13 位
    v_no = _VALIDATORS["fin_invoice_no"]
    assert v_no("12345678") is True
    assert v_no("1234567") is False
    assert v_no("123456789") is False
    assert v_no("00000000") is False


def test_validators_swift_bic():
    v = _VALIDATORS["fin_swift_bic"]
    assert v("ICBKCNBJSHA") is True                # 11 位
    assert v("ICBKCNBJ") is True                   # 8 位
    assert v("ICBKCNBJS") is False                 # 9 位
    # 8 位普通大写单词：第 5–6 位不是 ISO 国家码 → 拒绝
    # （"INDUSTRY" 的 "ST"、"DATABASE" 的 "BA" 恰好是真实国家码，会通过，
    #   所以还要靠 YAML 的 near 上下文压误报）
    assert v("PASSWORD") is False
    assert v("ICBKCNBJSHA7") is False              # 12 位


def test_validators_org_code():
    v = _VALIDATORS["fin_org_code"]
    # GB 11714 加权校验：12345678 的校验位是 8；1234567A 的校验位是 4
    assert v("12345678-8") is True
    assert v("123456788") is True
    assert v("12345678-9") is False                # 校验位错
    assert v("1234567A-4") is True
    assert v("1234567A-3") is False
    assert v("1234567") is False                   # 长度不够


def test_validators_business_and_tax():
    v_reg = _VALIDATORS["fin_business_reg_no"]
    assert v_reg("310000000012345") is True
    assert v_reg("31000000001234") is False        # 14 位
    assert v_reg("990000000012345") is False       # 非省级前缀
    assert v_reg("111111111111111") is False       # 全相同

    v_tax = _VALIDATORS["fin_tax_no"]
    assert v_tax("91350100M000100Y43") is True     # 18 位 = 合法 USCC
    assert v_tax("310000000012345") is True        # 15 位旧税号
    assert v_tax("31310000E12345678X") is False    # 18 位但 USCC 校验不过
    assert v_tax("9135010") is False               # 长度不对


def test_validators_account_license_and_audit():
    v_acc = _VALIDATORS["fin_account_license"]
    assert v_acc("J3100001234567") is True         # J + 13 位 = 14 位
    assert v_acc("J310000123456") is False         # J + 12 位
    assert v_acc("A3100001234567") is False        # 非 J 开头
    v_audit = _VALIDATORS["fin_audit_report_no"]
    assert v_audit("信会师报字[2025]第ZA12345号") is True
    assert v_audit("审计报告") is False            # 无数字
    assert v_audit("2025") is False                # 无"号/字/报告"


# ────────────────────────── 2. 规则包落盘 ──────────────────────────


def test_rule_pack_v2():
    assert registry.load_plugins() == [], "插件加载报错"
    packs = _packs()
    assert not rules_mod.load_errors(), rules_mod.load_errors()
    finance = [rule for rule in packs["finance-cn"] if rule.pack == "finance-cn"]
    ids = {rule.id for rule in finance}
    assert V1_RULE_IDS <= ids, f"v1 规则丢失：{V1_RULE_IDS - ids}"
    assert NEW_RULE_IDS <= ids, f"v2 新增规则缺失：{NEW_RULE_IDS - ids}"
    assert len(finance) == 23, f"finance-cn 自有规则应为 23 条，实际 {len(finance)}"

    # 所有 validate 名必须已注册，且不能是 unknown→none 的静默降级
    for rule in finance:
        for spec in rule.detect:
            name = str(spec.get("validate") or "").strip().lower()
            if name and name != "none":
                assert name in _VALIDATORS, f"{rule.id} 的 validate={name} 未注册"
    # 新增编号类规则必须配 validate（防止将来被改坏）
    validated = {
        "cn.biz_reg_no", "cn.org_code", "cn.tax_no", "cn.account_license_no",
        "cn.swift_bic", "cn.invoice_code", "cn.invoice_no", "cn.audit_report_no",
    }
    for rule in finance:
        if rule.id in validated:
            assert any(str(spec.get("validate") or "") for spec in rule.detect), \
                f"{rule.id} 缺少 validate"


def test_rule_pack_scopes():
    packs = _packs()
    allowed = {"pdf", "docx", "xlsx", "text", "image", "tiff", "mesh"}
    for rule in packs["finance-cn"]:
        if rule.pack != "finance-cn":
            continue
        assert set(rule.scope) <= allowed, f"{rule.id} scope 非法：{rule.scope}"
        assert isinstance(rule.priority, int), f"{rule.id} priority 非整数"


# ────────────────────────── 3. 财务字段端到端识别 ──────────────────────────


def test_uscc_false_positive_blocked():
    _doc, det, _session = _detect_text("uscc_cases.txt")
    uscc = _find(det.entities, etype="USCC")
    values = {entity.value for entity in uscc}
    assert values == {"91350100M000100Y43"}, f"USCC 误报未被拦截：{values}"
    for entity in det.entities:
        assert "31310000E12345678X" not in entity.value


def test_license_fields_detection():
    _doc, det, _session = _detect_text("license_fields.txt")
    ents = det.entities
    assert _one(ents, etype="USCC", value_contains="91350100M000100Y43").value
    assert _one(ents, etype="BIZ_REG_NO").value == "310000000012345"
    assert _one(ents, etype="ORG_CODE").value == "12345678-8"
    assert _one(ents, etype="TAX_NO").value == "310000000012345"
    assert _one(ents, etype="PERSON", label="法定代表人").value == "王建国"
    addr = _one(ents, etype="ADDRESS", label="注册地址")
    assert "世纪大道" in addr.value
    assert _one(ents, etype="LICENSE_NO").value.startswith("310000000012345")
    assert _one(ents, etype="ACCOUNT_LICENSE").value == "J3100001234567"
    assert _one(ents, etype="SWIFT_BIC").value == "ICBKCNBJSHA"
    assert _one(ents, etype="CONTRACT_NO").value == "HT-2025-0098"
    assert "信会师报字" in _one(ents, etype="AUDIT_NO").value
    assert _one(ents, etype="PERSON", label="签字注册会计师").value == "张伟"
    capital = _one(ents, etype="AMOUNT", label="注册资本")
    assert capital.value == "人民币5000万元"
    # 公开信息 L1：默认 keep（不参与金额缩放）
    assert capital.decision is not None and capital.decision.strategy == "keep"


def test_invoice_detection():
    _doc, det, session = _detect_text("invoice.txt")
    ents = det.entities
    assert _one(ents, etype="INVOICE_CODE").value == "031002100311"
    assert _one(ents, etype="INVOICE_NO").value == "12345678"
    assert _one(ents, etype="PERSON", label="开票人").value == "李静怡"
    # 编号替换：同长度、不泄漏原值
    for entity in ents:
        if entity.etype in ("INVOICE_CODE", "INVOICE_NO"):
            assert entity.decision is not None
            assert entity.decision.strategy == "surrogate"
            assert entity.decision.replacement != entity.value
            assert len(entity.decision.replacement) == len(entity.value)

    plan = _plan_of(_doc, det, session)
    amount = _one(ents, etype="CREDIT_LINE", value_contains="10,000,000.00")
    tax = _one_value_num(ents, "AMOUNT", Decimal("1300000"))
    total = _one_value_num(ents, "AMOUNT", Decimal("11300000"))
    new_amount = _num(plan[amount.id])
    new_tax = _num(plan[tax.id])
    new_total = _num(plan[total.id])
    assert new_total == new_amount + new_tax, (
        f"发票价税合计勾稽失败：{new_total} != {new_amount} + {new_tax}")
    assert new_amount != Decimal("10000000.00")
    assert "13%" not in " ".join(plan.values())


def test_bank_statement_amounts():
    _doc, det, session = _detect_text("bank_statement.txt")
    plan = _plan_of(_doc, det, session)
    account = _one(det.entities, etype="ACCOUNT")
    assert account.id not in plan, "银行账号不应参与金额缩放"

    amounts = [entity for entity in det.entities if entity.id in plan]
    assert len(amounts) >= 5, f"流水金额识别过少：{len(amounts)}"
    ratios = []
    for entity in amounts:
        new = _num(plan[entity.id])
        old = _num(entity.value)
        assert new is not None and old and old > 0
        ratios.append(new / old)
    # 同一文档同一因子（2 位小数舍入误差内）
    for ratio in ratios:
        assert abs(ratio - ratios[0]) < Decimal("0.000001"), f"因子不一致：{ratios}"
    # 余额 1,200,000 在"当日余额/期末余额"出现两次：同值合并成一个实体（多 span），
    # 只生成一个替换值（同值同映射）
    same = [entity for entity in amounts if "1,200,000.00" in entity.value]
    assert len(same) == 1 and len(same[0].spans) >= 2, (
        f"同值未合并或 span 丢失：{[(e.value, len(e.spans)) for e in same]}")
    assert _num(plan[same[0].id]) is not None


# ────────────────────────── 4. 金额一致性引擎 ──────────────────────────


def test_amount_plan_parts_total():
    e1 = _entity("e1", "5,000,000.00元")
    e2 = _entity("e2", "7,000,000.00元")
    e3 = _entity("e3", "12,000,000.00元")
    e4 = _entity("e4", "人民币12,000,000.00元")
    e5 = _entity("e5", "¥12,000,000.00")
    e6 = _entity("e6", "35%")
    e7 = _entity("e7", "13%")
    entities = [e1, e2, e3, e4, e5, e6, e7]
    rules = {
        "e1": _rule("t.line", "line"),
        "e2": _rule("t.line2", "line"),
        "e3": _rule("t.total", "total"),
        "e4": _rule("t.line3", "line"),
        "e5": _rule("t.line4", "line"),
        "e6": _rule("t.line5", "line"),
        "e7": _rule("t.line6", "line"),
    }
    session = pipeline.Session()
    plan = ac.build_amount_plan(entities, rules, session)

    assert "e6" not in plan and "e7" not in plan, "比率不应缩放"
    part1, part2, total = _num(plan["e1"]), _num(plan["e2"]), _num(plan["e3"])
    assert total == part1 + part2, f"合计 != 分项和：{total} != {part1}+{part2}"
    # 同一金额三种写法（含总计本身）→ 同一个假值数值
    assert _num(plan["e4"]) == total
    assert _num(plan["e5"]) == total
    # 保形：千分位/货币符号保留
    assert plan["e1"].endswith("元") and "," in plan["e1"]
    assert plan["e4"].startswith("人民币")
    assert plan["e5"].startswith("¥")

    factor = part1 / Decimal("5000000")
    assert Decimal("0.899") <= factor <= Decimal("1.101"), f"因子越界：{factor}"
    assert factor != Decimal(1), "因子不能为 1（否则等于没脱敏）"
    assert abs(part1 / Decimal("5000000") - part2 / Decimal("7000000")) < Decimal("1e-6"), \
        "分项未使用同一因子"
    # 确定性 + 会话映射（可逆导出）
    assert ac.build_amount_plan(entities, rules, pipeline.Session()) == plan
    assert session.map[("AMOUNT", "5,000,000.00元")] == plan["e1"]


def test_amount_plan_uppercase():
    total_numeric = _entity("u1", "11,300,000.00元")
    total_upper = _entity("u2", "壹仟壹佰叁拾万元整")
    part1 = _entity("u3", "10,000,000.00元")
    part2 = _entity("u4", "1,300,000.00元")
    entities = [total_numeric, total_upper, part1, part2]
    rules = {
        "u1": _rule("t.u1", "total"),
        "u2": _rule("t.u2", "total"),
        "u3": _rule("t.u3", "line"),
        "u4": _rule("t.u4", "line"),
    }
    plan = ac.build_amount_plan(entities, rules, pipeline.Session())
    assert plan["u1"] != total_numeric.value
    new_total = _num(plan["u1"])
    assert new_total == _num(plan["u3"]) + _num(plan["u4"])
    parsed_upper = ac.parse_cn_amount(plan["u2"])
    assert parsed_upper == new_total, f"中文大写未同步：{plan['u2']} -> {parsed_upper} != {new_total}"
    assert plan["u2"] != total_upper.value
    assert "元" in plan["u2"]


def test_amount_uppercase_roundtrip():
    for text in ("0", "1", "10", "105", "1006", "12000", "11300000", "12345678.91"):
        value = Decimal(text)
        formatted = ac.format_cn_amount(value)
        assert ac.parse_cn_amount(formatted) == value, f"{value} -> {formatted}"
    assert ac.parse_cn_amount("壹仟叁佰伍拾陆万元整") == Decimal("13560000")


def test_amount_units_and_ratio():
    wan = _entity("w1", "5,000万元")
    yuan = _entity("w2", "50,000,000.00元")
    plan = ac.build_amount_plan([wan, yuan], _rules_for([wan, yuan]), pipeline.Session())
    wan_text = plan["w1"]
    assert re.fullmatch(r"\d{1,3}(?:,\d{3})*万元", wan_text), f"万元格式被破坏：{wan_text}"
    wan_yuan = Decimal(wan_text.replace("万元", "").replace(",", "")) * 10000
    # 跨单位按"元"同值：5,000万元 与 50,000,000.00元 必须得到同一个元数值
    assert wan_yuan == _num(plan["w2"]), (
        f"跨单位同值不一致：{wan_text} vs {plan['w2']}")

    # 千分位/两位小数保形
    for value, pattern in (("1,234,567.89元", r"^\d{1,3}(?:,\d{3})+\.\d{2}元$"),
                           ("1234元", r"^\d+元$")):
        entity = _entity("f_" + value, value)
        plan = ac.build_amount_plan([entity], _rules_for([entity]), pipeline.Session())
        assert re.fullmatch(pattern, plan[entity.id]), f"{value} -> {plan[entity.id]}"


def test_amount_plan_skips_keep_and_rejected():
    keep = _entity("k1", "5,000,000.00元")
    keep_rule = _rule("t.keep", strategy="keep")
    rejected = _entity("r1", "7,000,000.00元", status="rejected")
    plan = ac.build_amount_plan([keep, rejected],
                                {"k1": keep_rule, "r1": _rule("t.r1")},
                                pipeline.Session())
    assert plan == {}, f"keep/rejected 实体不应被缩放：{plan}"
    assert ac.build_amount_plan([], {}, pipeline.Session()) == {}


# ────────────────────────── 5. 文本端到端（导出+自检） ──────────────────────────


def test_amounts_text_end_to_end():
    doc, det, session = _detect_text("amounts.txt")
    plan = _plan_of(doc, det, session)
    ents = [entity for entity in det.entities if entity.status != "rejected"]
    by_id = {entity.id: entity for entity in ents}

    total_entities = [e for e in ents
                      if e.id in plan and _num(e.value) == Decimal("12000000.00")]
    assert len(total_entities) >= 3, f"同一金额三种写法应都命中且参与缩放：{len(total_entities)}"
    total = _num(plan[total_entities[0].id])
    assert all(_num(plan[e.id]) == total for e in total_entities), (
        "同一金额多写法假值不一致")
    part1 = _one(ents, value_contains="5,000,000.00元")
    part2 = _one(ents, value_contains="7,000,000.00元")
    assert total == _num(plan[part1.id]) + _num(plan[part2.id]), "文本：合计 != 分项和"
    # 比率/税率不参与
    assert all("%" not in (by_id[eid].value) for eid in plan)

    ac.apply_plan(ents, plan, session)
    out = _TMP / "amounts_out.txt"
    _TP.export(doc, ents, pipeline.decisions_of(ents), out, {"mode": "ai"})
    verify = pipeline.verify(out, ents, _TP, _ctx())
    assert verify.get("residual") == 0, f"导出残留：{verify}"
    text = out.read_text(encoding="utf-8")
    for original in ("12,000,000.00", "5,000,000.00", "7,000,000.00"):
        assert original not in text, f"原文金额残留：{original}"


# ────────────────────────── 6. 既有样例不新增误报 ──────────────────────────


def test_xlsx_profit_statement_text_amounts():
    """利润表 xlsx（金额为"元"文本）走完整 XLSX 解析/检测/缩放/导出/复检。

    纯数值单元格（float）的金额检测需要表头驱动的 numeric_cell 检测器（Lead P1，
    见 docs/impl/finance.md §5.1）；本测试覆盖"文本金额"这一当前可完整闭环的
    xlsx 场景，并验证导出后单元格不残留原文、合计=分项和。
    """
    from openpyxl import Workbook, load_workbook
    from yinzhan.parsers.xlsx_parser import XlsxParser

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "利润表"
    sheet.append(["项目", "本期金额"])
    sheet.append(["收入A", "3,000,000.00元"])
    sheet.append(["收入B", "2,000,000.00元"])
    sheet.append(["合计", "5,000,000.00元"])
    src = _TMP / "利润表_文本金额.xlsx"
    workbook.save(src)

    parser = XlsxParser()
    doc = parser.parse(src, _ctx())
    rules = rules_mod.compile_rules(_packs(), "xlsx")
    det = pipeline.detect(doc, rules, "standard")
    session = pipeline.Session()
    pipeline.decide(det.entities, det.rules, session)
    plan = ac.build_amount_plan(det.entities, det.rules, session)
    assert len(plan) == 3, f"利润表应识别 3 个金额，实际 {len(plan)}"

    by_num = {_num(det.entities[i].value): det.entities[i] for i in range(len(det.entities))
              if det.entities[i].id in plan}
    part_a = by_num[Decimal("3000000.00")]
    part_b = by_num[Decimal("2000000.00")]
    total = by_num[Decimal("5000000.00")]
    assert _num(plan[total.id]) == _num(plan[part_a.id]) + _num(plan[part_b.id]), \
        "xlsx 利润表：合计 != 分项和"

    ents = [entity for entity in det.entities if entity.status != "rejected"]
    ac.apply_plan(ents, plan, session)
    out = _TMP / "利润表_out.xlsx"
    parser.export(doc, ents, pipeline.decisions_of(ents), out, {})
    verify = pipeline.verify(out, ents, parser, _ctx())
    assert verify.get("residual") == 0, f"xlsx 导出残留：{verify}"
    out_wb = load_workbook(out)
    out_sheet = out_wb[out_wb.sheetnames[0]]
    values = [str(cell.value) for row in out_sheet.iter_rows() for cell in row
              if cell.value is not None]
    for original in ("3,000,000.00", "2,000,000.00", "5,000,000.00"):
        assert not any(original in value for value in values), \
            f"xlsx 单元格原文残留：{original}"
    out_nums = [_num(value) for value in values if _num(value) is not None]
    assert any(num == part_a_new + part_b_new
               for num in out_nums for part_a_new in out_nums
               for part_b_new in out_nums
               if num > part_a_new and num > part_b_new and part_a_new != part_b_new), \
        f"xlsx 导出后合计≠分项和：{out_nums}"


def test_job_pipeline_integration():
    """走真实 JobManager 管线：确认 Lead 在 jobs.py 里接入金额引擎生效。

    这条测试保护"管线集成点"：如果 jobs.py 的接入被改坏/回退，金额会退回
    逐处 ±25% 随机，本测试会失败。
    """
    from yinzhan.jobs import MANAGER

    job = MANAGER.create(FIXTURES / "amounts.txt", strictness="standard")
    deadline = time.time() + 30
    while time.time() < deadline:
        snapshot = job.snapshot()
        if snapshot["stage"] == "ready":
            break
        if snapshot["stage"] == "error":
            raise AssertionError(f"Job 失败：{snapshot.get('error')}")
        time.sleep(0.05)
    else:
        raise AssertionError("Job 30s 未完成")
    assert any("金额一致性缩放" in warning for warning in job.warnings), (
        f"jobs.py 未接入金额引擎？warnings={job.warnings}")
    record = MANAGER.export(job.id, "ai")
    assert record.get("verify", {}).get("residual") == 0, record.get("verify")
    text = Path(record["out"]).read_text(encoding="utf-8")
    for original in ("12,000,000.00", "5,000,000.00", "7,000,000.00"):
        assert original not in text, f"Job 导出原文金额残留：{original}"
    numbers = [Decimal(x.replace(",", "")) for x in
               re.findall(r"(\d{1,3}(?:,\d{3})*\.\d{2})元", text)]
    assert len(numbers) >= 3, f"导出金额过少：{numbers}"
    # 合计 = 分项和（导出文本里至少有一对分项能加出另一个数）
    matched = any(
        num == a + b
        for num in numbers for a in numbers for b in numbers
        if num > a and num > b and a != b
    )
    assert matched, f"Job 导出后合计≠分项和：{numbers}"


def test_no_changes_on_baseline_samples():
    """在既有 5 个样例上对比"v1 规则 vs v2 规则"：

    - 不能少任何既有命中（功能红线 F1/F7）；
    - 不能新增任何误报（新增规则在样例上应 0 命中）。
    """
    from yinzhan.ocr import get_ocr

    registry.load_plugins()
    ocr = get_ocr()
    packs = _packs()
    workdir = str(_TMP / "samples")
    problems: list[str] = []
    for path in sorted(p for p in SAMPLES.iterdir() if p.is_file()):
        parser = registry.find_parser(path)
        assert parser is not None, f"没有解析器：{path.name}"
        doc = parser.parse(path, ParseContext(ocr=ocr, workdir=workdir))
        full = rules_mod.compile_rules(packs, doc.fmt)
        legacy = [rule for rule in full if rule.id not in NEW_RULE_IDS]
        d_old = pipeline.detect(doc, legacy, "standard")
        d_new = pipeline.detect(doc, full, "standard")
        old = {(entity.etype, entity.value) for entity in d_old.entities}
        new = {(entity.etype, entity.value) for entity in d_new.entities}
        missing = old - new
        added = new - old
        if missing:
            problems.append(f"{path.name} 少了 {sorted(missing)}")
        if added:
            problems.append(f"{path.name} 多了 {sorted(added)}")
    assert not problems, "既有样例误报/漏报检查失败：\n" + "\n".join(problems)


# ────────────────────────── 入口 ──────────────────────────


def main() -> int:
    registry.load_plugins()
    tests = [(name, obj) for name, obj in sorted(globals().items())
             if name.startswith("test_") and callable(obj)]
    failed = 0
    print(f"\n  数据隐站 · 财务 P0 测试（{len(tests)} 项）")
    print("  " + "─" * 58)
    for name, fn in tests:
        try:
            fn()
        except AssertionError as exc:
            failed += 1
            print(f"  ✗ {name}")
            print(f"    {exc}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"  ✗ {name}  [{type(exc).__name__}] {exc}")
            traceback.print_exc()
        else:
            print(f"  ✓ {name}")
    print("  " + "─" * 58)
    if failed:
        print(f"  失败 {failed}/{len(tests)}")
        return 1
    print(f"  全部通过（{len(tests)}/{len(tests)}）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
