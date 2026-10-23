# -*- coding: utf-8 -*-
"""PIPL 补充域 P0 测试：pipl-cn 文本规则 + XLSX 纯数值单元格。

运行（仓库任意目录，必须 -X utf8）：

    python -X utf8 app/tests/test_pipl_rules.py

只写系统临时目录，不污染仓库 data/；全部虚构数据。

覆盖（对应 task-17 验收）：
1. pipl-cn 规则包/校验器清单与有效性（pipl_* 命名空间、默认禁用的时间规则）；
2. 13 类文本规则正例（社交 ID/昵称/快递/票号/CVV/VIN/IMEI/MAC/SN/IP/门牌/
   航班车次座位/工牌/空格手机号），及替换后导出残留 0；
3. 5 个既有样例"无 pipl 包 vs 全量"对比：不丢既有命中；唯一预期新增是
   对账单.xlsx 数值金额列的 AMOUNT（origin=numeric_cell）；
4. XLSX 纯数值：利润表（标题行+多列表头）金额识别、合计=分项、导出写数字、
   公式保留 + fullCalcOnLoad、number_format 保留、verify residual 0；
5. 数值账号/日期：账号按字符串写回；日期默认 keep 不变；百分比列不处理；
6. 工资表/流水数值列：同一文档同一缩放因子、流水勾稽（期初+发生=期末）；
7. 身份证校验位策略（2026-10-23，隐私优先）：校验位合法仍 accepted；校验位/
   日期不合法但形状像身份证的 18 位串被 cn.idcard_shape 兜住 → ID_CARD、
   pending、surrogate 替换（human 涂黑 / AI 换名）；非身份证形 18 位不误报；
   与银行卡 Luhn 规则重叠时同值只保留一个实体；无上下文的裸号同样兜住。
   策略与评估见 docs/qa/idcard-checksum-policy-20261023.md。
"""

from __future__ import annotations

import datetime as dt
import json
import os
import re
import sys
import tempfile
from decimal import Decimal
from pathlib import Path

_TMP = Path(tempfile.mkdtemp(prefix="yz_pipl_test_"))
os.environ.setdefault("YINZHAN_DATA_DIR", str(_TMP / "data"))
os.environ.setdefault("YINZHAN_OUTPUT_DIR", str(_TMP / "out"))

REPO = Path(__file__).resolve().parents[2]
SERVER = REPO / "app" / "server"
FIXTURES = Path(__file__).resolve().parent / "fixtures" / "pipl"
SAMPLES = REPO / "samples"
sys.path.insert(0, str(SERVER))

from yinzhan import pipeline, registry, rules as rules_mod          # noqa: E402
from yinzhan.detectors.regex_det import _VALIDATORS                 # noqa: E402
from yinzhan.parsers.text_parser import TextParser                  # noqa: E402
from yinzhan.parsers.xlsx_parser import XlsxParser                  # noqa: E402
from yinzhan.replacers import amount_consistency as ac              # noqa: E402
from yinzhan.udm import ParseContext                                # noqa: E402

_TP = TextParser()
_CTX = ParseContext(ocr=None, workdir=str(_TMP / "work"))


# ────────────────────────── 小工具 ──────────────────────────


def _packs():
    return rules_mod.load_packs()


def _pipl_rules() -> list:
    return [rule for rule in _packs().get("pipl-cn", []) if rule.pack == "pipl-cn"]


def _detect_text(path: Path, strictness: str = "standard"):
    doc = _TP.parse(path, _CTX)
    rules = rules_mod.compile_rules(_packs(), "text")
    det = pipeline.detect(doc, rules, strictness)
    session = pipeline.Session()
    pipeline.decide(det.entities, det.rules, session)
    return doc, det, session


def _write_text_case(name: str, text: str) -> Path:
    """写一个临时 txt 用例，返回路径（只写系统临时目录）。"""
    path = _TMP / name
    path.write_text(text, "utf-8")
    return path


def _id_cards(det) -> list:
    return [entity for entity in det.entities if entity.etype == "ID_CARD"]


def _entity_set(entities) -> set[tuple[str, str]]:
    return {(entity.etype, entity.value) for entity in entities}


def _num(text) -> Decimal | None:
    match = re.search(r"-?\d[\d,]*(?:\.\d+)?", str(text or ""))
    if not match:
        return None
    return Decimal(match.group(0).replace(",", ""))


def _detect_xlsx(path: Path):
    parser = XlsxParser()
    doc = parser.parse(path, _CTX)
    rules = rules_mod.compile_rules(_packs(), "xlsx")
    det = pipeline.detect(doc, rules, "standard")
    session = pipeline.Session()
    pipeline.decide(det.entities, det.rules, session)
    return parser, doc, det, session


def _export_xlsx(parser, doc, det, session, out: Path):
    ents = [entity for entity in det.entities if entity.status != "rejected"]
    plan = ac.build_amount_plan(det.entities, det.rules, session)
    ac.apply_plan(ents, plan, session)
    parser.export(doc, ents, pipeline.decisions_of(ents), out, {})
    return plan, ents


# ────────────────────────── 1. 规则包/校验器 ──────────────────────────


def test_rule_pack_and_validators():
    assert registry.load_plugins() == [], "插件加载报错"
    packs = _packs()
    assert not rules_mod.load_errors(), rules_mod.load_errors()
    assert "pipl-cn" in packs, "缺少 pipl-cn 规则包"
    finance = [rule for rule in packs["pipl-cn"] if rule.pack == "pipl-cn"]
    assert len(finance) >= 15, f"pipl-cn 规则过少：{len(finance)}"
    ids = {rule.id for rule in finance}
    # 13 类 P0 的关键规则 id
    required = {
        "pipl.social_id", "pipl.nickname", "pipl.phone_grouped",
        "pipl.express_no", "pipl.ticket_no", "pipl.card_cvv",
        "pipl.vin", "pipl.imei", "pipl.mac", "pipl.device_sn",
        "pipl.ip_addr", "pipl.address_unit", "pipl.doorplate",
        "pipl.flight_no", "pipl.train_no", "pipl.seat_no", "pipl.badge",
        "pipl.numeric_amount", "pipl.numeric_account", "pipl.numeric_date",
    }
    assert required <= ids, f"缺少规则：{required - ids}"
    # 精确时间默认关闭（日期默认不动）
    datetime_rule = next(rule for rule in finance if rule.id == "pipl.datetime")
    assert datetime_rule.enabled is False, "pipl.datetime 必须默认 enabled: false"
    # 所有 validate 名已注册且为 pipl_ 前缀
    for rule in finance:
        for spec in rule.detect:
            name = str(spec.get("validate") or "").strip().lower()
            if not name:
                continue
            assert name in _VALIDATORS, f"{rule.id} 的 validate={name} 未注册"
            assert name.startswith("pipl_"), f"{rule.id} 校验器未用 pipl_ 前缀：{name}"


def test_validators():
    v = _VALIDATORS
    assert v["pipl_phone_digits"]("138 0013 8000") is True
    assert v["pipl_phone_digits"]("13800138000") is True
    assert v["pipl_phone_digits"]("12345678901") is False
    assert v["pipl_social_wxid"]("wxid_abc12345") is True
    assert v["pipl_social_generic"]("123456789") is True
    assert v["pipl_nickname"]("小明同学") is True
    assert v["pipl_nickname"]("\\") is False
    assert v["pipl_express_no"]("SF1234567890123") is True
    assert v["pipl_express_no"]("0000000000") is False
    assert v["pipl_ticket_no"]("7841234567890") is True
    assert v["pipl_ticket_no"]("0000000000") is False
    assert v["pipl_cvv"]("123") is True
    assert v["pipl_cvv"]("12345") is False
    # ISO 3779 已知合法 VIN 与改错一位
    assert v["pipl_vin"]("1M8GDM9AXKP042788") is True
    assert v["pipl_vin"]("1M8GDM9AXKP042789") is False
    assert v["pipl_vin"]("LSVAM4180E2123456") is False
    assert v["pipl_imei"]("490154203237518") is True
    assert v["pipl_imei"]("490154203237519") is False
    assert v["pipl_mac"]("00:1A:2B:3C:4D:5E") is True
    assert v["pipl_mac"]("00-1A-2B-3C-4D-5E") is True
    assert v["pipl_device_sn"]("SN-A1B2C3D4") is True
    assert v["pipl_ipv4"]("192.168.1.23") is True
    assert v["pipl_ipv4"]("256.1.1.1") is False
    assert v["pipl_ipv6"]("2001:db8::1") is True
    assert v["pipl_ipv6"]("::1") is True
    assert v["pipl_address_unit"]("3栋2单元501室") is True
    assert v["pipl_address_unit"]("12号") is False
    assert v["pipl_doorplate"]("1203") is True
    assert v["pipl_flight_no"]("CA1234") is True
    assert v["pipl_train_no"]("G1234") is True
    assert v["pipl_train_no"]("1234") is False
    assert v["pipl_seat_no"]("32A") is True
    assert v["pipl_seat_no"]("32G") is False
    assert v["pipl_employee_no"]("EMP2025001") is True


# ────────────────────────── 2. 文本规则正例/导出 ──────────────────────────


def test_text_rules_detection():
    _doc, det, session = _detect_text(FIXTURES / "text_p0.txt")
    got = _entity_set(det.entities)
    expected = {
        ("SOCIAL_ID", "wxid_abc12345"),
        ("SOCIAL_ID", "123456789"),
        ("SOCIAL_ID", "weibo_user01"),
        ("NICKNAME", "小明同学"),
        ("PHONE", "138 0013 8000"),
        ("EXPRESS_NO", "SF1234567890123"),
        ("EXPRESS_NO", "78123456789012"),
        ("TICKET_NO", "7841234567890"),
        ("CARD_CVV", "123"),
        ("VIN", "1M8GDM9AXKP042788"),
        ("DEVICE_ID", "490154203237518"),
        ("MAC", "00:1A:2B:3C:4D:5E"),
        ("DEVICE_SN", "SN-A1B2C3D4"),
        ("IP_ADDR", "192.168.1.23"),
        ("IP_ADDR", "2001:db8::1"),
        ("FLIGHT_NO", "CA1234"),
        ("TRAIN_NO", "G1234"),
        ("SEAT_NO", "32A"),
        ("EMPLOYEE_NO", "EMP2025001"),
    }
    missing = expected - got
    assert not missing, f"文本规则漏识别：{missing}"
    assert any(entity.etype == "ADDRESS" and "3栋2单元501室" in entity.value
               for entity in det.entities), "独立门牌未识别"
    assert any(entity.etype == "ADDRESS" and entity.value == "1203"
               for entity in det.entities), "房间号未识别"
    # 替换值必须变化（不泄漏原值）
    for entity in det.entities:
        if entity.status == "rejected" or entity.etype in ("ADDRESS",):
            continue
        assert entity.decision is not None
        if entity.decision.strategy == "keep":
            continue
        assert entity.decision.replacement != entity.value, f"替换未变化：{entity.value}"


def test_text_export_residual_zero():
    doc, det, session = _detect_text(FIXTURES / "text_p0.txt")
    ents = [entity for entity in det.entities if entity.status != "rejected"]
    out = _TMP / "pipl_text_out.txt"
    _TP.export(doc, ents, pipeline.decisions_of(ents), out, {"mode": "ai"})
    ver = pipeline.verify(out, ents, _TP, _CTX)
    assert ver.get("residual") == 0, f"文本导出残留：{ver}"
    text = out.read_text(encoding="utf-8")
    for original in ("wxid_abc12345", "SF1234567890123", "1M8GDM9AXKP042788",
                     "490154203237518", "00:1A:2B:3C:4D:5E", "192.168.1.23"):
        assert original not in text, f"原文残留：{original}"


# ────────────────────────── 3. 既有样例零非预期误报 ──────────────────────────


def test_baseline_samples_no_unexpected_fp():
    from yinzhan.ocr import get_ocr

    registry.load_plugins()
    ocr = get_ocr()
    packs = _packs()
    workdir = str(_TMP / "samples")
    added_samples: list[str] = []
    problems: list[str] = []
    for path in sorted(p for p in SAMPLES.iterdir() if p.is_file()):
        parser = registry.find_parser(path)
        assert parser is not None, f"没有解析器：{path.name}"
        doc = parser.parse(path, ParseContext(ocr=ocr, workdir=workdir))
        full = rules_mod.compile_rules(packs, doc.fmt)
        legacy = [rule for rule in full if rule.pack != "pipl-cn"]
        d_old = pipeline.detect(doc, legacy, "standard")
        d_new = pipeline.detect(doc, full, "standard")
        old = _entity_set(d_old.entities)
        new = _entity_set(d_new.entities)
        missing = old - new
        added = [entity for entity in d_new.entities
                 if (entity.etype, entity.value) not in old
                 and entity.status != "rejected"]
        if missing:
            problems.append(f"{path.name} 少了既有命中：{sorted(missing)[:5]}")
        if added:
            added_samples.append(path.name)
            for entity in added:
                # 唯一允许：纯数值金额列（numeric_det），且必须是对账单
                if not (path.name == "对账单.xlsx"
                        and entity.etype == "AMOUNT"
                        and entity.origin == "numeric_cell"):
                    problems.append(
                        f"{path.name} 非预期新增：{entity.etype}/{entity.value}"
                        f"（origin={entity.origin}）")
    allowed = {"对账单.xlsx"}
    unexpected_samples = set(added_samples) - allowed
    problems.extend(f"{name} 出现新增实体" for name in sorted(unexpected_samples))
    assert not problems, "既有样例误报检查失败：\n" + "\n".join(problems)


# ────────────────────────── 4. XLSX 纯数值金额 ──────────────────────────


def _make_profit_workbook(path: Path) -> dict:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "利润表"
    ws["A1"] = "2026 年度利润表（虚构）"
    ws.append(["项目", "本期金额", "校验"])
    ws.append(["收入A", 5_000_000, 0])
    ws.append(["收入B", 7_000_000, 0])
    ws.append(["合计", 12_000_000, 0])
    ws["C3"] = "=SUM(B3:B4)"          # 校验公式：合计应等于分项和
    for cell in ("B3", "B4", "B5"):
        ws[cell].number_format = "#,##0.00"
    wb.save(path)
    return {"total_cell": "B5", "part_cells": ("B3", "B4"),
            "formula_cell": "C3", "formula": "=SUM(B3:B4)"}


def test_numeric_profit_statement():
    from openpyxl import load_workbook

    src = _TMP / "利润表_数值.xlsx"
    meta = _make_profit_workbook(src)
    parser, doc, det, session = _detect_xlsx(src)
    amounts = [entity for entity in det.entities
               if entity.etype == "AMOUNT" and entity.origin == "numeric_cell"]
    values = {entity.value for entity in amounts}
    assert values == {"5000000", "7000000", "12000000"}, f"数值金额识别异常：{values}"

    out = _TMP / "利润表_数值_已隐去.xlsx"
    plan, ents = _export_xlsx(parser, doc, det, session, out)
    assert len(plan) == 3, f"金额引擎处理数量异常：{len(plan)}"
    by_value = {entity.value: plan[entity.id] for entity in amounts}
    part_sum = _num(by_value["5000000"]) + _num(by_value["7000000"])
    assert _num(by_value["12000000"]) == part_sum, "数值金额合计 != 分项和"

    wb = load_workbook(out)
    ws = wb["利润表"]
    for coordinate in ("B3", "B4", "B5"):
        value = ws[coordinate].value
        assert isinstance(value, (int, float)) and not isinstance(value, bool), (
            f"{coordinate} 未按数值写回：{value!r}（{type(value).__name__}）")
        assert ws[coordinate].number_format == "#,##0.00", (
            f"{coordinate} number_format 被破坏：{ws[coordinate].number_format}")
    assert Decimal(str(ws["B5"].value)) == Decimal(str(ws["B3"].value)) + Decimal(
        str(ws["B4"].value)), "成品合计 != 分项和"
    assert ws["C3"].value == meta["formula"], "公式被替换/破坏"
    assert bool(wb.calculation.fullCalcOnLoad) is True, "未开启全量重算"

    ver = pipeline.verify(out, ents, parser, _CTX)
    assert ver.get("residual") == 0, f"xlsx 导出残留：{ver}"


def test_numeric_percent_and_date_keep():
    from openpyxl import Workbook, load_workbook

    src = _TMP / "数值_日期_百分比.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "人员"
    ws.append(["姓名", "入职日期", "毛利率"])
    ws.append(["虚构甲", dt.date(2020, 3, 1), 0.35])
    ws.append(["虚构乙", dt.date(2021, 7, 15), 0.42])
    ws["C2"].number_format = "0%"
    ws["C3"].number_format = "0%"
    wb.save(src)
    parser, doc, det, session = _detect_xlsx(src)
    numeric = [entity for entity in det.entities if entity.origin == "numeric_cell"]
    dates = [entity for entity in numeric if entity.etype == "DATE"]
    percents = [entity for entity in numeric if entity.etype == "AMOUNT"]
    assert len(dates) == 2, f"日期列识别数量异常：{len(dates)}"
    assert not percents, f"百分比列不应识别为金额：{[e.value for e in percents]}"
    for entity in dates:
        assert entity.decision is not None and entity.decision.strategy == "keep"

    plan = ac.build_amount_plan(det.entities, det.rules, session)
    assert not plan, f"日期/百分比不应进入金额引擎：{plan}"
    out = _TMP / "数值_日期_百分比_out.xlsx"
    parser.export(doc, det.entities, pipeline.decisions_of(det.entities), out, {})
    wb2 = load_workbook(out)
    ws2 = wb2["人员"]
    assert ws2["C2"].value == 0.35 and ws2["C3"].value == 0.42, "百分比被改动"
    assert isinstance(ws2["B2"].value, (dt.date, dt.datetime)), "日期类型被破坏"
    kept = ws2["B2"].value
    assert (kept.year, kept.month, kept.day) == (2020, 3, 1), "日期默认应保留"


def test_numeric_account_writeback():
    from openpyxl import Workbook, load_workbook

    src = _TMP / "数值_账号.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["账号", "备注"])
    ws.append([622202112233445, "虚构"])
    wb.save(src)                       # 15 位数字账号（Excel 精度内）
    parser, doc, det, session = _detect_xlsx(src)
    accounts = [entity for entity in det.entities
                if entity.etype == "ACCOUNT" and entity.origin == "numeric_cell"]
    assert len(accounts) == 1, f"数值账号未识别：{len(accounts)}"
    ents = [entity for entity in det.entities if entity.status != "rejected"]
    out = _TMP / "数值_账号_out.xlsx"
    parser.export(doc, ents, pipeline.decisions_of(ents), out, {})
    wb2 = load_workbook(out)
    cell = wb2["Sheet"]["A2"]
    assert isinstance(cell.value, str), f"账号应写回字符串：{cell.value!r}"
    assert cell.value != "622202112233445", "账号未替换"


# ────────────────────────── 5. 工资表 / 流水 ──────────────────────────


def test_numeric_no_title_and_formula_in_data_row():
    """无标题行、且数据行含公式/文本关键词时，表头行不能被误判为数据行。"""
    from openpyxl import Workbook

    src = _TMP / "利润表_无标题.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["项目", "金额", "校验"])
    ws.append(["收入A", 5_000_000, 0])
    ws.append(["收入B", 7_000_000, 0])
    ws.append(["合计", 12_000_000, 0])
    ws["C2"] = "=B2*0"                # 数据行里有公式单元格
    wb.save(src)
    _parser, _doc, det, _session = _detect_xlsx(src)
    values = {entity.value for entity in det.entities
              if entity.etype == "AMOUNT" and entity.origin == "numeric_cell"}
    assert values == {"5000000", "7000000", "12000000"}, (
        f"含公式数据行时数值金额识别异常：{values}")


def test_numeric_salary_amounts():
    from openpyxl import Workbook

    src = _TMP / "工资表_数值.xlsx"
    wb = Workbook()
    salary = wb.active
    salary.title = "工资表"
    salary.append(["姓名", "应发工资", "实发工资"])
    salary.append(["虚构甲", 12000, 9000])
    salary.append(["虚构乙", 15000, 12000])
    wb.save(src)
    _parser, _doc, det, session = _detect_xlsx(src)
    amounts = [entity for entity in det.entities
               if entity.etype == "AMOUNT" and entity.origin == "numeric_cell"]
    # 12000 出现两次 → 同值合并为一个实体（多 span）
    assert len(amounts) == 3, f"工资表数值金额实体数异常：{len(amounts)}"
    assert {entity.value for entity in amounts} == {"12000", "9000", "15000"}
    plan = ac.build_amount_plan(det.entities, det.rules, session)
    assert len(plan) == len(amounts), f"金额计划数量异常：{len(plan)} != {len(amounts)}"
    ratios = []
    for entity in amounts:
        new, old = _num(plan[entity.id]), _num(entity.value)
        assert new is not None and old
        ratios.append(new / old)
    for ratio in ratios:
        # General 格式（0 小数位）按整数舍入，比例会有 ≤0.5/原值 的舍入误差
        assert abs(ratio - ratios[0]) < Decimal("0.0001"), (
            f"工资表缩放因子不一致：{ratios}")


def test_numeric_flow_reconciliation():
    from openpyxl import Workbook

    src = _TMP / "流水_数值.xlsx"
    wb = Workbook()
    flow = wb.active
    flow.title = "流水"
    flow.append(["项目", "金额"])
    flow.append(["期初余额", 100000])
    flow.append(["本期发生额", 20000])
    flow.append(["期末余额", 120000])
    flow.append(["其他", 0])
    wb.save(src)
    _parser, _doc, det, session = _detect_xlsx(src)
    amounts = [entity for entity in det.entities
               if entity.etype == "AMOUNT" and entity.origin == "numeric_cell"]
    assert len(amounts) == 4, f"流水数值金额数量异常：{len(amounts)}"
    plan = ac.build_amount_plan(det.entities, det.rules, session)
    by_value = {entity.value: plan[entity.id] for entity in amounts}
    assert _num(by_value["120000"]) == _num(by_value["100000"]) + _num(by_value["20000"]), (
        "流水勾稽（期初+发生=期末）失败")
    # 0 保持 0，不能被"必须变"逻辑改成 1
    assert _num(by_value["0"]) == Decimal(0), f"零金额被改写：{by_value['0']}"


# ────────────────────────── 6. 身份证校验位策略（2026-10-23） ──────────────────────────
# 背景：docs/qa/idcard-checksum-policy-20261023.md
#   * 合法校验位：cn.idcard（priority 120、validate cn_id_checksum、conf 0.99）
#     → accepted，保持既有行为；
#   * 校验位/日期不合法但形状像身份证：cn.idcard_shape（同 regex、
#     unless_validate cn_id_checksum、conf 0.66）→ ID_CARD、pending；
#   * 校验位只作为置信度信号，形状仍遮（human 涂黑 / AI 换假号 / redact_text 换号）。

# 既有样例里的合法身份证（GB 11643 校验位通过）
ID_VALID = "310104198703124510"
# 工单样例：北京 110101、1990-01-01、顺序码 123，校验位错误
ID_BAD_CHECKSUM = "110101199001011234"
# 形状通过正则但出生日期非法（1990-02-30）
ID_BAD_DATE = "110101199002301234"
# 形状通过正则但出生日期在未来（2099-01-01，cn_id_checksum 日期上限=今天）
ID_FUTURE_DATE = "110101209901011234"
# 构造值：身份证形状 + Luhn 合法 + 身份证校验位不合法（与 cn.bankcard 重叠）
ID_LUHN_SHAPE = "110101199001011233"


def test_idcard_shape_rule_contract():
    """规则契约：形状兜底必须复用同一条 regex，且只在身份证校验不通过时命中。"""
    common = [rule for rule in _packs()["common"] if rule.pack == "common"]
    by_id = {rule.id: rule for rule in common}
    assert "cn.idcard" in by_id, "缺少 cn.idcard 高置信规则"
    assert "cn.idcard_shape" in by_id, "缺少 cn.idcard_shape 形状兜底规则"

    valid_rule = by_id["cn.idcard"]
    shape_rule = by_id["cn.idcard_shape"]
    valid_spec = valid_rule.detect[0]
    shape_spec = shape_rule.detect[0]

    assert shape_rule.entity == "ID_CARD" and shape_rule.level == "L3"
    assert shape_rule.replace.get("strategy") == "surrogate", (
        "形状兜底必须仍走 surrogate（human 涂黑 / AI 换名）")
    assert str(shape_spec.get("pattern")) == str(valid_spec.get("pattern")), (
        "兜底规则必须复用 cn.idcard 的同一条 regex")
    assert shape_spec.get("type") == "regex"
    # 无硬校验（validate 缺省为 none），改用反向校验排除合法身份证，
    # 避免同一值被两条规则各出一个实体、也避免合法身份证被降成 pending。
    assert str(shape_spec.get("validate") or "none").lower() == "none", (
        "兜底规则不能设置硬 validate")
    assert str(shape_spec.get("unless_validate") or "").lower() == "cn_id_checksum", (
        "兜底规则必须 unless_validate: cn_id_checksum")
    confidence = float(shape_spec.get("confidence", 0.99))
    assert 0.60 <= confidence < 0.80, (
        f"confidence={confidence} 必须在 [0.60, 0.80) 才能 standard 档 pending")
    assert shape_rule.priority < valid_rule.priority, (
        "合法值优先走高置信 cn.idcard；兜底规则 priority 必须更低")


def test_idcard_valid_checksum_still_accepted():
    """合法身份证：仍 accepted，走高置信规则，且只产生一个实体。"""
    path = _write_text_case(
        "idcard_valid.txt",
        f"姓名：张伟，身份证号：{ID_VALID}，手机号：13812345678。")
    _doc, det, _session = _detect_text(path)
    cards = _id_cards(det)
    assert len(cards) == 1, f"合法身份证实体数异常：{[e.to_dict() for e in cards]}"
    entity = cards[0]
    assert entity.value == ID_VALID
    assert entity.status == "accepted", f"合法身份证不应 pending：{entity.status}"
    assert entity.confidence >= 0.99
    assert det.rules[entity.id].id == "cn.idcard", (
        f"合法身份证必须命中 cn.idcard：{det.rules[entity.id].id}")
    assert entity.decision is not None
    assert entity.decision.strategy == "surrogate", (
        "身份证替换策略必须是 surrogate（human 导出时 JobManager 统一改 black）")
    assert entity.decision.replacement != ID_VALID
    assert _VALIDATORS["cn_id_checksum"](entity.decision.replacement), (
        "换名后的假身份证必须自身校验位合法")


def test_idcard_bad_checksum_shape_pending_and_redacted():
    """工单复现串：校验位错误 → ID_CARD、pending、surrogate，MCP redact_text 实测替换。"""
    text = f"张伟，手机号13812345678，身份证号{ID_BAD_CHECKSUM}"
    path = _write_text_case("idcard_bad_checksum.txt", text)
    doc, det, _session = _detect_text(path)
    cards = [entity for entity in _id_cards(det) if entity.value == ID_BAD_CHECKSUM]
    assert len(cards) == 1, f"校验位错误的形状串未被唯一识别：{_id_cards(det)}"
    entity = cards[0]
    assert entity.status == "pending", f"低置信兜底应 pending：{entity.status}"
    assert 0.60 <= entity.confidence < 0.80, f"兜底置信度异常：{entity.confidence}"
    assert det.rules[entity.id].id == "cn.idcard_shape", (
        f"应命中 cn.idcard_shape：{det.rules[entity.id].id}")
    assert entity.decision is not None and entity.decision.strategy == "surrogate"
    assert entity.decision.replacement != ID_BAD_CHECKSUM
    assert _VALIDATORS["cn_id_checksum"](entity.decision.replacement)

    # 走真实 MCP redact_text 工具路径（DSH/Claude 使用同一入口）
    import mcp_server  # noqa: E402 - 延迟导入：只在需要实测工具路径时加载

    result = mcp_server.HANDLERS["redact_text"](
        {"text": text, "strictness": "standard"})
    payload = json.loads(result["content"][0]["text"])
    assert ID_BAD_CHECKSUM not in payload["redacted"], (
        f"redact_text 未遮住校验位错误的身份证号：{payload['redacted']}")
    assert payload["by_type"].get("身份证号") == 1, (
        f"redact_text 身份证号实体数异常：{payload['by_type']}")
    assert payload["entities"] >= 2, f"redact_text 实体数异常：{payload['entities']}"

    # human 涂黑版：JobManager 导出时会把 decision 统一改成 black，
    # 这里用同一规则模拟文本导出，确认原文被黑块整段覆盖（不是只靠换名）。
    from yinzhan.udm import Decision  # noqa: E402

    human_ents = [item for item in det.entities if item.status != "rejected"]
    human_decisions = {item.id: Decision(strategy="black", black=True,
                                         replacement=None) for item in human_ents}
    human_out = _TMP / "idcard_bad_checksum_human.txt"
    _TP.export(doc, human_ents, human_decisions, human_out, {"mode": "human"})
    human_text = human_out.read_text("utf-8")
    assert ID_BAD_CHECKSUM not in human_text, "human 涂黑版仍残留原身份证号"
    assert "█" * len(ID_BAD_CHECKSUM) in human_text, "human 涂黑版未按原长度盖黑块"


def test_idcard_shape_invalid_date_and_x_suffix_pending():
    """非法出生日期 / X 结尾但校验位错误：同样兜住为 pending。"""
    for value in (ID_BAD_CHECKSUM, ID_BAD_DATE, ID_FUTURE_DATE,
                  "11010119900101123X"):
        path = _write_text_case(
            f"idcard_shape_{value}.txt", f"身份证号：{value}。")
        _doc, det, _session = _detect_text(path)
        cards = [entity for entity in _id_cards(det) if entity.value == value]
        assert len(cards) == 1, f"{value} 形状串未被兜住：{_id_cards(det)}"
        assert cards[0].status == "pending", f"{value} 状态应为 pending"
        assert cards[0].decision.strategy == "surrogate"
        assert cards[0].decision.replacement != value


def test_idcard_shape_bare_number_no_context_pending():
    """全形状方案：无「身份证」上下文的裸号也必须兜住（隐私优先，不做 near 限定）。"""
    cases = (
        f"{ID_BAD_CHECKSUM} 是一串没有上下文提示的裸号码。",
        f"证件号 {ID_BAD_CHECKSUM}，请核对。",       # 弱上下文
        f"身份证号：{ID_BAD_CHECKSUM}。",             # 强上下文
    )
    for index, text in enumerate(cases):
        path = _write_text_case(f"idcard_bare_{index}.txt", text)
        _doc, det, _session = _detect_text(path)
        cards = [entity for entity in _id_cards(det) if entity.value == ID_BAD_CHECKSUM]
        assert len(cards) == 1, f"上下文案例 {index} 漏识别：{_id_cards(det)}"
        assert cards[0].status == "pending", (
            f"上下文案例 {index} 状态异常：{cards[0].status}")


def test_idcard_non_shape_random_no_false_positive():
    """非身份证形号码不误报：订单号/随机 18 位/19 位银行卡/USCC 均不产生 ID_CARD。"""
    cases = (
        "订单号 202610231234567890，金额 5000 元。",     # 前 10 位是时间戳，年份位不合法
        "随机 18 位 987654321012345678 不是身份证。",
        "银行卡 6222021122334456789 请勿外传。",          # 19 位（越界）
        "统一社会信用代码 91310115MA1K4X8L2T 请核对。",   # 含字母
    )
    for index, text in enumerate(cases):
        path = _write_text_case(f"idcard_negative_{index}.txt", text)
        _doc, det, _session = _detect_text(path)
        cards = _id_cards(det)
        assert not cards, f"非身份证形号码误报：{text} -> {[e.value for e in cards]}"


def test_idcard_overlap_bankcard_single_entity():
    """与 cn.bankcard 重叠：同值只保留一个实体，ID_CARD 形状优先、不重复。"""
    assert _VALIDATORS["luhn"](ID_LUHN_SHAPE), "构造值应通过 Luhn"
    assert not _VALIDATORS["cn_id_checksum"](ID_LUHN_SHAPE), "构造值身份证校验位应不合法"
    path = _write_text_case(
        "idcard_overlap_bank.txt", f"卡号：{ID_LUHN_SHAPE}，请核对。")
    _doc, det, _session = _detect_text(path)
    hits = [entity for entity in det.entities if entity.value == ID_LUHN_SHAPE]
    assert len(hits) == 1, (
        f"同值被多条规则重复命中：{[(e.etype, e.status) for e in hits]}")
    assert hits[0].etype == "ID_CARD", f"重叠时实体类型异常：{hits[0].etype}"
    assert hits[0].status == "pending"
    assert det.rules[hits[0].id].id == "cn.idcard_shape"
    assert not [entity for entity in det.entities
                if entity.etype == "BANK" and entity.value.strip() == ID_LUHN_SHAPE], (
        "重叠值不应再产生 BANK 实体")


def test_idcard_shape_strictness_matrix():
    """档位矩阵：standard/strict 进 pending；loose 按 min_conf=0.80 丢弃（已知取舍）。

    loose 档的产品语义是「少打扰」；0.66 的兜底会被丢弃，这是策略已知取舍，
    见 docs/qa/idcard-checksum-policy-20261023.md「档位矩阵」。默认档（Web/MCP
    均为 standard）已覆盖；若把本行为改动，必须同步更新本断言与 QA 文档。
    """
    text = f"身份证号：{ID_BAD_CHECKSUM}。"
    path = _write_text_case("idcard_strictness.txt", text)
    expected = {"standard": "pending", "strict": "pending", "loose": None}
    for strictness, status in expected.items():
        _doc, det, _session = _detect_text(path, strictness)
        cards = _id_cards(det)
        if status is None:
            assert not cards, (
                f"loose 档预期丢弃 0.66 兜底（已知取舍），实际：{cards}")
        else:
            assert len(cards) == 1 and cards[0].status == status, (
                f"{strictness} 档状态异常：{[(e.status, e.confidence) for e in cards]}")


# ────────────────────────── 入口 ──────────────────────────


def main() -> int:
    registry.load_plugins()
    tests = [(name, obj) for name, obj in sorted(globals().items())
             if name.startswith("test_") and callable(obj)]
    failed = 0
    print(f"\n  数据隐站 · PIPL 文本/数值 P0 测试（{len(tests)} 项）")
    print("  " + "─" * 58)
    for name, fn in tests:
        try:
            fn()
        except AssertionError as exc:
            failed += 1
            print(f"  ✗ {name}\n    {exc}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"  ✗ {name}  [{type(exc).__name__}] {exc}")
            import traceback

            traceback.print_exc()
        else:
            print(f"  ✓ {name}")
    print("  " + "─" * 58)
    if failed:
        print(f"  失败 {failed}/{len(tests)}")
        return 1
    print(f"  全部通过（{len(tests)}/{len(tests)}）")
    return 0


if __name__ == "__main__":
    sys.exit(main())
