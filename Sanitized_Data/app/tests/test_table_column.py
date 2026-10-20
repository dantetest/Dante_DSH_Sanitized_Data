# -*- coding: utf-8 -*-
"""表格表头列级检测器（``table_column``）P0 单测 —— task-18。

运行（仓库任意目录，必须 -X utf8）：

    python -X utf8 app/tests/test_table_column.py

也可用 pytest 收集（函数名 ``test_*``，内部惰性初始化引擎）。
只写系统临时目录（``YINZHAN_DATA_DIR`` / ``YINZHAN_OUTPUT_DIR`` 指向 temp），不污染仓库 data/。

覆盖（对应 task-18 验收）：
1. 检测器注册、规则包加载、scope/priority/策略；
2. 原型 xlsx 10/10 整列覆盖（含缺位身份证、带空格手机号、无后缀单位、文本金额）；
3. 原型 docx 8/8 整列覆盖；
4. 表头定位：表头在第 3 行可识别；"名称"等泛词不默认；重复表头/空表不产生实体；
5. 数值/公式单元格跳过（交给 numeric_cell）；纯数值手机号列不处理（P1 边界）；
6. 值形状守卫（"无/N/A"跳过）与 exclude_validate 让位（合法 USCC 不被当证件号）；
7. status: pending 能力；
8. 既有 5 样例：有/无 table-cn 对比，不丢失既有实体，新增实体仅限可解释白名单；
9. 导出 + 出口自检（residual=0）。
"""

from __future__ import annotations

import json
import os
import sys
import tempfile
import traceback
from pathlib import Path

# 必须在 import yinzhan.config 之前把数据/输出目录指向临时目录（不污染仓库）
_TMP = Path(tempfile.mkdtemp(prefix="yz_table_test_"))
os.environ.setdefault("YINZHAN_DATA_DIR", str(_TMP / "data"))
os.environ.setdefault("YINZHAN_OUTPUT_DIR", str(_TMP / "out"))

REPO = Path(__file__).resolve().parents[2]
SERVER = REPO / "app" / "server"
FIXTURES = Path(__file__).resolve().parent / "fixtures" / "table"
SAMPLES = REPO / "samples"
sys.path.insert(0, str(SERVER))

from openpyxl import Workbook                        # noqa: E402
from docx import Document as DocxDocument            # noqa: E402

from yinzhan import pipeline, registry, rules as rules_mod   # noqa: E402
from yinzhan.ocr import get_ocr                      # noqa: E402
from yinzhan.udm import DetectContext, ParseContext, Rule    # noqa: E402

#: table-cn 自有规则 id（11 条；新增/删除需同步本测试与 docs/impl/table.md）
RULE_IDS = {
    "table.person", "table.contact", "table.id_card", "table.phone",
    "table.address", "table.email", "table.org", "table.bank_card",
    "table.account", "table.amount", "table.txn_no",
}

#: 既有样例上"可解释的新增实体"白名单（见 docs/impl/table.md §5）：
#   - 王建国：客户名称列使 NER 低置信候选通过闸门（默认 strictness 下原本漏检）
#   - 6228…：账号列中 Luhn 不合法、正则漏掉的银行卡，由列语义补上
ALLOWED_NEW_ON_SAMPLES = {
    ("PERSON", "王建国"),
    ("ACCOUNT", "6228 4801 2233 4455 667"),
}

_ENGINE: tuple | None = None
_PASS = {"ok": 0, "fail": 0}


# ────────────────────────────── 断言工具 ──────────────────────────────


def check(name: str, cond: bool, detail: str = "") -> bool:
    if cond:
        _PASS["ok"] += 1
        print(f"  [PASS] {name}" + (f" - {detail}" if detail else ""))
        return True
    _PASS["fail"] += 1
    print(f"  [FAIL] {name}" + (f" - {detail}" if detail else ""))
    raise AssertionError(name + (f": {detail}" if detail else ""))


def section(title: str) -> None:
    print(f"\n── {title} ──")


def _engine():
    global _ENGINE
    if _ENGINE is None:
        errors = registry.load_plugins()
        if errors:
            raise RuntimeError(f"插件加载失败: {errors}")
        packs = rules_mod.load_packs()
        ctx = ParseContext(ocr=get_ocr(), workdir=str(_TMP / "work"))
        _ENGINE = (packs, ctx, rules_mod.load_errors())
    return _ENGINE


def _packs() -> dict:
    return _engine()[0]


def _ctx() -> ParseContext:
    return _engine()[1]


# ────────────────────────────── 夹具构造 ──────────────────────────────


def _load_fixture(name: str) -> dict:
    return json.loads((FIXTURES / name).read_text("utf-8"))


def _build_xlsx(path: Path, sheet: str, headers: list, rows: list,
                junk_rows: list | None = None) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    for junk in junk_rows or []:
        ws.append(junk)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(str(path))
    return path


def _build_docx(path: Path, headers: list, rows: list) -> Path:
    doc = DocxDocument()
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    for col, header in enumerate(headers):
        table.cell(0, col).text = str(header)
    for row_index, row in enumerate(rows, start=1):
        for col, value in enumerate(row):
            table.cell(row_index, col).text = str(value)
    doc.save(str(path))
    return path


def _parse(path: Path):
    return registry.find_parser(path).parse(path, _ctx())


def _detect(path: Path, table_only: bool = False):
    doc = _parse(path)
    rules = [rule for rule in rules_mod.compile_rules(_packs(), doc.fmt) if rule.enabled]
    if table_only:
        rules = [rule for rule in rules if rule.pack == "table-cn"]
    det = pipeline.detect(doc, rules, "standard")
    return doc, det


def _table_entities(det) -> list[tuple[str, str, str]]:
    """只取 table-cn 规则产生的实体：(rule_id, etype, value)。"""
    out: list[tuple[str, str, str]] = []
    for entity in det.entities:
        rule = det.rules.get(entity.id)
        if rule is not None and rule.pack == "table-cn":
            out.append((rule.id, entity.etype, entity.value))
    return out


def _work(name: str) -> Path:
    path = _TMP / name
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


# ────────────────────────────── 1. 规则包 / 检测器 ──────────────────────────────


def test_rule_pack_and_detector():
    section("1. 检测器注册与规则包加载")
    packs, _ctx_, load_errors = _engine()
    check("插件加载无错误", not load_errors, str(load_errors))
    check("table_column 检测器已注册", "table_column" in registry.all_detectors())
    check("table-cn 包存在", "table-cn" in packs)

    own = [rule for rule in packs["table-cn"] if rule.pack == "table-cn"]
    ids = {rule.id for rule in own}
    check(f"table-cn 自有规则 {len(RULE_IDS)} 条", len(own) == len(RULE_IDS), f"实际 {len(own)}")
    check("规则 id 与测试清单一致", ids == RULE_IDS, f"缺少 {sorted(RULE_IDS - ids)} 多出 {sorted(ids - RULE_IDS)}")
    for rule in own:
        check(f"{rule.id} scope 含 xlsx/docx/text（P1-B CSV/TSV）",
              set(rule.scope) >= {"xlsx", "docx", "text"}, str(rule.scope))
        check(f"{rule.id} priority < 15（让位既有规则）", rule.priority < 15, str(rule.priority))
        check(f"{rule.id} 替换策略 surrogate",
              rule.replace.get("strategy") == "surrogate", str(rule.replace))
        check(f"{rule.id} detect 类型 table_column",
              [item.get("type") for item in rule.detect] == ["table_column"],
              str([item.get("type") for item in rule.detect]))

    # 格式编译：table 规则在 xlsx/docx + P1-B 的 text(CSV/TSV) 生效；include 进来的 common 仍生效
    for fmt, expected in (("xlsx", len(RULE_IDS)), ("docx", len(RULE_IDS)),
                          ("text", len(RULE_IDS)),
                          ("pdf", 0), ("image", 0), ("tiff", 0), ("mesh", 0)):
        got = [rule for rule in rules_mod.compile_rules(packs, fmt) if rule.pack == "table-cn"]
        check(f"{fmt} 格式包含 {expected} 条 table 规则", len(got) == expected, f"实际 {len(got)}")
    selected = rules_mod.compile_rules(packs, "xlsx", ["table-cn"])
    selected_ids = {rule.id for rule in selected}
    check("只选 table-cn 包时 include common 生效",
          {"table.person", "cn.idcard", "cn.phone"} <= selected_ids,
          f"共 {len(selected_ids)} 条")


# ────────────────────────────── 2/3. 原型覆盖 ──────────────────────────────


def _check_prototype(fixture: str, builder) -> None:
    spec = _load_fixture(fixture)
    path = _work(fixture.replace(".json", "." + spec["kind"]))
    builder(path, spec)
    _doc, det = _detect(path, table_only=True)
    got = _table_entities(det)
    expected = [(item["rule"], item["etype"], item["value"]) for item in spec["expected"]]
    check(f"{fixture} 期望 {len(expected)} 个实体", len(got) == len(expected),
          f"实际 {len(got)}: {got}")
    for item in expected:
        check(f"{item[2]!r} → {item[0]}", item in got)
    # span 必须精确覆盖单元格文本（导出时替换不越界）
    by_value = {entity.value: entity for entity in det.entities}
    for item in expected:
        entity = by_value.get(item[2])
        if entity is None:
            check(f"{item[2]!r} 有实体可查 span", False)
            continue
        span = entity.spans[0]
        check(f"{item[2]!r} span 精确覆盖",
              span.start == 0 and span.end == len(item[2]),
              f"span=({span.start},{span.end}) len={len(item[2])}")
        check(f"{item[2]!r} origin=table", entity.origin == "table", entity.origin)
        check(f"{item[2]!r} confidence>=0.85", entity.confidence >= 0.85, str(entity.confidence))


def _build_prototype_xlsx(path: Path, spec: dict) -> Path:
    return _build_xlsx(path, spec["sheet"], spec["headers"], spec["rows"])


def _build_prototype_docx(path: Path, spec: dict) -> Path:
    return _build_docx(path, spec["headers"], spec["rows"])


def test_prototype_xlsx():
    section("2. 原型 xlsx 整列覆盖（10/10）")
    _check_prototype("prototype_xlsx.json", _build_prototype_xlsx)


def test_prototype_docx():
    section("3. 原型 docx 整列覆盖（8/8）")
    _check_prototype("prototype_docx.json", _build_prototype_docx)


# ────────────────────────────── 4. 表头定位 ──────────────────────────────


def test_header_locator():
    section("4. 表头定位与泛词排除")
    # 表头在第 3 行（前两行是标题/说明）
    path = _work("header_row3.xlsx")
    _build_xlsx(path, "表", ["姓名", "手机号"],
                [["张伟", "13817264451"], ["李四", "13901827733"]],
                junk_rows=[["2026 年人员名单"], []])
    _doc, det = _detect(path, table_only=True)
    values = {item[2] for item in _table_entities(det)}
    check("表头在第 3 行可识别", values == {"张伟", "李四", "13817264451", "13901827733"},
          str(values))

    # 泛词表头（名称/结果/备注）不得默认生效
    path = _work("generic_header.xlsx")
    _build_xlsx(path, "表", ["名称", "结果", "备注"], [["笔记本电脑", "正常", "无"]])
    _doc, det = _detect(path, table_only=True)
    check("泛词表头不产生实体", not _table_entities(det), str(_table_entities(det)))

    # 表头文本复现在数据行（分页/重复表头）→ 跳过，不生成实体
    path = _work("repeated_header.xlsx")
    _build_xlsx(path, "表", ["姓名"], [["姓名"], ["张伟"]])
    _doc, det = _detect(path, table_only=True)
    check("重复表头不生成实体", {item[2] for item in _table_entities(det)} == {"张伟"},
          str(_table_entities(det)))

    # 表头带空格/全角：归一化后精确匹配
    path = _work("space_header.xlsx")
    _build_xlsx(path, "表", [" 姓 名 ", "手机号码"], [["张伟", "13817264451"]])
    _doc, det = _detect(path, table_only=True)
    check("表头空白/全角可归一化匹配",
          {item[2] for item in _table_entities(det)} == {"张伟", "13817264451"},
          str(_table_entities(det)))

    # 无表格（只有段落）的 docx → 0
    path = _work("no_table.docx")
    doc = DocxDocument()
    doc.add_paragraph("姓名：张伟 手机号：13817264451")
    doc.save(str(path))
    _doc, det = _detect(path, table_only=True)
    check("无表格文档不产生列级实体", not _table_entities(det), str(_table_entities(det)))


# ────────────────────────────── 5. 数值/公式跳过 ──────────────────────────────


def test_numeric_and_formula_skip():
    section("5. 数值/公式/日期单元格跳过（交给 numeric_cell）")
    path = _work("numeric_skip.xlsx")
    _build_xlsx(path, "表", ["金额", "姓名"],
                [[28600, "=A2"],            # 数值金额 + 公式姓名
                 ["12,000.00", "李强"]])     # 文本金额
    _doc, det = _detect(path, table_only=True)
    values = {item[2] for item in _table_entities(det)}
    check("数值单元格被跳过、文本金额覆盖",
          values == {"12,000.00", "李强"}, str(values))

    # 纯数值手机号列：P0 不处理（numeric_cell 的账户/金额/日期模式不含手机号）
    path = _work("numeric_phone.xlsx")
    _build_xlsx(path, "表", ["手机号"], [[13817264451], [13901827733]])
    _doc, det = _detect(path, table_only=True)
    check("纯数值手机号列 P0 不处理（已知边界）", not _table_entities(det),
          str(_table_entities(det)))


# ────────────────────────────── 6. 值守卫 / 让位 ──────────────────────────────


def test_value_guard_and_yield():
    section("6. 值形状守卫与 exclude_validate 让位")
    path = _work("guard.xlsx")
    _build_xlsx(path, "表", ["联系电话", "证件号"],
                [["无", "91310000MA1FL1XX0G"],
                 ["N/A", "31010419870312451"],
                 ["138 1726 4451", "310104198703124510"]])
    _doc, det = _detect(path, table_only=True)
    got = _table_entities(det)
    phones = [item[2] for item in got if item[1] == "PHONE"]
    id_cards = [item[2] for item in got if item[1] == "ID_CARD"]
    check("电话列跳过 无/N/A，保留带空格号码", phones == ["138 1726 4451"], str(phones))
    check("证件号列跳过合法 USCC，保留缺位/正常身份证",
          set(id_cards) == {"31010419870312451", "310104198703124510"}, str(id_cards))

    # 全量规则下：USCC 仍由 common.cn.uscc 识别，未被列级规则抢走
    _doc, det_all = _detect(path)
    uscc = [(e.etype, e.value) for e in det_all.entities if e.value == "91310000MA1FL1XX0G"]
    check("合法 USCC 由 cn.uscc 识别（不被证件号列抢走）",
          ("USCC", "91310000MA1FL1XX0G") in uscc, str(uscc))


# ────────────────────────────── 7. pending ──────────────────────────────


def test_pending_status():
    section("7. status: pending 能力（高风险列人工确认）")
    from yinzhan.detectors.table_det import TableColumnDetector

    spec = _load_fixture("prototype_xlsx.json")
    path = _work("pending_proto.xlsx")
    _build_prototype_xlsx(path, spec)
    doc = _parse(path)
    detector = TableColumnDetector.from_spec(
        {"headers": ["姓名"], "header_rows": 3, "status": "pending"})
    rule = Rule(id="table.test_pending", entity="PERSON", label="姓名",
                level="L2", scope=["xlsx"], detect=[], replace={"strategy": "surrogate"})
    entities = detector.detect(doc, rule, DetectContext())
    check("pending 配置生效", bool(entities) and all(e.status == "pending" for e in entities),
          f"{len(entities)} 个, status={[e.status for e in entities]}")


# ────────────────────────────── 8. 既有样例隔离 ──────────────────────────────


def test_existing_samples():
    section("8. 既有 5 样例：不丢既有实体、新增仅限可解释白名单")
    packs, ctx, _ = _engine()
    found_new: dict[str, list] = {}
    for path in sorted(SAMPLES.iterdir()):
        if not path.is_file():
            continue
        doc = _parse(path)
        all_rules = [rule for rule in rules_mod.compile_rules(packs, doc.fmt) if rule.enabled]
        no_table = [rule for rule in all_rules if rule.pack != "table-cn"]
        base = pipeline.detect(doc, no_table, "standard")
        every = pipeline.detect(doc, all_rules, "standard")
        base_keys = {(entity.etype, entity.value) for entity in base.entities}
        every_keys = {(entity.etype, entity.value) for entity in every.entities}
        gone = base_keys - every_keys
        new = every_keys - base_keys
        check(f"{path.name} 既有实体不减少", not gone, f"减少 {sorted(gone)}")
        found_new[path.name] = sorted(new)
        if path.name != "对账单.xlsx":
            check(f"{path.name} 无新增实体", not new, f"新增 {sorted(new)}")
        else:
            unexpected = new - ALLOWED_NEW_ON_SAMPLES
            check("对账单 新增实体均在白名单（可解释）", not unexpected,
                  f"不可解释 {sorted(unexpected)}；实际新增 {sorted(new)}")
    print(f"    既有样例新增实体明细：{json.dumps(found_new, ensure_ascii=False)}")


# ────────────────────────────── 9. 导出冒烟 ──────────────────────────────


def test_export_verify():
    section("9. 导出 + 出口自检（residual=0）")
    spec = _load_fixture("prototype_xlsx.json")
    path = _work("export_proto.xlsx")
    _build_prototype_xlsx(path, spec)
    doc = _parse(path)
    rules = [rule for rule in rules_mod.compile_rules(_packs(), doc.fmt) if rule.enabled]
    det = pipeline.detect(doc, rules, "standard")
    pipeline.decide(det.entities, det.rules, pipeline.Session())
    parser = registry.find_parser(path)
    out = _work("out/prototype_ai.xlsx")
    out.parent.mkdir(parents=True, exist_ok=True)
    pipeline.export(doc, parser, det.entities, out, {"mode": "ai"})
    check("导出文件存在且非空", out.is_file() and out.stat().st_size > 0)
    info = pipeline.verify(out, det.entities, parser, _ctx())
    check("导出复验通过", info.get("verified") is True, str(info))
    check("残留 = 0", info.get("residual") == 0, f"residual={info.get('residual')}")
    # 再解析成品：原值不应还出现在文本里
    re_doc = _parse(out)
    text = "\n".join(block.text or "" for page in re_doc.pages for block in page.blocks)
    leaked = [item["value"] for item in spec["expected"] if item["value"] in text]
    check("成品中无原值残留（抽样）", not leaked, f"残留 {leaked}")


# ────────────────────────────── 主流程 ──────────────────────────────

_TESTS = [
    test_rule_pack_and_detector,
    test_prototype_xlsx,
    test_prototype_docx,
    test_header_locator,
    test_numeric_and_formula_skip,
    test_value_guard_and_yield,
    test_pending_status,
    test_existing_samples,
    test_export_verify,
]


def main() -> int:
    print("=" * 68)
    print("  表格表头列级检测器（table_column）P0 单测 · task-18")
    print("=" * 68)
    failed = 0
    for test in _TESTS:
        try:
            test()
        except AssertionError as exc:
            failed += 1
            print(f"  [FAIL] {test.__name__}: {exc}")
        except Exception:
            failed += 1
            print(f"  [ERROR] {test.__name__}")
            traceback.print_exc()
    total = _PASS["ok"] + _PASS["fail"]
    print("\n" + "=" * 68)
    if failed and _PASS["fail"] == 0:
        _PASS["fail"] = failed
    if failed:
        print(f"  失败 {failed}/{len(_TESTS)} 个测试组；断言 {_PASS['ok']} 通过 / {_PASS['fail']} 失败")
        print("=" * 68)
        return 1
    print(f"  全部通过：{len(_TESTS)} 个测试组，{_PASS['ok']} 项断言全绿")
    print("=" * 68)
    return 0


if __name__ == "__main__":
    sys.exit(main())
