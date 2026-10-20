# -*- coding: utf-8 -*-
"""P1-B：Office 内嵌图片自动脱敏 + 隐藏对象枚举测试。

运行（Windows 控制台请带 -X utf8）：

    python -X utf8 app/tests/test_office_media.py
    cd app\\server && python -X utf8 ..\\tests\\test_office_media.py

设计说明
--------
* 纯离线、不联网；夹具默认生成到临时目录（tempfile.mkdtemp()）。
  想保留夹具人工查看时设置 ``YINZHAN_WRITE_FIXTURES=1``，会写到
  ``app/tests/fixtures/office_media/``（二进制夹具不入库，按需生成）。
* 内嵌图片全部 PIL/cv2 合成：QR 码用 cv2.QRCodeEncoder 生成（mediascan 可稳定检出），
  普通插图不含敏感目标（用于验证“无命中保留原图”）。
* 隐藏对象：DOCX 注入 chart 部件 + OLE embedding（带 rels，让 python-docx 保存时保留），
  XLSX 用 openpyxl 真图表 + 注入 OLE；不可解码图片用伪造 EMF。
* 所有 PII / marker 均为虚构值。
"""
from __future__ import annotations

import io
import json
import os
import shutil
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
SERVER = REPO / "app" / "server"
if str(SERVER) not in sys.path:
    sys.path.insert(0, str(SERVER))

import cv2                                                  # noqa: E402
import numpy as np                                          # noqa: E402
import openpyxl                                             # noqa: E402
import docx                                                 # noqa: E402
from PIL import Image, ImageDraw, ImageFont                   # noqa: E402

from yinzhan import mediascan, media_text, pipeline, registry   # noqa: E402
from yinzhan.jobs import _load_rules                        # noqa: E402
from yinzhan.ocr import get_ocr                             # noqa: E402
from yinzhan.parsers import docx_parser as docx_mod         # noqa: E402
from yinzhan.udm import ParseContext, VerifyError           # noqa: E402

FIXTURE_DIR = Path(__file__).resolve().parent / "fixtures" / "office_media"
BODY_PHONE = "13800138000"
OLE_MARKER = b"OLE-FAKE-BYTES-SECRET"
EMF_MARKER = b"EMF-FAKE-BYTES-SECRET"
TEXTBOX_MARKER = "TEXTBOX-SECRET-MARKER"
CHART_TITLE = "CHARTTITLE-SECRET"
QR_TEXT = "https://example.com/SECRET-QR-9f3a"

# P1.5 文字 OCR 较慢：除“文字脱敏”专项用例外，本测试模块默认关闭文字 OCR，
# 保证 P1-B 视觉/隐藏对象用例的耗时稳定（与产品默认开不冲突，产品侧见 docs）。
_TEXT_OCR_ENV = "YINZHAN_OFFICE_TEXT_OCR"
_saved_text_ocr: str | None = None


def setUpModule() -> None:  # noqa: N802 - unittest 钩子
    global _saved_text_ocr
    _saved_text_ocr = os.environ.get(_TEXT_OCR_ENV)
    os.environ[_TEXT_OCR_ENV] = "0"


def tearDownModule() -> None:  # noqa: N802 - unittest 钩子
    if _saved_text_ocr is None:
        os.environ.pop(_TEXT_OCR_ENV, None)
    else:
        os.environ[_TEXT_OCR_ENV] = _saved_text_ocr


def _text_ocr_on() -> None:
    os.environ[_TEXT_OCR_ENV] = "1"


def _text_ocr_off() -> None:
    os.environ[_TEXT_OCR_ENV] = "0"


# ────────────────────────── 夹具 ──────────────────────────

def _fixtures_root() -> Path:
    if os.environ.get("YINZHAN_WRITE_FIXTURES") == "1":
        FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
        return FIXTURE_DIR
    return Path(tempfile.mkdtemp(prefix="yinzhan_office_media_"))


def _replace_retry(source: Path, target: Path) -> None:
    """Windows 上刚写出的 office 文件偶发被杀软/索引器占用，重试替换。"""
    for attempt in range(10):
        try:
            os.replace(source, target)
            return
        except PermissionError:
            time.sleep(0.2 * (attempt + 1))
    shutil.copyfile(source, target)
    source.unlink()


def _zip_edit(path: Path, *, add: dict[str, bytes] | None = None,
              edit_text: dict[str, callable] | None = None,
              rel_additions: dict[str, str] | None = None) -> None:
    """zip 级编辑：加部件 / 文本转换 / 加 relationship。"""
    with zipfile.ZipFile(path) as archive:
        data = {item.filename: archive.read(item.filename)
                for item in archive.infolist()}
    if rel_additions:
        for rels_name, fragment in rel_additions.items():
            text = data.get(rels_name, b"").decode("utf-8")
            if "</Relationships>" not in text:
                raise AssertionError(f"{rels_name} 不是合法 rels")
            text = text.replace("</Relationships>", fragment + "</Relationships>")
            data[rels_name] = text.encode("utf-8")
    if edit_text:
        for name, transform in edit_text.items():
            data[name] = transform(data[name])
    if add:
        data.update(add)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, blob in data.items():
            archive.writestr(name, blob)
    _replace_retry(tmp, path)


def _make_qr_image(path: Path) -> Path:
    encoder = cv2.QRCodeEncoder_create()
    matrix = encoder.encode(QR_TEXT)
    image = Image.fromarray(matrix).convert("RGB").resize((240, 240), Image.NEAREST)
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)
    return path


def _make_normal_image(path: Path) -> Path:
    image = Image.new("RGB", (240, 160), (235, 232, 222))
    draw = ImageDraw.Draw(image)
    draw.rectangle([10, 10, 230, 150], outline=(30, 30, 30), width=2)
    draw.text((20, 60), "ordinary illustration / no target", fill=(20, 20, 20))
    image.save(path)
    return path


#: P1.5 文字 OCR 测试用的虚构 PII（PIL 默认字体只画 ASCII）
TEXT_PII_LINES = (
    "TEL: 13800138000",
    "ID: 11010519491231002X",
    "EMAIL: test.user@example.com",
    "AMOUNT: 12,345.67 CNY",
)


def _text_font():
    try:
        return ImageFont.load_default(size=40)
    except TypeError:  # pragma: no cover - 老 Pillow
        return ImageFont.load_default()


def _make_text_image(path: Path) -> Path:
    """画一张含 ASCII 手机号/身份证号/邮箱/金额的截图（P1.5 文字 OCR 用）。"""
    image = Image.new("RGB", (900, 520), (250, 248, 242))
    draw = ImageDraw.Draw(image)
    font = _text_font()
    y = 40
    for line in TEXT_PII_LINES:
        draw.text((40, y), line, font=font, fill=(15, 15, 15))
        y += 100
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path)
    return path


def _make_qr_text_image(path: Path, qr: Path) -> Path:
    """QR 码 + 文字叠加图：验证“视觉 + 文字”两条链路叠加脱敏。"""
    image = Image.new("RGB", (1000, 560), (250, 248, 242))
    with Image.open(qr) as qr_image:
        image.paste(qr_image.convert("RGB").resize((240, 240)), (40, 40))
    draw = ImageDraw.Draw(image)
    font = _text_font()
    draw.text((40, 320), "TEL: 13800138000", font=font, fill=(15, 15, 15))
    draw.text((40, 400), "EMAIL: overlay@example.com", font=font, fill=(15, 15, 15))
    image.save(path)
    return path


def _content_types_add(overrides: str = "", defaults: str = "") -> callable:
    def transform(raw: bytes) -> bytes:
        text = raw.decode("utf-8")
        addition = overrides + defaults
        if "</Types>" not in text:
            return raw
        return text.replace("</Types>", addition + "</Types>").encode("utf-8")
    return transform


def _build_docx_hidden(path: Path, qr: Path, normal: Path) -> Path:
    """DOCX：正文 PII + QR 图 + 普通图 + chart/OLE/EMF 隐藏对象。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.add_paragraph(f"委托方：张三 手机号 {BODY_PHONE} 正文内容")
    document.add_picture(str(qr))
    document.add_picture(str(normal))
    document.save(str(path))
    chart_xml = (
        '<?xml version="1.0"?>'
        '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">'
        '<c:chart><c:title><c:tx><c:rich>'
        '<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:r><a:t>{CHART_TITLE}</a:t></a:r></a:p></c:rich></c:tx></c:title>'
        '</c:chart></c:chartSpace>').encode("utf-8")
    _zip_edit(
        path,
        add={
            "word/charts/chart1.xml": chart_xml,
            "word/embeddings/oleObject1.bin": OLE_MARKER,
            "word/media/image9.emf": EMF_MARKER,
        },
        edit_text={"[Content_Types].xml": _content_types_add(
            overrides=(
                '<Override PartName="/word/charts/chart1.xml" '
                'ContentType="application/vnd.openxmlformats-officedocument'
                '.drawingml.chart+xml"/>'
                '<Override PartName="/word/embeddings/oleObject1.bin" '
                'ContentType="application/vnd.openxmlformats-officedocument'
                '.oleObject"/>'),
            defaults='<Default Extension="emf" ContentType="image/x-emf"/>')},
        rel_additions={"word/_rels/document.xml.rels": (
            '<Relationship Id="rIdChart1" Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/chart" Target="charts/chart1.xml"/>'
            '<Relationship Id="rIdOle1" Type="http://schemas.openxmlformats.org/'
            'officeDocument/2006/relationships/oleObject" '
            'Target="embeddings/oleObject1.bin"/>')},
    )
    return path


def _build_docx_clean(path: Path, normal: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.add_paragraph("普通文档：只有一张插图")
    document.add_picture(str(normal))
    document.save(str(path))
    return path


def _build_docx_textbox(path: Path) -> Path:
    """DOCX：document.xml 里注入 w:txbxContent（文本框不可自动脱敏）。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.add_paragraph("正文：文本框测试")
    document.save(str(path))

    def inject(document_raw: bytes) -> bytes:
        text = document_raw.decode("utf-8")
        fragment = (
            '<w:p><w:r><w:txbxContent><w:p><w:r>'
            f'<w:t>{TEXTBOX_MARKER}</w:t>'
            '</w:r></w:p></w:txbxContent></w:r></w:p>')
        return text.replace("</w:body>", fragment + "</w:body>").encode("utf-8")

    _zip_edit(
        path,
        edit_text={"word/document.xml": inject},
        add={"word/drawings/drawing9.xml": (
            '<?xml version="1.0"?><w:drawing xmlns:w="http://schemas.'
            'openxmlformats.org/wordprocessingml/2006/main"><wps:txbx '
            'xmlns:wps="http://schemas.microsoft.com/office/word/2010/'
            'wordprocessingShape"><wps:txbxContent/></wps:txbx></w:drawing>'
        ).encode("utf-8")},
    )
    return path


def _build_xlsx_hidden(path: Path, qr: Path, normal: Path) -> Path:
    """XLSX：单元格 PII + QR 图 + 普通图 + 真图表 + 注入 OLE。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "客户手机号"
    sheet["A2"] = BODY_PHONE
    from openpyxl.chart import BarChart, Reference
    from openpyxl.drawing.image import Image as XLImage

    sheet.add_image(XLImage(str(qr)), "C2")
    sheet.add_image(XLImage(str(normal)), "C20")
    chart = BarChart()
    chart.title = CHART_TITLE
    chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=2))
    sheet.add_chart(chart, "H2")
    workbook.save(str(path))
    _zip_edit(
        path,
        add={"xl/embeddings/oleObject1.bin": OLE_MARKER},
        edit_text={"[Content_Types].xml": _content_types_add(
            overrides=(
                '<Override PartName="/xl/embeddings/oleObject1.bin" '
                'ContentType="application/vnd.openxmlformats-officedocument'
                '.oleObject"/>'))},
    )
    return path


def _build_xlsx_clean(path: Path, normal: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "普通表格"
    from openpyxl.drawing.image import Image as XLImage

    sheet.add_image(XLImage(str(normal)), "C2")
    workbook.save(str(path))
    return path


def _build_docx_text(path: Path, text_image: Path, normal: Path) -> Path:
    """DOCX：一张“文字截图” + 一张普通图（P1.5 文字 OCR 自动脱敏）。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.add_paragraph("正文：内嵌一张含文字信息的截图")
    document.add_picture(str(text_image))
    document.add_picture(str(normal))
    document.save(str(path))
    return path


def _build_docx_qr_text(path: Path, qr_text_image: Path) -> Path:
    """DOCX：QR + 文字叠加图（视觉与文字链路都要命中）。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.add_paragraph("正文：二维码与文字混排截图")
    document.add_picture(str(qr_text_image))
    document.save(str(path))
    return path


def _build_xlsx_text(path: Path, text_image: Path, normal: Path) -> Path:
    """XLSX：一张“文字截图” + 一张普通图。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = "内嵌文字截图测试"
    from openpyxl.drawing.image import Image as XLImage

    sheet.add_image(XLImage(str(text_image)), "C2")
    sheet.add_image(XLImage(str(normal)), "C20")
    workbook.save(str(path))
    return path


class Fixtures:
    root: Path | None = None

    @classmethod
    def get(cls) -> Path:
        if cls.root is None:
            cls.root = _fixtures_root()
            cls.qr = _make_qr_image(cls.root / "qr.png")
            cls.normal = _make_normal_image(cls.root / "normal.png")
        return cls.root

    @classmethod
    def docx_hidden(cls) -> Path:
        root = cls.get()
        path = root / "docx_hidden.docx"
        if not path.exists():
            _build_docx_hidden(path, cls.qr, cls.normal)
        return path

    @classmethod
    def docx_clean(cls) -> Path:
        root = cls.get()
        path = root / "docx_clean.docx"
        if not path.exists():
            _build_docx_clean(path, cls.normal)
        return path

    @classmethod
    def docx_textbox(cls) -> Path:
        root = cls.get()
        path = root / "docx_textbox.docx"
        if not path.exists():
            _build_docx_textbox(path)
        return path

    @classmethod
    def xlsx_hidden(cls) -> Path:
        root = cls.get()
        path = root / "xlsx_hidden.xlsx"
        if not path.exists():
            _build_xlsx_hidden(path, cls.qr, cls.normal)
        return path

    @classmethod
    def xlsx_clean(cls) -> Path:
        root = cls.get()
        path = root / "xlsx_clean.xlsx"
        if not path.exists():
            _build_xlsx_clean(path, cls.normal)
        return path

    # ── P1.5 文字 OCR 夹具 ──
    @classmethod
    def text_image(cls) -> Path:
        root = cls.get()
        path = root / "text_pii.png"
        if not path.exists():
            _make_text_image(path)
        return path

    @classmethod
    def qr_text_image(cls) -> Path:
        root = cls.get()
        path = root / "qr_text.png"
        if not path.exists():
            _make_qr_text_image(path, cls.qr)
        return path

    @classmethod
    def docx_text(cls) -> Path:
        root = cls.get()
        path = root / "docx_text.docx"
        if not path.exists():
            _build_docx_text(path, cls.text_image(), cls.normal)
        return path

    @classmethod
    def docx_qr_text(cls) -> Path:
        root = cls.get()
        path = root / "docx_qr_text.docx"
        if not path.exists():
            _build_docx_qr_text(path, cls.qr_text_image())
        return path

    @classmethod
    def xlsx_text(cls) -> Path:
        root = cls.get()
        path = root / "xlsx_text.xlsx"
        if not path.exists():
            _build_xlsx_text(path, cls.text_image(), cls.normal)
        return path


# ────────────────────────── 工具 ──────────────────────────

def _ctx() -> ParseContext:
    return ParseContext(ocr=get_ocr(), workdir=str(Fixtures.get()))


def _parser(path: Path):
    registry.load_plugins()
    return registry.find_parser(path)


def _detect(doc):
    rules, _pack, _index = _load_rules(doc.fmt)
    det = pipeline.detect(doc, rules, "standard")
    pipeline.decide(det.entities, det.rules, pipeline.Session())
    return det


def _zip_media_parts(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return [name for name in archive.namelist()
                if "/media/" in name.lower()
                and name.lower().endswith((".png", ".jpg", ".jpeg", ".gif",
                                           ".bmp", ".tif", ".tiff", ".webp"))]


def _part_bytes(path: Path, part: str) -> bytes:
    with zipfile.ZipFile(path) as archive:
        return archive.read(part)


def _scan_marks(image_bytes: bytes) -> list[dict]:
    with Image.open(io.BytesIO(image_bytes)) as image:
        frame = image.convert("RGB").copy()
    marks, _warnings = mediascan.analyze(
        frame, options={"text_boxes": [], "fmt": "image"})
    return marks


def _find_qr_part(path: Path) -> tuple[str, list[dict]]:
    for part in _zip_media_parts(path):
        marks = _scan_marks(_part_bytes(path, part))
        kinds = {str(mark.get("kind")) for mark in marks}
        if "qrcode" in kinds or "barcode" in kinds:
            return part, marks
    raise AssertionError("源文件里找不到可检出的二维码图片")


def _find_clean_part(path: Path, exclude: str) -> str:
    for part in _zip_media_parts(path):
        if part == exclude:
            continue
        if not _scan_marks(_part_bytes(path, part)):
            return part
    raise AssertionError("源文件里找不到无敏感目标的普通图片")


def _image_mean(image_bytes: bytes) -> float:
    with Image.open(io.BytesIO(image_bytes)) as image:
        return float(np.asarray(image.convert("L"), dtype=np.float64).mean())


def _zip_has_part(path: Path, part: str) -> bool:
    with zipfile.ZipFile(path) as archive:
        return part in archive.namelist()


def _text_hits(image_bytes: bytes, suffix: str = ".png") -> int:
    """直接调用 Lead 的 media_text helper 检测文字敏感命中数。"""
    registry.load_plugins()
    _new_bytes, hits, _warns = media_text.redact_image_text(image_bytes, suffix)
    return int(hits or 0)


def _find_text_part(path: Path) -> str:
    """找出 zip 里“含文字敏感信息”的图片部件（P1.5 helper 可检出）。"""
    for part in _zip_media_parts(path):
        suffix = Path(part).suffix.lower()
        if suffix not in media_text.SUPPORTED_SUFFIXES:
            continue
        if _text_hits(_part_bytes(path, part), suffix) > 0:
            return part
    raise AssertionError("源文件里找不到 helper 可检出文字敏感信息的图片")


# ────────────────────────── 1. 解析：meta 字段 ──────────────────────────

class TestDocxParseMeta(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.parser = _parser(Fixtures.docx_hidden())
        cls.doc = cls.parser.parse(Fixtures.docx_hidden(), _ctx())
        cls.media = (cls.doc.meta or {}).get("embedded_media") or {}

    def test_meta_shape(self):
        media = self.media
        self.assertGreaterEqual(int(media.get("image_count") or 0), 2)
        self.assertGreaterEqual(int(media.get("hidden_object_count") or 0), 2)
        self.assertEqual(int(media.get("count") or 0),
                         len(media.get("hidden_objects") or []))
        self.assertTrue(media.get("auto_redact"))
        self.assertIn("word/charts/chart1.xml", media.get("hidden_objects") or [])
        self.assertIn("word/embeddings/oleObject1.bin", media.get("hidden_objects") or [])
        self.assertTrue(any(name.lower().endswith(".emf")
                            for name in (media.get("unsupported_images") or [])))

    def test_hidden_objects_top_level(self):
        hidden = (self.doc.meta or {}).get("hidden_objects")
        self.assertIsInstance(hidden, list)
        self.assertGreaterEqual(len(hidden), 2)

    def test_meta_json_serializable(self):
        json.dumps(self.doc.meta, ensure_ascii=False)

    def test_warnings(self):
        joined = " ".join(self.doc.warnings)
        self.assertIn("内嵌图片", joined)
        self.assertIn("隐藏对象", joined)
        self.assertIn("自动扫描", joined)


class TestXlsxParseMeta(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.parser = _parser(Fixtures.xlsx_hidden())
        cls.doc = cls.parser.parse(Fixtures.xlsx_hidden(), _ctx())
        cls.media = (cls.doc.meta or {}).get("embedded_media") or {}

    def test_meta_shape(self):
        media = self.media
        self.assertGreaterEqual(int(media.get("image_count") or 0), 2)
        self.assertGreaterEqual(int(media.get("hidden_object_count") or 0), 2)
        self.assertEqual(int(media.get("count") or 0),
                         len(media.get("hidden_objects") or []))
        hidden = media.get("hidden_objects") or []
        self.assertTrue(any("charts" in name.lower() for name in hidden))
        self.assertTrue(any("embeddings" in name.lower() for name in hidden))
        self.assertTrue(media.get("auto_redact"))

    def test_warnings(self):
        joined = " ".join(self.doc.warnings)
        self.assertIn("内嵌图片", joined)
        self.assertIn("隐藏对象", joined)


# ────────────────────────── 2. DOCX 导出 ──────────────────────────

class TestDocxExportDefault(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source = Fixtures.docx_hidden()
        cls.parser = _parser(cls.source)
        cls.doc = cls.parser.parse(cls.source, _ctx())
        cls.det = _detect(cls.doc)
        cls.qr_part, _marks = _find_qr_part(cls.source)
        cls.clean_part = _find_clean_part(cls.source, exclude=cls.qr_part)
        cls.out = Fixtures.get() / "docx_hidden_default_out.docx"
        cls.res = pipeline.export(cls.doc, cls.parser, cls.det.entities,
                                  cls.out, {"mode": "human"})
        cls.stats = (cls.doc.extra or {}).get("office_media") or {}

    def test_images_redacted(self):
        marks = _scan_marks(_part_bytes(self.out, self.qr_part))
        self.assertEqual(marks, [], f"二维码图片未脱敏：{marks}")
        self.assertLess(_image_mean(_part_bytes(self.out, self.qr_part)), 10.0)
        self.assertGreaterEqual(int(self.stats.get("images_redacted") or 0), 1)
        self.assertIn("已自动脱敏", " ".join(self.res.warnings))

    def test_clean_image_preserved(self):
        self.assertEqual(_part_bytes(self.source, self.clean_part),
                         _part_bytes(self.out, self.clean_part),
                         "无敏感目标的普通图片不应被改动")
        self.assertIn("未检出", " ".join(self.res.warnings))

    def test_hidden_objects_kept_by_default(self):
        self.assertGreaterEqual(int(self.stats.get("hidden_kept") or 0), 1)
        output_meta = self.parser.parse(self.out, _ctx()).meta
        media = output_meta.get("embedded_media") or {}
        self.assertGreaterEqual(int(media.get("hidden_object_count") or 0), 1)
        self.assertIn("保留", " ".join(self.res.warnings))

    def test_verify_needs_confirmation(self):
        with self.assertRaises(VerifyError):
            pipeline.verify(self.out, self.det.entities, self.parser, _ctx())
        ver = pipeline.verify(self.out, self.det.entities, self.parser, _ctx(),
                              options={"allow_unprocessed_media": True})
        self.assertTrue(ver.get("verified"))
        self.assertEqual(ver.get("residual"), 0)

    def test_office_can_open(self):
        document = docx.Document(str(self.out))
        self.assertTrue(document.paragraphs)

    def test_pii_redacted(self):
        document = docx.Document(str(self.out))
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertNotIn(BODY_PHONE, text)


class TestDocxExportStrict(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source = Fixtures.docx_hidden()
        cls.parser = _parser(cls.source)
        cls.doc = cls.parser.parse(cls.source, _ctx())
        cls.qr_part, _marks = _find_qr_part(cls.source)
        cls.out = Fixtures.get() / "docx_hidden_strict_out.docx"
        cls.res = pipeline.export(cls.doc, cls.parser, [], cls.out,
                                  {"mode": "human", "strict_office_media": True})
        cls.stats = (cls.doc.extra or {}).get("office_media") or {}

    def test_hidden_objects_stripped(self):
        output_meta = self.parser.parse(self.out, _ctx()).meta
        media = output_meta.get("embedded_media") or {}
        self.assertEqual(int(media.get("hidden_object_count") or 0), 0)
        self.assertFalse(_zip_has_part(self.out, "word/charts/chart1.xml"))
        self.assertFalse(_zip_has_part(self.out, "word/embeddings/oleObject1.bin"))
        self.assertGreaterEqual(int(self.stats.get("hidden_stripped") or 0), 1)

    def test_all_images_blanked(self):
        for part in _zip_media_parts(self.out):
            self.assertEqual(_scan_marks(_part_bytes(self.out, part)), [])
            self.assertLess(_image_mean(_part_bytes(self.out, part)), 10.0,
                            f"strict 模式下 {part} 未置空白")

    def test_verify_passes_without_confirmation(self):
        ver = pipeline.verify(self.out, [], self.parser, _ctx())
        self.assertTrue(ver.get("verified"))
        self.assertEqual(ver.get("media", {}).get("count"), 0)

    def test_office_can_open(self):
        self.assertTrue(docx.Document(str(self.out)).paragraphs)


class TestDocxCleanMedia(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source = Fixtures.docx_clean()
        cls.parser = _parser(cls.source)
        cls.doc = cls.parser.parse(cls.source, _ctx())
        cls.out = Fixtures.get() / "docx_clean_out.docx"
        cls.res = pipeline.export(cls.doc, cls.parser, [], cls.out, {"mode": "human"})

    def test_parse_not_counted(self):
        media = (self.doc.meta or {}).get("embedded_media") or {}
        self.assertEqual(int(media.get("count") or 0), 0)
        self.assertEqual(int(media.get("image_count") or 0), 1)

    def test_clean_image_untouched(self):
        parts = _zip_media_parts(self.source)
        self.assertEqual(len(parts), 1)
        self.assertEqual(_part_bytes(self.source, parts[0]),
                         _part_bytes(self.out, parts[0]))
        self.assertIn("未检出", " ".join(self.res.warnings))

    def test_verify_passes(self):
        ver = pipeline.verify(self.out, [], self.parser, _ctx())
        self.assertTrue(ver.get("verified"))


class TestDocxTextbox(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source = Fixtures.docx_textbox()
        cls.parser = _parser(cls.source)
        cls.doc = cls.parser.parse(cls.source, _ctx())

    def test_textbox_counted(self):
        media = (self.doc.meta or {}).get("embedded_media") or {}
        self.assertGreaterEqual(int(media.get("textboxes") or 0), 1)
        hidden = media.get("hidden_objects") or []
        self.assertTrue(any(str(item).startswith("textbox:") for item in hidden))

    def test_strict_removes_textbox_xml(self):
        out = Fixtures.get() / "docx_textbox_strict_out.docx"
        parser = _parser(self.source)
        doc = parser.parse(self.source, _ctx())
        pipeline.export(doc, parser, [], out,
                        {"mode": "human", "strict_office_media": True})
        raw = _part_bytes(out, "word/document.xml")
        self.assertNotIn(b"txbxContent", raw)
        output_media = (parser.parse(out, _ctx()).meta or {}).get("embedded_media") or {}
        self.assertEqual(int(output_media.get("textboxes") or 0), 0)
        self.assertTrue(docx.Document(str(out)).paragraphs)


# ────────────────────────── 3. XLSX 导出 ──────────────────────────

class TestXlsxExportDefault(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source = Fixtures.xlsx_hidden()
        cls.parser = _parser(cls.source)
        cls.doc = cls.parser.parse(cls.source, _ctx())
        cls.qr_part, _marks = _find_qr_part(cls.source)
        cls.clean_part = _find_clean_part(cls.source, exclude=cls.qr_part)
        cls.out = Fixtures.get() / "xlsx_hidden_default_out.xlsx"
        cls.res = pipeline.export(cls.doc, cls.parser, [], cls.out, {"mode": "human"})
        cls.stats = (cls.doc.extra or {}).get("office_media") or {}

    def test_qr_redacted(self):
        self.assertEqual(_scan_marks(_part_bytes(self.out, self.qr_part)), [])
        self.assertLess(_image_mean(_part_bytes(self.out, self.qr_part)), 10.0)

    def test_clean_image_preserved(self):
        self.assertEqual(_part_bytes(self.source, self.clean_part),
                         _part_bytes(self.out, self.clean_part))

    def test_chart_kept_needs_confirmation(self):
        self.assertGreaterEqual(int(self.stats.get("hidden_kept") or 0), 1)
        with self.assertRaises(VerifyError):
            pipeline.verify(self.out, [], self.parser, _ctx())
        ver = pipeline.verify(self.out, [], self.parser, _ctx(),
                              options={"allow_unprocessed_media": True})
        self.assertTrue(ver.get("verified"))

    def test_excel_can_open(self):
        workbook = openpyxl.load_workbook(str(self.out))
        self.assertIn("客户手机号", [cell.value for row in workbook.active.iter_rows()
                                    for cell in row if cell.value])


class TestXlsxExportStrict(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source = Fixtures.xlsx_hidden()
        cls.parser = _parser(cls.source)
        cls.doc = cls.parser.parse(cls.source, _ctx())
        cls.out = Fixtures.get() / "xlsx_hidden_strict_out.xlsx"
        cls.res = pipeline.export(cls.doc, cls.parser, [], cls.out,
                                  {"mode": "human", "strict_office_media": True})

    def test_hidden_stripped(self):
        output_meta = (self.parser.parse(self.out, _ctx()).meta or {}).get(
            "embedded_media") or {}
        self.assertEqual(int(output_meta.get("hidden_object_count") or 0), 0)
        with zipfile.ZipFile(self.out) as archive:
            names = [name.lower() for name in archive.namelist()]
        self.assertFalse(any("/charts/" in name for name in names))
        self.assertFalse(any("/embeddings/" in name for name in names))

    def test_images_blanked_and_verify(self):
        for part in _zip_media_parts(self.out):
            self.assertLess(_image_mean(_part_bytes(self.out, part)), 10.0)
            self.assertEqual(_scan_marks(_part_bytes(self.out, part)), [])
        ver = pipeline.verify(self.out, [], self.parser, _ctx())
        self.assertTrue(ver.get("verified"))
        self.assertTrue(openpyxl.load_workbook(str(self.out)).active)


# ────────────────────────── 4. 自检 / 开关 / 负样本 ──────────────────────────

class TestSelfCheckBlocks(unittest.TestCase):
    def setUp(self):
        self.parser = _parser(Fixtures.docx_hidden())

    def test_export_blocks_when_paint_fails(self):
        doc = self.parser.parse(Fixtures.docx_hidden(), _ctx())
        original = docx_mod._office_paint_marks
        docx_mod._office_paint_marks = lambda pil, marks, warnings: (pil, [])  # type: ignore[assignment]
        try:
            out = Fixtures.get() / "docx_selfcheck_out.docx"
            with self.assertRaises(Exception) as caught:
                pipeline.export(doc, self.parser, [], out, {"mode": "human"})
            self.assertIn("复检失败", str(caught.exception))
            self.assertFalse(out.exists(), "自检失败时应删除成品")
        finally:
            docx_mod._office_paint_marks = original  # type: ignore[assignment]


class TestRedactSwitch(unittest.TestCase):
    def setUp(self):
        self.parser = _parser(Fixtures.docx_hidden())

    def test_export_option_disables_redaction(self):
        doc = self.parser.parse(Fixtures.docx_hidden(), _ctx())
        qr_part, _marks = _find_qr_part(Fixtures.docx_hidden())
        out = Fixtures.get() / "docx_redact_off.docx"
        res = pipeline.export(doc, self.parser, [], out,
                              {"mode": "human", "redact_office_media": False})
        self.assertIn("自动脱敏已关闭", " ".join(res.warnings))
        self.assertTrue(_scan_marks(_part_bytes(out, qr_part)),
                        "关闭自动脱敏后二维码图片应保持原样")

    def test_parse_env_switch_counts_images(self):
        original = docx_mod.OFFICE_MEDIA_REDACT_DEFAULT
        docx_mod.OFFICE_MEDIA_REDACT_DEFAULT = False
        try:
            parser = _parser(Fixtures.docx_hidden())
            doc = parser.parse(Fixtures.docx_hidden(), _ctx())
            media = (doc.meta or {}).get("embedded_media") or {}
            self.assertFalse(media.get("auto_redact"))
            self.assertGreaterEqual(int(media.get("count") or 0),
                                    int(media.get("image_count") or 0))
        finally:
            docx_mod.OFFICE_MEDIA_REDACT_DEFAULT = original


# ────────────────────────── P1.5：内嵌图片纯文字 OCR 规则脱敏 ──────────────────────────

class TestDocxTextRedaction(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _text_ocr_on()
        cls.source = Fixtures.docx_text()
        cls.parser = _parser(cls.source)
        cls.doc = cls.parser.parse(cls.source, _ctx())
        cls.out = Fixtures.get() / "docx_text_out.docx"
        cls.res = pipeline.export(cls.doc, cls.parser, [], cls.out, {"mode": "human"})
        cls.stats = (cls.doc.extra or {}).get("office_media") or {}
        cls.text_part = _find_text_part(cls.source)

    @classmethod
    def tearDownClass(cls):
        _text_ocr_off()

    def test_parse_announces_text_ocr(self):
        media = (self.doc.meta or {}).get("embedded_media") or {}
        self.assertTrue(media.get("text_ocr_supported"))
        self.assertIn("OCR 规则脱敏", " ".join(self.doc.warnings))

    def test_text_hits_recorded(self):
        self.assertGreaterEqual(int(self.stats.get("text_hits") or 0), 1)
        self.assertGreaterEqual(int(self.stats.get("text_images_processed") or 0), 1)
        self.assertGreaterEqual(int(self.stats.get("text_images_redacted") or 0), 1)
        self.assertIn("内嵌图片文字 OCR", " ".join(self.res.warnings))

    def test_text_image_redacted_and_rescan_zero(self):
        source_bytes = _part_bytes(self.source, self.text_part)
        output_bytes = _part_bytes(self.out, self.text_part)
        self.assertNotEqual(source_bytes, output_bytes, "文字截图应被重写脱敏")
        self.assertEqual(_text_hits(output_bytes, Path(self.text_part).suffix), 0,
                         "成品文字截图重扫 helper 仍有命中")

    def test_normal_image_preserved(self):
        others = [part for part in _zip_media_parts(self.source)
                  if part != self.text_part]
        self.assertTrue(others)
        for part in others:
            self.assertEqual(_part_bytes(self.source, part),
                             _part_bytes(self.out, part),
                             f"无文字/视觉目标的图片 {part} 不应被改动")

    def test_verify_and_office_open(self):
        ver = pipeline.verify(self.out, [], self.parser, _ctx())
        self.assertTrue(ver.get("verified"))
        self.assertEqual(ver.get("residual"), 0)
        self.assertEqual(int((self.stats or {}).get("text_verify_hits") or 0), 0)
        self.assertTrue(docx.Document(str(self.out)).paragraphs)


class TestXlsxTextRedaction(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _text_ocr_on()
        cls.source = Fixtures.xlsx_text()
        cls.parser = _parser(cls.source)
        cls.doc = cls.parser.parse(cls.source, _ctx())
        cls.out = Fixtures.get() / "xlsx_text_out.xlsx"
        cls.res = pipeline.export(cls.doc, cls.parser, [], cls.out, {"mode": "human"})
        cls.stats = (cls.doc.extra or {}).get("office_media") or {}
        cls.text_part = _find_text_part(cls.source)

    @classmethod
    def tearDownClass(cls):
        _text_ocr_off()

    def test_text_hits_recorded(self):
        self.assertGreaterEqual(int(self.stats.get("text_hits") or 0), 1)
        self.assertEqual(int(self.stats.get("text_verify_hits") or 0), 0)

    def test_image_redacted(self):
        output_bytes = _part_bytes(self.out, self.text_part)
        self.assertEqual(_text_hits(output_bytes, Path(self.text_part).suffix), 0)

    def test_verify_and_excel_open(self):
        ver = pipeline.verify(self.out, [], self.parser, _ctx())
        self.assertTrue(ver.get("verified"))
        self.assertTrue(openpyxl.load_workbook(str(self.out)).active)


class TestVisualPlusTextOverlap(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _text_ocr_on()
        cls.source = Fixtures.docx_qr_text()
        cls.parser = _parser(cls.source)
        cls.doc = cls.parser.parse(cls.source, _ctx())
        cls.out = Fixtures.get() / "docx_qr_text_out.docx"
        cls.res = pipeline.export(cls.doc, cls.parser, [], cls.out, {"mode": "human"})
        cls.stats = (cls.doc.extra or {}).get("office_media") or {}
        cls.part = _zip_media_parts(cls.source)[0]

    @classmethod
    def tearDownClass(cls):
        _text_ocr_off()

    def test_both_channels_hit(self):
        self.assertGreaterEqual(int(self.stats.get("images_redacted") or 0), 1)
        self.assertIn("qrcode", self.stats.get("hits_by_kind") or {})
        self.assertGreaterEqual(int(self.stats.get("text_hits") or 0), 1)

    def test_output_rescan_zero(self):
        data = _part_bytes(self.out, self.part)
        self.assertEqual(_scan_marks(data), [], "成品仍有视觉/码类命中")
        self.assertEqual(_text_hits(data, Path(self.part).suffix), 0,
                         "成品仍有文字敏感命中")
        ver = pipeline.verify(self.out, [], self.parser, _ctx())
        self.assertTrue(ver.get("verified"))


class TestTextBudget(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _text_ocr_on()

    @classmethod
    def tearDownClass(cls):
        _text_ocr_off()

    def test_seconds_budget_skips_remaining(self):
        source = Fixtures.docx_text()
        parser = _parser(source)
        doc = parser.parse(source, _ctx())
        saved = os.environ.get("YINZHAN_OFFICE_TEXT_MAX_SECONDS")
        os.environ["YINZHAN_OFFICE_TEXT_MAX_SECONDS"] = "0.001"
        try:
            out = Fixtures.get() / "docx_text_budget_out.docx"
            res = pipeline.export(doc, parser, [], out, {"mode": "human"})
        finally:
            if saved is None:
                os.environ.pop("YINZHAN_OFFICE_TEXT_MAX_SECONDS", None)
            else:
                os.environ["YINZHAN_OFFICE_TEXT_MAX_SECONDS"] = saved
        stats = (doc.extra or {}).get("office_media") or {}
        self.assertGreaterEqual(int(stats.get("text_skipped") or 0), 1)
        self.assertIn("文字 OCR 预算", " ".join(res.warnings))

    def test_image_count_budget(self):
        source = Fixtures.docx_text()
        parser = _parser(source)
        doc = parser.parse(source, _ctx())
        saved = os.environ.get("YINZHAN_OFFICE_TEXT_MAX_IMAGES")
        os.environ["YINZHAN_OFFICE_TEXT_MAX_IMAGES"] = "1"
        try:
            out = Fixtures.get() / "docx_text_budget2_out.docx"
            pipeline.export(doc, parser, [], out, {"mode": "human"})
        finally:
            if saved is None:
                os.environ.pop("YINZHAN_OFFICE_TEXT_MAX_IMAGES", None)
            else:
                os.environ["YINZHAN_OFFICE_TEXT_MAX_IMAGES"] = saved
        stats = (doc.extra or {}).get("office_media") or {}
        self.assertGreaterEqual(int(stats.get("text_skipped") or 0), 1)


class TestTextSelfCheckBlocks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        _text_ocr_on()

    @classmethod
    def tearDownClass(cls):
        _text_ocr_off()

    def test_export_blocks_when_text_residual(self):
        source = Fixtures.docx_text()
        parser = _parser(source)
        doc = parser.parse(source, _ctx())
        original = docx_mod._office_redact_text_bytes
        # 模拟“涂黑无效”：helper 永远报告命中但返回原图字节 → 自检必须拦截
        docx_mod._office_redact_text_bytes = (  # type: ignore[assignment]
            lambda data, part, workdir: (data, 1, []))
        try:
            out = Fixtures.get() / "docx_text_selfcheck_out.docx"
            with self.assertRaises(Exception) as caught:
                pipeline.export(doc, parser, [], out, {"mode": "human"})
            self.assertIn("文字复检失败", str(caught.exception))
            self.assertFalse(out.exists(), "文字自检失败时应删除成品")
        finally:
            docx_mod._office_redact_text_bytes = original  # type: ignore[assignment]


# ────────────────────────── main ──────────────────────────

def main() -> int:
    print("数据隐站 · Office 内嵌媒体自动脱敏 P1-B 测试")
    print("-" * 64)
    print(f"fixtures={FIXTURE_DIR if os.environ.get('YINZHAN_WRITE_FIXTURES') == '1' else 'tempdir'}")
    suite = unittest.defaultTestLoader.loadTestsFromModule(sys.modules[__name__])
    result = unittest.TextTestRunner(verbosity=2).run(suite)
    print("-" * 64)
    if result.wasSuccessful():
        print(f"通过：{result.testsRun} 项（skipped={len(result.skipped)}）")
        return 0
    print(f"失败：failures={len(result.failures)} errors={len(result.errors)}")
    return 1


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception:  # pragma: no cover
        import traceback

        traceback.print_exc()
        sys.exit(3)
